import openpyxl, shutil, os

def apply_updates():
    targets = [
        ('Gastos/2027-Budget-Codiagro P&L.xlsx', 'PL Mensual 2027'),
        ('Gastos/Desglose 2027 Presupuestos V5.xlsx', 'PL 2027 Desglose V5'),
        ('Desglose 2027 Presupuestos V5.xlsx', 'PL 2027 Desglose V5')
    ]

    # Formacion: 53.215 repartido mensual
    # 11 meses de -4434.58 y diciembre -4434.62
    formacion_vals = [-4434.58]*11 + [-4434.62]
    assert abs(sum(formacion_vals) - (-53215.0)) < 0.01

    # Rules for commercials:
    # Ene: 0.75, Feb-Jul: 1.0, Ago: 0.5, Sep-Nov: 1.0, Dic: 0.75
    factors = [0.75, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 0.5, 1.0, 1.0, 1.0, 0.75]
    assert sum(factors) == 11.0

    def calc_com_months(base_monthly):
        # Expenses are negative
        return [-round(base_monthly * f, 2) for f in factors]

    commercial_bases = {
        '629100001': (2000.0, 'Pere Porta'),
        '629100002': (2000.0, 'Ricardo Perez'),
        '629100004': (2000.0, 'Garcia'),
        '629100003': (1200.0, 'Javier Paredes'),
        '629100005': (600.0, 'Irene'),
        '629100008': (600.0, 'Pedro Medina')
    }

    for file_path, sheet_name in targets:
        if not os.path.exists(file_path):
            print(f"Skipping non-existent: {file_path}")
            continue

        print(f"\nProcessing {file_path} [{sheet_name}]...")
        # Create a backup
        backup_path = file_path.replace('.xlsx', '_BACKUP_BEFORE_EXPENSES_UPDATE.xlsx')
        if not os.path.exists(backup_path):
            shutil.copy2(file_path, backup_path)
            print(f"Backup created: {backup_path}")

        wb = openpyxl.load_workbook(file_path, data_only=False)
        if sheet_name not in wb.sheetnames:
            print(f"Sheet {sheet_name} not found in {file_path}!")
            continue

        ws = wb[sheet_name]
        updated_ctas = []

        for r in range(1, ws.max_row + 1):
            val_col9 = str(ws.cell(r, 9).value or ws.cell(r, 1).value or '')
            
            # Check Formacion
            if '629000013' in val_col9:
                for idx, v in enumerate(formacion_vals):
                    ws.cell(r, 11 + idx).value = v
                ws.cell(r, 23).value = sum(formacion_vals)
                ws.cell(r, 24).value = 'Presupuesto anual formación no bonificada (53.215,00 €) con reparto mensual lineal homogéneo.'
                updated_ctas.append(('629000013', sum(formacion_vals)))

            # Check Commercials
            for cta_code, (base_amt, name) in commercial_bases.items():
                if cta_code in val_col9:
                    m_vals = calc_com_months(base_amt)
                    for idx, v in enumerate(m_vals):
                        ws.cell(r, 11 + idx).value = v
                    tot_year = sum(m_vals)
                    ws.cell(r, 23).value = tot_year
                    ws.cell(r, 24).value = f'Gastos comerciales {name}: {base_amt:,.2f} €/mes (Agosto al 50%, Enero y Diciembre al 75%). Total año: {abs(tot_year):,.2f} €.'
                    updated_ctas.append((cta_code, tot_year))

        wb.save(file_path)
        print(f"Successfully saved {file_path}! Updated accounts: {updated_ctas}")

if __name__ == '__main__':
    apply_updates()
