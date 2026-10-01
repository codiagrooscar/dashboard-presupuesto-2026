import os
import re
import shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import pandas as pd
import unicodedata
from datetime import datetime

base_dir = r"c:\Users\oscar.ocampo\OneDrive - Agroquímica Codiagro S.L\Escritorio\Presupuesto 2026\Calsif Comerciales"
os.chdir(base_dir)

def norm(s):
    if not isinstance(s, str): return ''
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn').upper().strip()

def map_client_name(c):
    if not isinstance(c, str): return c
    cn = norm(c)
    if 'DOGA TARIM' in cn:
        return 'ÇİTAR ÇİÇEK TARIM TIC. LTD. STI'
    return c

print("1. Cargando datos de presupuestos y pedidos reales...")

# 1. Cargar Budget original sin aplanar
src_budget = 'Presupuesto_Ventas_2027_Original_Sin_Aplanar.xlsx' if os.path.exists('Presupuesto_Ventas_2027_Original_Sin_Aplanar.xlsx') else 'Presupuesto_Ventas_2027.xlsx'
shutil.copy2(src_budget, 'temp_budget_input.xlsx')
df_bud = pd.read_excel('temp_budget_input.xlsx', sheet_name='Previsión Matriz Horizontal')
df_bud = df_bud[df_bud['Comercial'].notna() & (~df_bud['Comercial'].astype(str).str.contains('TOTAL', case=False))].copy()
df_bud['Cliente'] = df_bud['Cliente'].apply(map_client_name)
df_bud['CLIENTE_NORM'] = df_bud['Cliente'].apply(norm)

def clean_budget_sku(row):
    sku = str(row['Código Artículo']).strip().upper()
    pais = str(row['País']).strip().upper()
    if pais != 'ESPAÑA' and len(sku) > 2:
        return sku[:-2]
    return sku

df_bud['SKU_CLEAN'] = df_bud.apply(clean_budget_sku, axis=1)

client_map = df_bud.groupby('CLIENTE_NORM')['Comercial'].first().to_dict()
client_map[norm('FINCA DOÑA ANA C.B.')] = 'Javier'
client_map[norm('FITOSANITARIOS CARCAIXENT, S.L.')] = 'Javier'
client_map[norm('ALMENDRALIA IBÉRICA, S.L.U.')] = 'Javier'
client_map[norm('TÉCNICAS AGRÍCOLAS, S.A.')] = 'Ricardo'

# 2. Cargar Pedidos reales de septiembre
def get_latest_pedidos_file():
    candidates = [
        f for f in os.listdir('.')
        if f.startswith('Pedidos') and f.endswith('.xlsx') and not f.startswith(('temp_', '~$'))
    ]
    if not candidates:
        return 'Pedidos 01.10 comerciales.xlsx'
    def sort_key(f):
        m = re.search(r'(\d{1,2})\.(\d{1,2})', f)
        if m:
            return (int(m.group(2)), int(m.group(1)), os.path.getmtime(f))
        return (0, 0, os.path.getmtime(f))
    candidates.sort(key=sort_key, reverse=True)
    return candidates[0]

pedidos_file = get_latest_pedidos_file()
print(f"Usando pedidos de cierre: {pedidos_file}")
shutil.copy2(pedidos_file, 'temp_pedidos_input.xlsx')
df_ped = pd.read_excel('temp_pedidos_input.xlsx')
df_ped['Cliente'] = df_ped['Cliente'].apply(map_client_name)
df_ped['CLIENTE_NORM'] = df_ped['Cliente'].apply(norm)
df_ped['SKU_CLEAN'] = df_ped['CodigoArticulo'].astype(str).str.strip().str.upper()
df_ped = df_ped[~df_ped['CLIENTE_NORM'].str.contains('SUSTAINABLE', case=False, na=False)].copy()
df_ped['Comercial'] = df_ped['CLIENTE_NORM'].map(client_map).fillna('Sin Asignar')

comerciales = ['Alfonso', 'García', 'Irene', 'Javier', 'Mehmet', 'Pedro', 'Ricardo']

# Estilos Openpyxl
font_title = Font(name='Segoe UI', size=15, bold=True, color='FFFFFF')
font_subtitle = Font(name='Segoe UI', size=10, italic=True, color='94A3B8')
font_instructions = Font(name='Segoe UI', size=9, bold=True, color='1E293B')

font_kpi_lbl = Font(name='Segoe UI', size=9, bold=True, color='64748B')
font_group_hdr = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
font_col_hdr = Font(name='Segoe UI', size=9, bold=True, color='FFFFFF')
font_regular = Font(name='Segoe UI', size=9, color='1E293B')
font_bold = Font(name='Segoe UI', size=9, bold=True, color='0F172A')
font_editable = Font(name='Segoe UI', size=9, bold=True, color='065F46')
font_total = Font(name='Segoe UI', size=10, bold=True, color='0F172A')

fill_dark = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
fill_subbanner = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
fill_instructions = PatternFill(start_color='EFF6FF', end_color='EFF6FF', fill_type='solid')

fill_grp_info = PatternFill(start_color='334155', end_color='334155', fill_type='solid')
fill_grp_sep = PatternFill(start_color='1D4ED8', end_color='1D4ED8', fill_type='solid')
fill_grp_q4 = PatternFill(start_color='047857', end_color='047857', fill_type='solid')
fill_grp_tot = PatternFill(start_color='475569', end_color='475569', fill_type='solid')

fill_col_hdr = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
fill_col_sep = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
fill_col_q4 = PatternFill(start_color='059669', end_color='059669', fill_type='solid')

fill_card_kpi = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
fill_card_accent = PatternFill(start_color='ECFDF5', end_color='ECFDF5', fill_type='solid')
fill_card_blue = PatternFill(start_color='EFF6FF', end_color='EFF6FF', fill_type='solid')
fill_card_warn = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')

fill_prev_hist = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
fill_editable = PatternFill(start_color='ECFDF5', end_color='ECFDF5', fill_type='solid')
fill_comment_sep = PatternFill(start_color='FEF9C3', end_color='FEF9C3', fill_type='solid')
fill_comment_q4 = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
fill_total_row = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')

border_thin = Border(
    left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1')
)
border_editable = Border(
    left=Side(style='thin', color='10B981'), right=Side(style='thin', color='10B981'),
    top=Side(style='thin', color='10B981'), bottom=Side(style='thin', color='10B981')
)
border_double = Border(
    top=Side(style='thin', color='94A3B8'), bottom=Side(style='double', color='0F172A'),
    left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1')
)

align_center = Alignment(horizontal='center', vertical='center')
align_left = Alignment(horizontal='left', vertical='center')
align_right = Alignment(horizontal='right', vertical='center')

# Validación de Comentarios Sep
dv_coment_sep = DataValidation(
    type="list",
    formula1='"Adelanto de pedidos,Atraso de pedidos,Pedido extra no previsto,Campaña habitual,Cancelado / Sin siembra,Otro motivo"',
    allow_blank=True
)
dv_coment_sep.error = 'Por favor seleccione un motivo de la lista o escriba su comentario'
dv_coment_sep.errorTitle = 'Motivo de Desviación'

out_dir = "Revision_Previsiones_Q4_2026"
os.makedirs(out_dir, exist_ok=True)

def populate_sheet_commercial(ws, c, is_individual=False):
    ws.views.sheetView[0].showGridLines = True
    
    sub_ped = df_ped[df_ped['Comercial'] == c]
    sub_bud = df_bud[df_bud['Comercial'] == c]
    
    # 1. Título y Banner
    ws.merge_cells('B2:V2')
    ws['B2'] = f"CODIAGRO · ACTUALIZACIÓN DE PREVISIONES COMERCIALES Q4 2026 (OCTUBRE - DICIEMBRE)"
    ws['B2'].font = font_title
    ws['B2'].fill = fill_dark
    ws['B2'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[2].height = 36
    
    ws.merge_cells('B3:V3')
    ws['B3'] = f"DELEGADO COMERCIAL: {c.upper()} | Cierre Septiembre Oficial (01/10/2026) vs Previsiones a Final de Año | Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws['B3'].font = font_subtitle
    ws['B3'].fill = fill_subbanner
    ws['B3'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[3].height = 20
    
    # 2. Caja de Instrucciones
    ws.merge_cells('B4:V4')
    ws['B4'] = "📌 INSTRUCCIONES: 1) Revise el bloque de SEPTIEMBRE (Previsto, Real, Error Absoluto y Forecast Accuracy) e indique en col. L el motivo de la desviación (adelanto, atraso, extra). 2) En las columnas verdes (N, P, R), revise y actualice las cantidades estimadas de Octubre, Noviembre y Diciembre si prevé cambios. 3) Comente en col. V cualquier nota de campaña."
    ws['B4'].font = font_instructions
    ws['B4'].fill = fill_instructions
    ws['B4'].alignment = Alignment(horizontal='left', vertical='center', indent=1, wrap_text=True)
    ws.row_dimensions[4].height = 26
    
    # 3. KPI Header Cards (Fila 5 y 6)
    kpis = [
        ('B', 'D', 'PREVISTO SEP (U)', None, fill_card_blue, '1D4ED8', 'F'),
        ('E', 'G', 'REAL PEDIDOS SEP (U)', None, fill_card_accent, '047857', 'G'),
        ('H', 'I', '% CONSECUCIÓN SEP', None, fill_card_kpi, '0F172A', None),
        ('J', 'L', 'FORECAST ACCURACY SEP', None, fill_card_kpi, '0F172A', None),
        ('M', 'O', 'PREVISIÓN INICIAL Q4 (U)', None, fill_card_blue, '1D4ED8', 'S'),
        ('P', 'R', 'PREVISIÓN REVISADA Q4 (U)', None, fill_card_accent, '047857', 'T'),
        ('S', 'V', 'VARIACIÓN NETA Q4 (U)', None, fill_card_warn, 'B45309', 'U')
    ]
    
    ws.row_dimensions[5].height = 16
    ws.row_dimensions[6].height = 26
    
    # Construcción de datos por línea
    keys_set = set()
    for _, r in sub_bud.iterrows():
        keys_set.add((r['Cliente'], r['SKU_CLEAN']))
    for _, r in sub_ped.iterrows():
        keys_set.add((r['Cliente'], r['SKU_CLEAN']))
        
    keys_sorted = sorted(list(keys_set), key=lambda x: (x[0], x[1]))
    
    # Cabeceras Agrupadas (Fila 8)
    ws.merge_cells('B8:E8')
    ws['B8'] = "DATOS CLIENTE Y ARTÍCULO"
    ws['B8'].font = font_group_hdr
    ws['B8'].fill = fill_grp_info
    ws['B8'].alignment = align_center
    
    ws.merge_cells('F8:L8')
    ws['F8'] = "CIERRE SEPTIEMBRE 2026 (CALIBRACIÓN Y ANÁLISIS DE EXACTITUD)"
    ws['F8'].font = font_group_hdr
    ws['F8'].fill = fill_grp_sep
    ws['F8'].alignment = align_center
    
    ws.merge_cells('M8:R8')
    ws['M8'] = "REVISIÓN PREVISIONES MES A MES (HASTA FINAL DE AÑO 2026)"
    ws['M8'].font = font_group_hdr
    ws['M8'].fill = fill_grp_q4
    ws['M8'].alignment = align_center
    
    ws.merge_cells('S8:V8')
    ws['S8'] = "IMPACTO Q4 Y JUSTIFICACIÓN"
    ws['S8'].font = font_group_hdr
    ws['S8'].fill = fill_grp_tot
    ws['S8'].alignment = align_center
    
    ws.row_dimensions[8].height = 22
    
    # Cabeceras Columnas (Fila 9)
    col_headers = [
        ('B', 'País', fill_col_hdr),
        ('C', 'Cliente', fill_col_hdr),
        ('D', 'Cód. SKU', fill_col_hdr),
        ('E', 'Descripción Artículo', fill_col_hdr),
        ('F', 'Sep Previsto (u)', fill_col_sep),
        ('G', 'Sep Real (u)', fill_col_sep),
        ('H', 'Desv. Sep (u)', fill_col_sep),
        ('I', 'Error Abs. Sep (u)', fill_col_sep),
        ('J', '% Consec.', fill_col_sep),
        ('K', 'Forecast Accuracy', fill_col_sep),
        ('L', 'Comentario Sep (¿Adelanto/Atraso/Extra?)', fill_col_sep),
        ('M', 'Oct Anterior (u)', fill_col_q4),
        ('N', 'Oct Revisado (u)', fill_col_q4),
        ('O', 'Nov Anterior (u)', fill_col_q4),
        ('P', 'Nov Revisado (u)', fill_col_q4),
        ('Q', 'Dic Anterior (u)', fill_col_q4),
        ('R', 'Dic Revisado (u)', fill_col_q4),
        ('S', 'Total Q4 Ant (u)', fill_col_hdr),
        ('T', 'Total Q4 Rev (u)', fill_col_hdr),
        ('U', 'Dif Q4 (u)', fill_col_hdr),
        ('V', 'Comentarios / Justificación Q4', fill_col_hdr)
    ]
    
    for c_letter, h_text, fill_c in col_headers:
        c_cell = ws[f"{c_letter}9"]
        c_cell.value = h_text
        c_cell.font = font_col_hdr
        c_cell.fill = fill_c
        c_cell.alignment = align_center
        c_cell.border = border_thin
    ws.row_dimensions[9].height = 28
    
    row_num = 10
    start_data_row = 10
    
    for cli, sku in keys_sorted:
        b_m = sub_bud[(sub_bud['Cliente'] == cli) & (sub_bud['SKU_CLEAN'] == sku)]
        p_m = sub_ped[(sub_ped['Cliente'] == cli) & (sub_ped['SKU_CLEAN'] == sku)]
        
        s_b = float(b_m['Sep-26 (u)'].sum()) if len(b_m) > 0 else 0.0
        s_r = float(p_m['Pedidas'].sum()) if len(p_m) > 0 else 0.0
        oct_ant = float(b_m['Oct-26 (u)'].sum()) if len(b_m) > 0 else 0.0
        nov_ant = float(b_m['Nov-26 (u)'].sum()) if len(b_m) > 0 else 0.0
        dic_ant = float(b_m['Dic-26 (u)'].sum()) if len(b_m) > 0 else 0.0
        
        # Si no hay previsto ni real en sep ni en q4, omitir
        if s_b == 0 and s_r == 0 and oct_ant == 0 and nov_ant == 0 and dic_ant == 0:
            continue
            
        desc = ""
        pais = "ESPAÑA"
        if len(b_m) > 0:
            if pd.notna(b_m['Descripción Artículo'].iloc[0]):
                desc = str(b_m['Descripción Artículo'].iloc[0])
            if pd.notna(b_m['País'].iloc[0]):
                pais = str(b_m['País'].iloc[0]).strip()
        elif len(p_m) > 0:
            if pd.notna(p_m['Articulo'].iloc[0]):
                desc = str(p_m['Articulo'].iloc[0])
                
        r_str = str(row_num)
        
        # B: País
        ws[f"B{r_str}"] = pais
        ws[f"B{r_str}"].font = font_regular
        ws[f"B{r_str}"].alignment = align_center
        ws[f"B{r_str}"].border = border_thin
        
        # C: Cliente
        ws[f"C{r_str}"] = cli
        ws[f"C{r_str}"].font = font_bold
        ws[f"C{r_str}"].alignment = align_left
        ws[f"C{r_str}"].border = border_thin
        
        # D: SKU
        ws[f"D{r_str}"] = sku
        ws[f"D{r_str}"].font = font_regular
        ws[f"D{r_str}"].alignment = align_center
        ws[f"D{r_str}"].border = border_thin
        
        # E: Desc
        ws[f"E{r_str}"] = desc
        ws[f"E{r_str}"].font = font_regular
        ws[f"E{r_str}"].alignment = align_left
        ws[f"E{r_str}"].border = border_thin
        
        # F: Sep Previsto (u)
        ws[f"F{r_str}"] = s_b
        ws[f"F{r_str}"].font = font_regular
        ws[f"F{r_str}"].alignment = align_right
        ws[f"F{r_str}"].number_format = '#,##0'
        ws[f"F{r_str}"].border = border_thin
        
        # G: Sep Real (u)
        ws[f"G{r_str}"] = s_r
        ws[f"G{r_str}"].font = font_bold
        ws[f"G{r_str}"].alignment = align_right
        ws[f"G{r_str}"].number_format = '#,##0'
        ws[f"G{r_str}"].border = border_thin
        
        # H: Desv Sep (u) = Real - Previsto
        ws[f"H{r_str}"] = f"=G{r_str}-F{r_str}"
        ws[f"H{r_str}"].font = font_regular
        ws[f"H{r_str}"].alignment = align_right
        ws[f"H{r_str}"].number_format = '+#,##0;-#,##0;0'
        ws[f"H{r_str}"].border = border_thin
        
        # I: Error Abs. Sep (u) = |Real - Previsto| = |(Ventas + Pendientes) - Estimación|
        ws[f"I{r_str}"] = f"=ABS(G{r_str}-F{r_str})"
        ws[f"I{r_str}"].font = font_regular
        ws[f"I{r_str}"].alignment = align_right
        ws[f"I{r_str}"].number_format = '#,##0'
        ws[f"I{r_str}"].border = border_thin

        # J: % Consecución = Real / Previsto
        ws[f"J{r_str}"] = f'=IF(F{r_str}>0, G{r_str}/F{r_str}, IF(G{r_str}>0, 1.0, 0.0))'
        ws[f"J{r_str}"].font = font_regular
        ws[f"J{r_str}"].alignment = align_right
        ws[f"J{r_str}"].number_format = '0.0%'
        ws[f"J{r_str}"].border = border_thin
        
        # K: Forecast Accuracy = 1 - ErrorAbs / Previsto
        ws[f"K{r_str}"] = f'=IF(F{r_str}>0, 1.0 - (I{r_str}/F{r_str}), IF(G{r_str}=0, 1.0, 0.0))'
        ws[f"K{r_str}"].font = font_bold
        ws[f"K{r_str}"].alignment = align_right
        ws[f"K{r_str}"].number_format = '0.0%'
        ws[f"K{r_str}"].border = border_thin
        
        # L: Comentario Sep
        ws[f"L{r_str}"] = "Pedido extra no previsto" if (s_b == 0 and s_r > 0) else ""
        ws[f"L{r_str}"].font = font_regular
        ws[f"L{r_str}"].fill = fill_comment_sep
        ws[f"L{r_str}"].alignment = align_left
        ws[f"L{r_str}"].border = border_thin
        
        # M: Oct Ant
        ws[f"M{r_str}"] = oct_ant
        ws[f"M{r_str}"].font = font_regular
        ws[f"M{r_str}"].fill = fill_prev_hist
        ws[f"M{r_str}"].alignment = align_right
        ws[f"M{r_str}"].number_format = '#,##0'
        ws[f"M{r_str}"].border = border_thin
        
        # N: Oct Rev (precargado con oct_ant)
        ws[f"N{r_str}"] = oct_ant
        ws[f"N{r_str}"].font = font_editable
        ws[f"N{r_str}"].fill = fill_editable
        ws[f"N{r_str}"].alignment = align_right
        ws[f"N{r_str}"].number_format = '#,##0'
        ws[f"N{r_str}"].border = border_editable
        
        # O: Nov Ant
        ws[f"O{r_str}"] = nov_ant
        ws[f"O{r_str}"].font = font_regular
        ws[f"O{r_str}"].fill = fill_prev_hist
        ws[f"O{r_str}"].alignment = align_right
        ws[f"O{r_str}"].number_format = '#,##0'
        ws[f"O{r_str}"].border = border_thin
        
        # P: Nov Rev
        ws[f"P{r_str}"] = nov_ant
        ws[f"P{r_str}"].font = font_editable
        ws[f"P{r_str}"].fill = fill_editable
        ws[f"P{r_str}"].alignment = align_right
        ws[f"P{r_str}"].number_format = '#,##0'
        ws[f"P{r_str}"].border = border_editable
        
        # Q: Dic Ant
        ws[f"Q{r_str}"] = dic_ant
        ws[f"Q{r_str}"].font = font_regular
        ws[f"Q{r_str}"].fill = fill_prev_hist
        ws[f"Q{r_str}"].alignment = align_right
        ws[f"Q{r_str}"].number_format = '#,##0'
        ws[f"Q{r_str}"].border = border_thin
        
        # R: Dic Rev
        ws[f"R{r_str}"] = dic_ant
        ws[f"R{r_str}"].font = font_editable
        ws[f"R{r_str}"].fill = fill_editable
        ws[f"R{r_str}"].alignment = align_right
        ws[f"R{r_str}"].number_format = '#,##0'
        ws[f"R{r_str}"].border = border_editable
        
        # S: Total Q4 Ant = M + O + Q
        ws[f"S{r_str}"] = f"=M{r_str}+O{r_str}+Q{r_str}"
        ws[f"S{r_str}"].font = font_regular
        ws[f"S{r_str}"].alignment = align_right
        ws[f"S{r_str}"].number_format = '#,##0'
        ws[f"S{r_str}"].border = border_thin
        
        # T: Total Q4 Rev = N + P + R
        ws[f"T{r_str}"] = f"=N{r_str}+P{r_str}+R{r_str}"
        ws[f"T{r_str}"].font = font_bold
        ws[f"T{r_str}"].alignment = align_right
        ws[f"T{r_str}"].number_format = '#,##0'
        ws[f"T{r_str}"].border = border_thin
        
        # U: Dif Q4 = T - S
        ws[f"U{r_str}"] = f"=T{r_str}-S{r_str}"
        ws[f"U{r_str}"].font = font_regular
        ws[f"U{r_str}"].alignment = align_right
        ws[f"U{r_str}"].number_format = '+#,##0;-#,##0;0'
        ws[f"U{r_str}"].border = border_thin
        
        # V: Comentario Q4
        ws[f"V{r_str}"] = ""
        ws[f"V{r_str}"].font = font_regular
        ws[f"V{r_str}"].fill = fill_comment_q4
        ws[f"V{r_str}"].alignment = align_left
        ws[f"V{r_str}"].border = border_thin
        
        ws.row_dimensions[row_num].height = 20
        row_num += 1
        
    last_data_row = row_num - 1
    
    # Fila de Totales
    tot_row = row_num
    tr_str = str(tot_row)
    
    ws.merge_cells(f"B{tr_str}:E{tr_str}")
    ws[f"B{tr_str}"] = f"TOTAL CONSOLIDADO - {c.upper()} (UNIDADES)"
    ws[f"B{tr_str}"].font = font_total
    ws[f"B{tr_str}"].fill = fill_total_row
    ws[f"B{tr_str}"].alignment = Alignment(horizontal='right', vertical='center', indent=1)
    
    for c_let in ['B', 'C', 'D', 'E']:
        ws[f"{c_let}{tr_str}"].border = border_double
        
    # Sumas de cantidades (F: Prev, G: Real, H: Desv, I: Error Absoluto, M, N, O, P, Q, R, S, T, U)
    for c_let in ['F', 'G', 'H', 'I', 'M', 'N', 'O', 'P', 'Q', 'R', 'S', 'T', 'U']:
        ws[f"{c_let}{tr_str}"] = f"=SUM({c_let}{start_data_row}:{c_let}{last_data_row})"
        ws[f"{c_let}{tr_str}"].font = font_total
        ws[f"{c_let}{tr_str}"].fill = fill_total_row
        ws[f"{c_let}{tr_str}"].alignment = align_right
        ws[f"{c_let}{tr_str}"].border = border_double
        ws[f"{c_let}{tr_str}"].number_format = '+#,##0;-#,##0;0' if c_let in ['H', 'U'] else '#,##0'
        
    # Consecución Total
    ws[f"J{tr_str}"] = f'=IF(F{tr_str}>0, G{tr_str}/F{tr_str}, 1.0)'
    ws[f"J{tr_str}"].font = font_total
    ws[f"J{tr_str}"].fill = fill_total_row
    ws[f"J{tr_str}"].alignment = align_right
    ws[f"J{tr_str}"].number_format = '0.0%'
    ws[f"J{tr_str}"].border = border_double
    
    # FÓRMULA OFICIAL DEL JEFE: 1 - SUM(|Ventas + Pendientes - Estimación|) / SUM(Estimación)
    # I_tot es SUM(Error Absoluto), F_tot es SUM(Estimación)
    ws[f"K{tr_str}"] = f'=IF(F{tr_str}>0, 1.0 - (I{tr_str}/F{tr_str}), 1.0)'
    ws[f"K{tr_str}"].font = font_total
    ws[f"K{tr_str}"].fill = fill_total_row
    ws[f"K{tr_str}"].alignment = align_right
    ws[f"K{tr_str}"].number_format = '0.0%'
    ws[f"K{tr_str}"].border = border_double
    
    # Vacíos para comentarios en total
    for c_let in ['L', 'V']:
        ws[f"{c_let}{tr_str}"].border = border_double
        ws[f"{c_let}{tr_str}"].fill = fill_total_row
        
    ws.row_dimensions[tot_row].height = 25
    
    # Llenar fórmulas en KPI Cards de cabecera
    # B:D -> F_tot
    ws.merge_cells("B5:D5")
    ws.merge_cells("B6:D6")
    ws['B5'] = "PREVISTO SEP 2026 (U)"
    ws['B5'].font = font_kpi_lbl
    ws['B5'].fill = fill_card_blue
    ws['B5'].alignment = align_center
    ws['B6'] = f"=F{tr_str}"
    ws['B6'].font = Font(name='Segoe UI', size=13, bold=True, color='1D4ED8')
    ws['B6'].fill = fill_card_blue
    ws['B6'].alignment = align_center
    ws['B6'].number_format = '#,##0 "u"'
    
    # E:G -> G_tot
    ws.merge_cells("E5:G5")
    ws.merge_cells("E6:G6")
    ws['E5'] = "REAL PEDIDOS SEP 2026 (U)"
    ws['E5'].font = font_kpi_lbl
    ws['E5'].fill = fill_card_accent
    ws['E5'].alignment = align_center
    ws['E6'] = f"=G{tr_str}"
    ws['E6'].font = Font(name='Segoe UI', size=13, bold=True, color='047857')
    ws['E6'].fill = fill_card_accent
    ws['E6'].alignment = align_center
    ws['E6'].number_format = '#,##0 "u"'
    
    # H:I -> J_tot (% Consecución)
    ws.merge_cells("H5:I5")
    ws.merge_cells("H6:I6")
    ws['H5'] = "% CONSECUCIÓN SEP"
    ws['H5'].font = font_kpi_lbl
    ws['H5'].fill = fill_card_kpi
    ws['H5'].alignment = align_center
    ws['H6'] = f"=J{tr_str}"
    ws['H6'].font = Font(name='Segoe UI', size=13, bold=True, color='0F172A')
    ws['H6'].fill = fill_card_kpi
    ws['H6'].alignment = align_center
    ws['H6'].number_format = '0.0%'
    
    # J:L -> K_tot (Forecast Accuracy Oficial Dirección)
    ws.merge_cells("J5:L5")
    ws.merge_cells("J6:L6")
    ws['J5'] = "FORECAST ACCURACY SEP (OFICIAL)"
    ws['J5'].font = font_kpi_lbl
    ws['J5'].fill = fill_card_kpi
    ws['J5'].alignment = align_center
    ws['J6'] = f"=K{tr_str}"
    ws['J6'].font = Font(name='Segoe UI', size=13, bold=True, color='0F172A')
    ws['J6'].fill = fill_card_kpi
    ws['J6'].alignment = align_center
    ws['J6'].number_format = '0.0%'
    
    # M:O -> S_tot (Total Q4 Ant)
    ws.merge_cells("M5:O5")
    ws.merge_cells("M6:O6")
    ws['M5'] = "PREVISIÓN INICIAL Q4 (U)"
    ws['M5'].font = font_kpi_lbl
    ws['M5'].fill = fill_card_blue
    ws['M5'].alignment = align_center
    ws['M6'] = f"=S{tr_str}"
    ws['M6'].font = Font(name='Segoe UI', size=13, bold=True, color='1D4ED8')
    ws['M6'].fill = fill_card_blue
    ws['M6'].alignment = align_center
    ws['M6'].number_format = '#,##0 "u"'
    
    # P:R -> T_tot (Total Q4 Rev)
    ws.merge_cells("P5:R5")
    ws.merge_cells("P6:R6")
    ws['P5'] = "PREVISIÓN REVISADA Q4 (U)"
    ws['P5'].font = font_kpi_lbl
    ws['P5'].fill = fill_card_accent
    ws['P5'].alignment = align_center
    ws['P6'] = f"=T{tr_str}"
    ws['P6'].font = Font(name='Segoe UI', size=13, bold=True, color='047857')
    ws['P6'].fill = fill_card_accent
    ws['P6'].alignment = align_center
    ws['P6'].number_format = '#,##0 "u"'
    
    # S:V -> U_tot (Diferencia Q4)
    ws.merge_cells("S5:V5")
    ws.merge_cells("S6:V6")
    ws['S5'] = "VARIACIÓN NETA Q4 (U)"
    ws['S5'].font = font_kpi_lbl
    ws['S5'].fill = fill_card_warn
    ws['S5'].alignment = align_center
    ws['S6'] = f"=U{tr_str}"
    ws['S6'].font = Font(name='Segoe UI', size=13, bold=True, color='B45309')
    ws['S6'].fill = fill_card_warn
    ws['S6'].alignment = align_center
    ws['S6'].number_format = '+#,##0 "u";-#,##0 "u";0 "u"'
    
    for c1, c2 in [('B','D'), ('E','G'), ('H','I'), ('J','L'), ('M','O'), ('P','R'), ('S','V')]:
        start_c = ord(c1)
        end_c = ord(c2)
        for cc in range(start_c, end_c + 1):
            letter = chr(cc)
            ws[f"{letter}5"].border = border_thin
            ws[f"{letter}6"].border = border_thin
            
    # Añadir Data Validation a columna L (Comentario Sep)
    ws.add_data_validation(dv_coment_sep)
    dv_coment_sep.add(f"L{start_data_row}:L{last_data_row}")
    
    # Anchos de columna optimizados para visualización perfecta
    col_widths = {
        'A': 3,
        'B': 14, # País
        'C': 34, # Cliente
        'D': 14, # SKU
        'E': 32, # Desc
        'F': 16, # Sep Prev
        'G': 15, # Sep Real
        'H': 14, # Desv Sep
        'I': 16, # Error Abs. Sep
        'J': 13, # % Consec
        'K': 17, # Accuracy Oficial
        'L': 32, # Coment Sep
        'M': 15, # Oct Ant
        'N': 16, # Oct Rev
        'O': 15, # Nov Ant
        'P': 16, # Nov Rev
        'Q': 15, # Dic Ant
        'R': 16, # Dic Rev
        'S': 17, # Total Q4 Ant
        'T': 17, # Total Q4 Rev
        'U': 15, # Dif Q4
        'V': 34  # Coment Q4
    }
    for col_l, w in col_widths.items():
        ws.column_dimensions[col_l].width = w

print("2. Generando archivos individuales para cada comercial...")

for c in comerciales:
    wb_ind = openpyxl.Workbook()
    wb_ind.remove(wb_ind.active)
    ws_ind = wb_ind.create_sheet(title=f"Previsiones_{c}")
    populate_sheet_commercial(ws_ind, c, is_individual=True)
    
    # Guardar archivo individual
    safe_name = norm(c).capitalize()
    filepath = os.path.join(out_dir, f"Revision_Previsiones_Q4_2026_{safe_name}.xlsx")
    wb_ind.save(filepath)
    print(f" -> Guardado: {filepath}")

print("3. Generando archivo CONSOLIDADO Master para Dirección...")

wb_con = openpyxl.Workbook()
wb_con.remove(wb_con.active)

# Hoja 1: Resumen General Consolidado
ws_resumen = wb_con.create_sheet(title="Resumen Consolidado Dirección")
ws_resumen.views.sheetView[0].showGridLines = True

ws_resumen.merge_cells('B2:N2')
ws_resumen['B2'] = "CODIAGRO · RESUMEN DIRECCIÓN: REVISIÓN DE PREVISIONES Q4 2026"
ws_resumen['B2'].font = font_title
ws_resumen['B2'].fill = fill_dark
ws_resumen['B2'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
ws_resumen.row_dimensions[2].height = 36

ws_resumen.merge_cells('B3:N3')
ws_resumen['B3'] = f"Consolidado de Desviaciones de Septiembre y Previsiones hasta Final de Año (Octubre a Diciembre) | Cierre Sep: 01/10/2026"
ws_resumen['B3'].font = font_subtitle
ws_resumen['B3'].fill = fill_subbanner
ws_resumen['B3'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
ws_resumen.row_dimensions[3].height = 20

# Resumen tabla cabeceras (13 columnas: B a N)
res_headers = [
    ('B', 'Comercial', fill_col_hdr),
    ('C', 'Clientes Activos', fill_col_hdr),
    ('D', 'Sep Previsto (u)', fill_col_sep),
    ('E', 'Sep Real (u)', fill_col_sep),
    ('F', 'Desv. Sep (u)', fill_col_sep),
    ('G', 'Error Abs. Sep (u)', fill_col_sep),
    ('H', '% Consec.', fill_col_sep),
    ('I', 'Forecast Accuracy', fill_col_sep),
    ('J', 'Oct Previsto (u)', fill_col_q4),
    ('K', 'Nov Previsto (u)', fill_col_q4),
    ('L', 'Dic Previsto (u)', fill_col_q4),
    ('M', 'Total Q4 Previsto (u)', fill_col_q4),
    ('N', 'Estado Cierre Sep', fill_col_hdr)
]

ws_resumen.merge_cells('B5:C5')
ws_resumen['B5'] = "INFORMACIÓN COMERCIAL"
ws_resumen['B5'].font = font_group_hdr
ws_resumen['B5'].fill = fill_grp_info
ws_resumen['B5'].alignment = align_center

ws_resumen.merge_cells('D5:I5')
ws_resumen['D5'] = "CIERRE SEPTIEMBRE 2026 (REAL vs PREVISTO Y EXACTITUD)"
ws_resumen['D5'].font = font_group_hdr
ws_resumen['D5'].fill = fill_grp_sep
ws_resumen['D5'].alignment = align_center

ws_resumen.merge_cells('J5:M5')
ws_resumen['J5'] = "PREVISIÓN CARTERA HASTA FINAL DE AÑO (Q4 2026)"
ws_resumen['J5'].font = font_group_hdr
ws_resumen['J5'].fill = fill_grp_q4
ws_resumen['J5'].alignment = align_center

ws_resumen.cell(5, 14).value = "VALIDACIÓN"
ws_resumen.cell(5, 14).font = font_group_hdr
ws_resumen.cell(5, 14).fill = fill_grp_tot
ws_resumen.cell(5, 14).alignment = align_center
ws_resumen.row_dimensions[5].height = 22

for c_l, h_t, f_c in res_headers:
    cell = ws_resumen[f"{c_l}6"]
    cell.value = h_t
    cell.font = font_col_hdr
    cell.fill = f_c
    cell.alignment = align_center
    cell.border = border_thin
ws_resumen.row_dimensions[6].height = 26

r_res = 7
for c in comerciales:
    s_c = f"='{c}'!"
    r_str = str(r_res)
    
    ws_resumen[f"B{r_str}"] = c
    ws_resumen[f"B{r_str}"].font = font_bold
    ws_resumen[f"B{r_str}"].alignment = align_left
    ws_resumen[f"B{r_str}"].border = border_thin
    
    # Clientes únicos
    n_cli = df_bud[df_bud['Comercial'] == c]['Cliente'].nunique()
    ws_resumen[f"C{r_str}"] = n_cli
    ws_resumen[f"C{r_str}"].font = font_regular
    ws_resumen[f"C{r_str}"].alignment = align_center
    ws_resumen[f"C{r_str}"].border = border_thin
    
    # Links a fórmulas del total de la pestaña
    sub_ped = df_ped[df_ped['Comercial'] == c]
    sub_bud = df_bud[df_bud['Comercial'] == c]
    keys_c = set()
    for _, r in sub_bud.iterrows(): keys_c.add((r['Cliente'], r['SKU_CLEAN']))
    for _, r in sub_ped.iterrows(): keys_c.add((r['Cliente'], r['SKU_CLEAN']))
    count_lines = 0
    for cli, sku in keys_c:
        bm = sub_bud[(sub_bud['Cliente'] == cli) & (sub_bud['SKU_CLEAN'] == sku)]
        pm = sub_ped[(sub_ped['Cliente'] == cli) & (sub_ped['SKU_CLEAN'] == sku)]
        sb = float(bm['Sep-26 (u)'].sum()) if len(bm) > 0 else 0
        sr = float(pm['Pedidas'].sum()) if len(pm) > 0 else 0
        o_ant = float(bm['Oct-26 (u)'].sum()) if len(bm) > 0 else 0
        n_ant = float(bm['Nov-26 (u)'].sum()) if len(bm) > 0 else 0
        d_ant = float(bm['Dic-26 (u)'].sum()) if len(bm) > 0 else 0
        if sb > 0 or sr > 0 or o_ant > 0 or n_ant > 0 or d_ant > 0:
            count_lines += 1
    t_row_sheet = 10 + count_lines
    
    # D: Sep Previsto (col F en hoja individual)
    ws_resumen[f"D{r_str}"] = f"{s_c}F{t_row_sheet}"
    ws_resumen[f"D{r_str}"].font = font_regular
    ws_resumen[f"D{r_str}"].alignment = align_right
    ws_resumen[f"D{r_str}"].number_format = '#,##0'
    ws_resumen[f"D{r_str}"].border = border_thin
    
    # E: Sep Real (col G en hoja individual)
    ws_resumen[f"E{r_str}"] = f"{s_c}G{t_row_sheet}"
    ws_resumen[f"E{r_str}"].font = font_bold
    ws_resumen[f"E{r_str}"].alignment = align_right
    ws_resumen[f"E{r_str}"].number_format = '#,##0'
    ws_resumen[f"E{r_str}"].border = border_thin
    
    # F: Desv Sep (col H en hoja individual)
    ws_resumen[f"F{r_str}"] = f"{s_c}H{t_row_sheet}"
    ws_resumen[f"F{r_str}"].font = font_regular
    ws_resumen[f"F{r_str}"].alignment = align_right
    ws_resumen[f"F{r_str}"].number_format = '+#,##0;-#,##0;0'
    ws_resumen[f"F{r_str}"].border = border_thin
    
    # G: Error Abs. Sep (col I en hoja individual)
    ws_resumen[f"G{r_str}"] = f"{s_c}I{t_row_sheet}"
    ws_resumen[f"G{r_str}"].font = font_regular
    ws_resumen[f"G{r_str}"].alignment = align_right
    ws_resumen[f"G{r_str}"].number_format = '#,##0'
    ws_resumen[f"G{r_str}"].border = border_thin
    
    # H: % Consec (col J en hoja individual)
    ws_resumen[f"H{r_str}"] = f"{s_c}J{t_row_sheet}"
    ws_resumen[f"H{r_str}"].font = font_bold
    ws_resumen[f"H{r_str}"].alignment = align_right
    ws_resumen[f"H{r_str}"].number_format = '0.0%'
    ws_resumen[f"H{r_str}"].border = border_thin
    
    # I: Accuracy Oficial Dirección (col K en hoja individual)
    ws_resumen[f"I{r_str}"] = f"{s_c}K{t_row_sheet}"
    ws_resumen[f"I{r_str}"].font = font_bold
    ws_resumen[f"I{r_str}"].alignment = align_right
    ws_resumen[f"I{r_str}"].number_format = '0.0%'
    ws_resumen[f"I{r_str}"].border = border_thin
    
    # J: Oct Previsto (col N en hoja individual)
    ws_resumen[f"J{r_str}"] = f"{s_c}N{t_row_sheet}"
    ws_resumen[f"J{r_str}"].font = font_editable
    ws_resumen[f"J{r_str}"].alignment = align_right
    ws_resumen[f"J{r_str}"].number_format = '#,##0'
    ws_resumen[f"J{r_str}"].border = border_thin
    
    # K: Nov Previsto (col P en hoja individual)
    ws_resumen[f"K{r_str}"] = f"{s_c}P{t_row_sheet}"
    ws_resumen[f"K{r_str}"].font = font_editable
    ws_resumen[f"K{r_str}"].alignment = align_right
    ws_resumen[f"K{r_str}"].number_format = '#,##0'
    ws_resumen[f"K{r_str}"].border = border_thin
    
    # L: Dic Previsto (col R en hoja individual)
    ws_resumen[f"L{r_str}"] = f"{s_c}R{t_row_sheet}"
    ws_resumen[f"L{r_str}"].font = font_editable
    ws_resumen[f"L{r_str}"].alignment = align_right
    ws_resumen[f"L{r_str}"].number_format = '#,##0'
    ws_resumen[f"L{r_str}"].border = border_thin
    
    # M: Total Q4 Previsto (col T en hoja individual)
    ws_resumen[f"M{r_str}"] = f"{s_c}T{t_row_sheet}"
    ws_resumen[f"M{r_str}"].font = font_bold
    ws_resumen[f"M{r_str}"].alignment = align_right
    ws_resumen[f"M{r_str}"].number_format = '#,##0'
    ws_resumen[f"M{r_str}"].border = border_thin
    
    # N: Estado
    ws_resumen[f"N{r_str}"] = f'=IF(H{r_str}>=1.0, "Superado (+)", IF(H{r_str}>=0.7, "En Rango", "Revisar"))'
    ws_resumen[f"N{r_str}"].font = font_bold
    ws_resumen[f"N{r_str}"].alignment = align_center
    ws_resumen[f"N{r_str}"].border = border_thin
    
    ws_resumen.row_dimensions[r_res].height = 22
    r_res += 1

tot_r_res = str(r_res)
ws_resumen[f"B{tot_r_res}"] = "TOTAL CODIAGRO"
ws_resumen[f"B{tot_r_res}"].font = font_total
ws_resumen[f"B{tot_r_res}"].fill = fill_total_row
ws_resumen[f"B{tot_r_res}"].alignment = Alignment(horizontal='left', vertical='center', indent=1)

ws_resumen[f"C{tot_r_res}"] = f"=SUM(C7:C{r_res-1})"
ws_resumen[f"C{tot_r_res}"].font = font_total
ws_resumen[f"C{tot_r_res}"].fill = fill_total_row
ws_resumen[f"C{tot_r_res}"].alignment = align_center
ws_resumen[f"C{tot_r_res}"].border = border_double

# Sumas de Previsto (D), Real (E), Desv (F), Error Absoluto (G), Oct (J), Nov (K), Dic (L), Total Q4 (M)
for cl in ['D', 'E', 'F', 'G', 'J', 'K', 'L', 'M']:
    ws_resumen[f"{cl}{tot_r_res}"] = f"=SUM({cl}7:{cl}{r_res-1})"
    ws_resumen[f"{cl}{tot_r_res}"].font = font_total
    ws_resumen[f"{cl}{tot_r_res}"].fill = fill_total_row
    ws_resumen[f"{cl}{tot_r_res}"].alignment = align_right
    ws_resumen[f"{cl}{tot_r_res}"].border = border_double
    ws_resumen[f"{cl}{tot_r_res}"].number_format = '+#,##0;-#,##0;0' if cl == 'F' else '#,##0'

# Total Consecución = Real / Previsto
ws_resumen[f"H{tot_r_res}"] = f'=E{tot_r_res}/D{tot_r_res}'
ws_resumen[f"H{tot_r_res}"].font = font_total
ws_resumen[f"H{tot_r_res}"].fill = fill_total_row
ws_resumen[f"H{tot_r_res}"].alignment = align_right
ws_resumen[f"H{tot_r_res}"].border = border_double
ws_resumen[f"H{tot_r_res}"].number_format = '0.0%'

# Total Accuracy Oficial Dirección: 1 - SUM(|Ventas + Pendientes - Estimación|) / SUM(Estimación)
# G{tot_r_res} es la suma de los errores absolutos de todos los comerciales, D{tot_r_res} es la suma del presupuesto
ws_resumen[f"I{tot_r_res}"] = f'=1.0 - (G{tot_r_res}/D{tot_r_res})'
ws_resumen[f"I{tot_r_res}"].font = font_total
ws_resumen[f"I{tot_r_res}"].fill = fill_total_row
ws_resumen[f"I{tot_r_res}"].alignment = align_right
ws_resumen[f"I{tot_r_res}"].border = border_double
ws_resumen[f"I{tot_r_res}"].number_format = '0.0%'

ws_resumen[f"N{tot_r_res}"] = "CONSOLIDADO"
ws_resumen[f"N{tot_r_res}"].font = font_total
ws_resumen[f"N{tot_r_res}"].fill = fill_total_row
ws_resumen[f"N{tot_r_res}"].alignment = align_center
ws_resumen[f"N{tot_r_res}"].border = border_double
ws_resumen[f"B{tot_r_res}"].border = border_double

ws_resumen.row_dimensions[r_res].height = 25

res_widths = {'A': 3, 'B': 18, 'C': 16, 'D': 17, 'E': 17, 'F': 16, 'G': 17, 'H': 14, 'I': 17, 'J': 16, 'K': 16, 'L': 16, 'M': 19, 'N': 18}
for c_l, w in res_widths.items():
    ws_resumen.column_dimensions[c_l].width = w

# Añadir hojas de cada comercial al consolidado
for c in comerciales:
    ws_c = wb_con.create_sheet(title=c)
    populate_sheet_commercial(ws_c, c, is_individual=False)

consolidado_path = os.path.join(out_dir, "Revision_Previsiones_Q4_2026_CONSOLIDADO.xlsx")
wb_con.save(consolidado_path)
print(f" -> Guardado Consolidado Master: {consolidado_path}")
print("¡Proceso completado con éxito!")
