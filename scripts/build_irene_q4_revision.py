# -*- coding: utf-8 -*-
"""
scripts/build_irene_q4_revision.py
Genera la plantilla de revisión de previsiones Q4 para Irene:
- Octubre 2026 en curso: Real a fecha de hoy (27/10/2026) vs Previsión vigente
- Última versión de previsiones: Previsiones revisadas vigentes de Q4 (Oct 32.958 u, Nov 23.970 u, Dic 29.118 u)
- Columnas editables en verde para Cierre Estimado de Octubre, Noviembre Revisado y Diciembre Revisado
- Fórmulas automáticas de impacto Q4, validación de datos y KPIs directivos en cabecera
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import pandas as pd
import unicodedata
from datetime import datetime

base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(base_dir)

def norm(s):
    if not isinstance(s, str): return ''
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn').upper().strip()

print("1. Cargando previsiones vigentes de Irene y pedidos reales al 27/10...")

# 1. Cargar la última versión de previsiones de Irene
src_prev = 'temp_Copia de Revision_Previsiones_Q4_2026_Irene.xlsx'
wb_prev = openpyxl.load_workbook(src_prev, data_only=True)
ws_prev = wb_prev['Previsiones_Irene']

irene_lines = []
existing_keys = set()

for r in range(10, ws_prev.max_row):
    pais = ws_prev.cell(r, 2).value or 'ESPAÑA'
    cli = ws_prev.cell(r, 3).value
    sku = ws_prev.cell(r, 4).value
    desc = ws_prev.cell(r, 5).value or ''
    if not cli or not sku:
        continue
    
    # Previsiones revisadas vigentes (de su última entrega)
    oct_vig = float(ws_prev.cell(r, 12).value or 0.0)
    nov_vig = float(ws_prev.cell(r, 14).value or 0.0)
    dic_vig = float(ws_prev.cell(r, 16).value or 0.0)
    coment_ant = ws_prev.cell(r, 20).value or ''
    
    k = (norm(cli), norm(sku))
    existing_keys.add(k)
    
    irene_lines.append({
        'pais': pais,
        'cliente': cli,
        'sku': sku,
        'desc': desc,
        'oct_prev_vigente': oct_vig,
        'nov_prev_vigente': nov_vig,
        'dic_prev_vigente': dic_vig,
        'coment_ant': coment_ant,
        'es_extra': False
    })

print(f"Previsiones vigentes cargadas: {len(irene_lines)} líneas de Irene.")

# 2. Cargar pedidos reales de octubre al 27/10
df_ped = pd.read_excel('Pedidos27.10 comerciales.xlsx')
df_ped['CLI_NORM'] = df_ped['Cliente'].apply(norm)
df_ped['SKU_NORM'] = df_ped['CodigoArticulo'].apply(norm)

irene_clients_norm = [norm(c) for c in ['CAMPOEJIDO S.C.A.', 'COOP.AGRIC. SAN ISIDRO S.C.A.', 'GUIVARTO AGRICOLA, S.L.', 'SUCA, S.C.A.']]
df_ped_irene = df_ped[df_ped['CLI_NORM'].apply(lambda x: any(ic in x or x in ic for ic in irene_clients_norm))].copy()

# Mapear pedidos por (cliente_norm, sku_norm)
ped_map = {}
for _, row in df_ped_irene.iterrows():
    k = (row['CLI_NORM'], row['SKU_NORM'])
    if k not in ped_map:
        ped_map[k] = {
            'cliente_orig': row['Cliente'],
            'sku_orig': row['CodigoArticulo'],
            'desc_orig': row['Articulo'],
            'pedidas': 0.0,
            'servidas': 0.0,
            'pendientes': 0.0,
            'importe': 0.0
        }
    ped_map[k]['pedidas'] += float(row['Pedidas'] or 0.0)
    ped_map[k]['servidas'] += float(row['Servidas'] or 0.0)
    ped_map[k]['pendientes'] += float(row['Pendientes'] or 0.0)
    ped_map[k]['importe'] += float(row['ImporteNeto'] or 0.0)

# Asignar pedidos reales a las líneas existentes
matched_keys = set()
for item in irene_lines:
    k = (norm(item['cliente']), norm(item['sku']))
    if k in ped_map:
        item['oct_real'] = ped_map[k]['pedidas']
        item['oct_servidas'] = ped_map[k]['servidas']
        item['oct_pendientes'] = ped_map[k]['pendientes']
        item['oct_importe'] = ped_map[k]['importe']
        matched_keys.add(k)
    else:
        item['oct_real'] = 0.0
        item['oct_servidas'] = 0.0
        item['oct_pendientes'] = 0.0
        item['oct_importe'] = 0.0

# Incorporar pedidos extras de octubre no previstos
unmatched_count = 0
for k, pdata in ped_map.items():
    if k not in existing_keys:
        irene_lines.append({
            'pais': 'ESPAÑA',
            'cliente': pdata['cliente_orig'],
            'sku': pdata['sku_orig'],
            'desc': pdata['desc_orig'],
            'oct_prev_vigente': 0.0,
            'nov_prev_vigente': 0.0,
            'dic_prev_vigente': 0.0,
            'coment_ant': 'Pedido extra no previsto registrado en octubre',
            'oct_real': pdata['pedidas'],
            'oct_servidas': pdata['servidas'],
            'oct_pendientes': pdata['pendientes'],
            'oct_importe': pdata['importe'],
            'es_extra': True
        })
        unmatched_count += 1

print(f"Pedidos reales procesados: {len(ped_map)} combinaciones ({len(matched_keys)} coincidentes con presupuesto, {unmatched_count} pedidos extras añadidos).")
print(f"Total líneas finales para Irene: {len(irene_lines)}")

# Ordenar por Cliente y SKU
irene_lines.sort(key=lambda x: (x['cliente'], x['sku']))

# ==============================================================================
# CONSTRUCCIÓN DEL LIBRO EXCEL CON ESTILOS PROFESIONALES
# ==============================================================================
print("2. Construyendo archivo Excel para Irene...")
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Previsiones_Irene"
ws.views.sheetView[0].showGridLines = True

# Paleta y Estilos
font_title = Font(name='Segoe UI', size=15, bold=True, color='FFFFFF')
font_subtitle = Font(name='Segoe UI', size=10, italic=True, color='94A3B8')
font_instructions = Font(name='Segoe UI', size=9, bold=True, color='1E293B')

font_kpi_lbl = Font(name='Segoe UI', size=9, bold=True, color='64748B')
font_group_hdr = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
font_col_hdr = Font(name='Segoe UI', size=9, bold=True, color='FFFFFF')
font_regular = Font(name='Segoe UI', size=9, color='1E293B')
font_bold = Font(name='Segoe UI', size=9, bold=True, color='0F172A')
font_editable = Font(name='Segoe UI', size=9, bold=True, color='065F46')
font_total = Font(name='Segoe UI', size=10, bold=True, color='0F172A')

fill_dark = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
fill_subbanner = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
fill_instructions = PatternFill(start_color='EFF6FF', end_color='EFF6FF', fill_type='solid')

fill_grp_info = PatternFill(start_color='334155', end_color='334155', fill_type='solid')
fill_grp_oct = PatternFill(start_color='1D4ED8', end_color='1D4ED8', fill_type='solid')
fill_grp_q4 = PatternFill(start_color='047857', end_color='047857', fill_type='solid')
fill_grp_tot = PatternFill(start_color='475569', end_color='475569', fill_type='solid')

fill_col_hdr = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
fill_col_oct = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
fill_col_q4 = PatternFill(start_color='059669', end_color='059669', fill_type='solid')

fill_card_kpi = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
fill_card_accent = PatternFill(start_color='ECFDF5', end_color='ECFDF5', fill_type='solid')
fill_card_blue = PatternFill(start_color='EFF6FF', end_color='EFF6FF', fill_type='solid')
fill_card_warn = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')

fill_prev_hist = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
fill_editable = PatternFill(start_color='ECFDF5', end_color='ECFDF5', fill_type='solid')
fill_comment_oct = PatternFill(start_color='FEF9C3', end_color='FEF9C3', fill_type='solid')
fill_comment_q4 = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
fill_total_row = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')

border_thin = Border(
    left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1')
)
border_editable = Border(
    left=Side(style='thin', color='10B981'), right=Side(style='thin', color='10B981'),
    top=Side(style='thin', color='10B981'), bottom=Side(style='thin', color='10B981')
)
border_double = Border(
    top=Side(style='thin', color='94A3B8'), bottom=Side(style='double', color='0F172A'),
    left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1')
)

align_center = Alignment(horizontal='center', vertical='center')
align_left = Alignment(horizontal='left', vertical='center')
align_right = Alignment(horizontal='right', vertical='center')

# Validación de Comentarios de Octubre
dv_coment_oct = DataValidation(
    type="list",
    formula1='"Cierre según previsto,Pedidos pendientes de entrar,Atraso a Noviembre,Pedido extra no previsto,Sin más demanda este mes,Otro motivo"',
    allow_blank=True
)
dv_coment_oct.error = 'Por favor seleccione un motivo de la lista o escriba su comentario'
dv_coment_oct.errorTitle = 'Motivo de Desviación / Estado'

# 1. Título y Banner
ws.merge_cells('B2:T2')
ws['B2'] = "CODIAGRO · ACTUALIZACIÓN DE PREVISIONES COMERCIALES Q4 2026 (OCTUBRE - DICIEMBRE)"
ws['B2'].font = font_title
ws['B2'].fill = fill_dark
ws['B2'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
ws.row_dimensions[2].height = 36

now_str = datetime.now().strftime('%d/%m/%Y %H:%M')
ws.merge_cells('B3:T3')
ws['B3'] = f"DELEGADA COMERCIAL: IRENE | Octubre 2026 en Curso (Real al 27/10/2026) vs Previsiones Vigentes Q4 | Generado: {now_str}"
ws['B3'].font = font_subtitle
ws['B3'].fill = fill_subbanner
ws['B3'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
ws.row_dimensions[3].height = 20

# 2. Caja de Instrucciones
ws.merge_cells('B4:T4')
ws['B4'] = "📌 INSTRUCCIONES: 1) Revise el bloque de OCTUBRE EN CURSO (Previsión Vigente, Real al 27/10, Pendiente y % Consecución) e indique en col. J si prevé pedidos adicionales antes de fin de mes. 2) En las columnas verdes (L, N, P), revise y actualice el Cierre Estimado de Octubre, y las cantidades estimadas de Noviembre y Diciembre si prevé cambios. 3) Comente en col. T cualquier nota de campaña o justificación."
ws['B4'].font = font_instructions
ws['B4'].fill = fill_instructions
ws['B4'].alignment = Alignment(horizontal='left', vertical='center', indent=1, wrap_text=True)
ws.row_dimensions[4].height = 26

ws.row_dimensions[5].height = 16
ws.row_dimensions[6].height = 26

# Cabeceras Agrupadas (Fila 8)
ws.merge_cells('B8:E8')
ws['B8'] = "DATOS CLIENTE Y ARTÍCULO"
ws['B8'].font = font_group_hdr
ws['B8'].fill = fill_grp_info
ws['B8'].alignment = align_center

ws.merge_cells('F8:J8')
ws['F8'] = "OCTUBRE 2026 (EN CURSO - REAL AL 27/10/2026)"
ws['F8'].font = font_group_hdr
ws['F8'].fill = fill_grp_oct
ws['F8'].alignment = align_center

ws.merge_cells('K8:P8')
ws['K8'] = "REVISIÓN PREVISIONES MES A MES (HASTA FINAL DE AÑO 2026)"
ws['K8'].font = font_group_hdr
ws['K8'].fill = fill_grp_q4
ws['K8'].alignment = align_center

ws.merge_cells('Q8:T8')
ws['Q8'] = "IMPACTO Q4 Y JUSTIFICACIÓN"
ws['Q8'].font = font_group_hdr
ws['Q8'].fill = fill_grp_tot
ws['Q8'].alignment = align_center

ws.row_dimensions[8].height = 22

# Cabeceras de Columnas (Fila 9)
col_headers = [
    ('B', 'País', fill_col_hdr),
    ('C', 'Cliente', fill_col_hdr),
    ('D', 'Cód. SKU', fill_col_hdr),
    ('E', 'Descripción Artículo', fill_col_hdr),
    ('F', 'Oct Prev. Vigente (u)', fill_col_oct),
    ('G', 'Oct Real a 27/10 (u)', fill_col_oct),
    ('H', 'Pendiente / Desv. (u)', fill_col_oct),
    ('I', '% Consec. a 27/10', fill_col_oct),
    ('J', 'Comentario Octubre (¿Cierre previsto / Atraso / Extra?)', fill_col_oct),
    ('K', 'Oct Prev. Anterior (u)', fill_col_q4),
    ('L', 'Oct Cierre Estimado (u)', fill_col_q4),
    ('M', 'Nov Prev. Anterior (u)', fill_col_q4),
    ('N', 'Nov Revisado (u)', fill_col_q4),
    ('O', 'Dic Prev. Anterior (u)', fill_col_q4),
    ('P', 'Dic Revisado (u)', fill_col_q4),
    ('Q', 'Total Q4 Ant (u)', fill_col_hdr),
    ('R', 'Total Q4 Rev (u)', fill_col_hdr),
    ('S', 'Dif Q4 (u)', fill_col_hdr),
    ('T', 'Comentarios / Justificación Q4', fill_col_hdr)
]

for c_letter, h_text, fill_c in col_headers:
    c_cell = ws[f"{c_letter}9"]
    c_cell.value = h_text
    c_cell.font = font_col_hdr
    c_cell.fill = fill_c
    c_cell.alignment = align_center
    c_cell.border = border_thin
ws.row_dimensions[9].height = 28

# Rellenar filas de datos
start_data_row = 10
row_num = 10

for item in irene_lines:
    r_str = str(row_num)
    
    # B: País
    ws[f"B{r_str}"] = item['pais']
    ws[f"B{r_str}"].font = font_regular
    ws[f"B{r_str}"].alignment = align_center
    ws[f"B{r_str}"].border = border_thin
    
    # C: Cliente
    ws[f"C{r_str}"] = item['cliente']
    ws[f"C{r_str}"].font = font_regular
    ws[f"C{r_str}"].alignment = align_left
    ws[f"C{r_str}"].border = border_thin
    
    # D: SKU
    ws[f"D{r_str}"] = item['sku']
    ws[f"D{r_str}"].font = font_regular
    ws[f"D{r_str}"].alignment = align_center
    ws[f"D{r_str}"].border = border_thin
    
    # E: Descripción
    ws[f"E{r_str}"] = item['desc']
    ws[f"E{r_str}"].font = font_regular
    ws[f"E{r_str}"].alignment = align_left
    ws[f"E{r_str}"].border = border_thin
    
    # F: Oct Previsión Vigente (u)
    ws[f"F{r_str}"] = item['oct_prev_vigente']
    ws[f"F{r_str}"].font = font_regular
    ws[f"F{r_str}"].alignment = align_right
    ws[f"F{r_str}"].number_format = '#,##0'
    ws[f"F{r_str}"].border = border_thin
    
    # G: Oct Real a 27/10 (u)
    ws[f"G{r_str}"] = item['oct_real']
    ws[f"G{r_str}"].font = font_bold
    ws[f"G{r_str}"].alignment = align_right
    ws[f"G{r_str}"].number_format = '#,##0'
    ws[f"G{r_str}"].border = border_thin
    
    # H: Pendiente / Desviación Octubre = |Real - Previsto|
    ws[f"H{r_str}"] = f"=ABS(G{r_str}-F{r_str})"
    ws[f"H{r_str}"].font = font_regular
    ws[f"H{r_str}"].alignment = align_right
    ws[f"H{r_str}"].number_format = '#,##0'
    ws[f"H{r_str}"].border = border_thin
    
    # I: % Consecución a fecha de hoy = Real / Previsto
    ws[f"I{r_str}"] = f'=IF(F{r_str}>0, G{r_str}/F{r_str}, IF(G{r_str}>0, 1.0, 0.0))'
    ws[f"I{r_str}"].font = font_regular
    ws[f"I{r_str}"].alignment = align_right
    ws[f"I{r_str}"].number_format = '0.0%'
    ws[f"I{r_str}"].border = border_thin
    
    # J: Comentario Octubre
    default_coment = ""
    if item['es_extra'] or (item['oct_prev_vigente'] == 0 and item['oct_real'] > 0):
        default_coment = "Pedido extra no previsto"
    elif item['oct_prev_vigente'] > 0 and item['oct_real'] >= item['oct_prev_vigente']:
        default_coment = "Cierre según previsto"
    ws[f"J{r_str}"] = default_coment
    ws[f"J{r_str}"].font = font_regular
    ws[f"J{r_str}"].fill = fill_comment_oct
    ws[f"J{r_str}"].alignment = align_left
    ws[f"J{r_str}"].border = border_thin
    
    # K: Oct Prev. Anterior (u) (Fórmula referenciando F)
    ws[f"K{r_str}"] = f"=F{r_str}"
    ws[f"K{r_str}"].font = font_regular
    ws[f"K{r_str}"].fill = fill_prev_hist
    ws[f"K{r_str}"].alignment = align_right
    ws[f"K{r_str}"].number_format = '#,##0'
    ws[f"K{r_str}"].border = border_thin
    
    # L: Oct Cierre Estimado (u) (Editable verde: de partida igual a Oct Anterior para que empiece en 0 variación hasta que el comercial revise)
    ws[f"L{r_str}"] = item['oct_prev_vigente']
    ws[f"L{r_str}"].font = font_editable
    ws[f"L{r_str}"].fill = fill_editable
    ws[f"L{r_str}"].alignment = align_right
    ws[f"L{r_str}"].number_format = '#,##0'
    ws[f"L{r_str}"].border = border_editable
    
    # M: Nov Prev. Anterior (u)
    ws[f"M{r_str}"] = item['nov_prev_vigente']
    ws[f"M{r_str}"].font = font_regular
    ws[f"M{r_str}"].fill = fill_prev_hist
    ws[f"M{r_str}"].alignment = align_right
    ws[f"M{r_str}"].number_format = '#,##0'
    ws[f"M{r_str}"].border = border_thin
    
    # N: Nov Revisado (u) (Editable verde: precargado con nov_prev_vigente)
    ws[f"N{r_str}"] = item['nov_prev_vigente']
    ws[f"N{r_str}"].font = font_editable
    ws[f"N{r_str}"].fill = fill_editable
    ws[f"N{r_str}"].alignment = align_right
    ws[f"N{r_str}"].number_format = '#,##0'
    ws[f"N{r_str}"].border = border_editable
    
    # O: Dic Prev. Anterior (u)
    ws[f"O{r_str}"] = item['dic_prev_vigente']
    ws[f"O{r_str}"].font = font_regular
    ws[f"O{r_str}"].fill = fill_prev_hist
    ws[f"O{r_str}"].alignment = align_right
    ws[f"O{r_str}"].number_format = '#,##0'
    ws[f"O{r_str}"].border = border_thin
    
    # P: Dic Revisado (u) (Editable verde: precargado con dic_prev_vigente)
    ws[f"P{r_str}"] = item['dic_prev_vigente']
    ws[f"P{r_str}"].font = font_editable
    ws[f"P{r_str}"].fill = fill_editable
    ws[f"P{r_str}"].alignment = align_right
    ws[f"P{r_str}"].number_format = '#,##0'
    ws[f"P{r_str}"].border = border_editable
    
    # Q: Total Q4 Ant = K + M + O
    ws[f"Q{r_str}"] = f"=K{r_str}+M{r_str}+O{r_str}"
    ws[f"Q{r_str}"].font = font_regular
    ws[f"Q{r_str}"].alignment = align_right
    ws[f"Q{r_str}"].number_format = '#,##0'
    ws[f"Q{r_str}"].border = border_thin
    
    # R: Total Q4 Rev = L + N + P
    ws[f"R{r_str}"] = f"=L{r_str}+N{r_str}+P{r_str}"
    ws[f"R{r_str}"].font = font_bold
    ws[f"R{r_str}"].alignment = align_right
    ws[f"R{r_str}"].number_format = '#,##0'
    ws[f"R{r_str}"].border = border_thin
    
    # S: Dif Q4 = R - Q
    ws[f"S{r_str}"] = f"=R{r_str}-Q{r_str}"
    ws[f"S{r_str}"].font = font_regular
    ws[f"S{r_str}"].alignment = align_right
    ws[f"S{r_str}"].number_format = '+#,##0;-#,##0;0'
    ws[f"S{r_str}"].border = border_thin
    
    # T: Comentarios Q4
    ws[f"T{r_str}"] = item['coment_ant'] if item['es_extra'] else ""
    ws[f"T{r_str}"].font = font_regular
    ws[f"T{r_str}"].fill = fill_comment_q4
    ws[f"T{r_str}"].alignment = align_left
    ws[f"T{r_str}"].border = border_thin
    
    ws.row_dimensions[row_num].height = 20
    row_num += 1

last_data_row = row_num - 1

# Fila de Totales
tot_row = row_num
tr_str = str(tot_row)

ws.merge_cells(f"B{tr_str}:E{tr_str}")
ws[f"B{tr_str}"] = "TOTAL CONSOLIDADO - IRENE (UNIDADES)"
ws[f"B{tr_str}"].font = font_total
ws[f"B{tr_str}"].fill = fill_total_row
ws[f"B{tr_str}"].alignment = Alignment(horizontal='right', vertical='center', indent=1)

for c_let in ['B', 'C', 'D', 'E']:
    ws[f"{c_let}{tr_str}"].border = border_double

# Sumas automáticas de columnas numéricas
for c_let in ['F', 'G', 'H', 'K', 'L', 'M', 'N', 'O', 'P', 'Q', 'R', 'S']:
    ws[f"{c_let}{tr_str}"] = f"=SUM({c_let}{start_data_row}:{c_let}{last_data_row})"
    ws[f"{c_let}{tr_str}"].font = font_total
    ws[f"{c_let}{tr_str}"].fill = fill_total_row
    ws[f"{c_let}{tr_str}"].alignment = align_right
    ws[f"{c_let}{tr_str}"].border = border_double
    ws[f"{c_let}{tr_str}"].number_format = '+#,##0;-#,##0;0' if c_let == 'S' else '#,##0'

# Consecución Total
ws[f"I{tr_str}"] = f'=IF(F{tr_str}>0, G{tr_str}/F{tr_str}, 1.0)'
ws[f"I{tr_str}"].font = font_total
ws[f"I{tr_str}"].fill = fill_total_row
ws[f"I{tr_str}"].alignment = align_right
ws[f"I{tr_str}"].number_format = '0.0%'
ws[f"I{tr_str}"].border = border_double

for c_let in ['J', 'T']:
    ws[f"{c_let}{tr_str}"].border = border_double
    ws[f"{c_let}{tr_str}"].fill = fill_total_row

ws.row_dimensions[tot_row].height = 25

# Llenar Fórmulas en KPI Cards de Cabecera (Filas 5 y 6)
# B:D -> F_tot (Previsión Octubre Vigente)
ws.merge_cells("B5:D5")
ws.merge_cells("B6:D6")
ws['B5'] = "PREVISIÓN OCTUBRE VIGENTE (U)"
ws['B5'].font = font_kpi_lbl
ws['B5'].fill = fill_card_blue
ws['B5'].alignment = align_center
ws['B6'] = f"=F{tr_str}"
ws['B6'].font = Font(name='Segoe UI', size=13, bold=True, color='1D4ED8')
ws['B6'].fill = fill_card_blue
ws['B6'].alignment = align_center
ws['B6'].number_format = '#,##0 "u"'

# E:G -> G_tot (Real Octubre a 27/10)
ws.merge_cells("E5:G5")
ws.merge_cells("E6:G6")
ws['E5'] = "REAL OCTUBRE A 27/10 (U)"
ws['E5'].font = font_kpi_lbl
ws['E5'].fill = fill_card_accent
ws['E5'].alignment = align_center
ws['E6'] = f"=G{tr_str}"
ws['E6'].font = Font(name='Segoe UI', size=13, bold=True, color='047857')
ws['E6'].fill = fill_card_accent
ws['E6'].alignment = align_center
ws['E6'].number_format = '#,##0 "u"'

# H:J -> I_tot (% Consecución Octubre a 27/10)
ws.merge_cells("H5:J5")
ws.merge_cells("H6:J6")
ws['H5'] = "% CONSECUCIÓN A FECHA DE HOY"
ws['H5'].font = font_kpi_lbl
ws['H5'].fill = fill_card_kpi
ws['H5'].alignment = align_center
ws['H6'] = f"=I{tr_str}"
ws['H6'].font = Font(name='Segoe UI', size=13, bold=True, color='0F172A')
ws['H6'].fill = fill_card_kpi
ws['H6'].alignment = align_center
ws['H6'].number_format = '0.0%'

# K:M -> Q_tot (Total Q4 Previsión Anterior)
ws.merge_cells("K5:M5")
ws.merge_cells("K6:M6")
ws['K5'] = "PREVISIÓN VIGENTE Q4 (U)"
ws['K5'].font = font_kpi_lbl
ws['K5'].fill = fill_card_blue
ws['K5'].alignment = align_center
ws['K6'] = f"=Q{tr_str}"
ws['K6'].font = Font(name='Segoe UI', size=13, bold=True, color='1D4ED8')
ws['K6'].fill = fill_card_blue
ws['K6'].alignment = align_center
ws['K6'].number_format = '#,##0 "u"'

# N:P -> R_tot (Total Q4 Previsión Revisada)
ws.merge_cells("N5:P5")
ws.merge_cells("N6:P6")
ws['N5'] = "PREVISIÓN REVISADA Q4 (U)"
ws['N5'].font = font_kpi_lbl
ws['N5'].fill = fill_card_accent
ws['N5'].alignment = align_center
ws['N6'] = f"=R{tr_str}"
ws['N6'].font = Font(name='Segoe UI', size=13, bold=True, color='047857')
ws['N6'].fill = fill_card_accent
ws['N6'].alignment = align_center
ws['N6'].number_format = '#,##0 "u"'

# Q:T -> S_tot (Diferencia Neta Q4)
ws.merge_cells("Q5:T5")
ws.merge_cells("Q6:T6")
ws['Q5'] = "VARIACIÓN NETA Q4 (U)"
ws['Q5'].font = font_kpi_lbl
ws['Q5'].fill = fill_card_warn
ws['Q5'].alignment = align_center
ws['Q6'] = f"=S{tr_str}"
ws['Q6'].font = Font(name='Segoe UI', size=13, bold=True, color='B45309')
ws['Q6'].fill = fill_card_warn
ws['Q6'].alignment = align_center
ws['Q6'].number_format = '+#,##0 "u";-#,##0 "u";0 "u"'

for c1, c2 in [('B','D'), ('E','G'), ('H','J'), ('K','M'), ('N','P'), ('Q','T')]:
    start_c = ord(c1)
    end_c = ord(c2)
    for cc in range(start_c, end_c + 1):
        letter = chr(cc)
        ws[f"{letter}5"].border = border_thin
        ws[f"{letter}6"].border = border_thin

# Validación de datos en columna J
ws.add_data_validation(dv_coment_oct)
dv_coment_oct.add(f"J{start_data_row}:J{last_data_row}")

# Anchos de columna óptimos
col_widths = {
    'A': 3,
    'B': 14, # País
    'C': 34, # Cliente
    'D': 14, # SKU
    'E': 32, # Desc
    'F': 18, # Oct Prev Vigente
    'G': 18, # Oct Real a 27/10
    'H': 18, # Pendiente / Desv
    'I': 16, # % Consec a 27/10
    'J': 34, # Coment Octubre
    'K': 18, # Oct Prev Ant
    'L': 18, # Oct Cierre Estimado
    'M': 18, # Nov Prev Ant
    'N': 18, # Nov Revisado
    'O': 18, # Dic Prev Ant
    'P': 18, # Dic Revisado
    'Q': 18, # Total Q4 Ant
    'R': 18, # Total Q4 Rev
    'S': 16, # Dif Q4
    'T': 34  # Coment Q4
}
for col_l, w in col_widths.items():
    ws.column_dimensions[col_l].width = w

# Guardar en carpeta y en raíz
out_dir = "Revision_Previsiones_Q4_2026"
os.makedirs(out_dir, exist_ok=True)
filepath_dir = os.path.join(out_dir, "Revision_Previsiones_Q4_2026_Irene.xlsx")
filepath_root = "Revision_Previsiones_Q4_2026_Irene.xlsx"

wb.save(filepath_dir)
wb.save(filepath_root)

print(f"Archivo guardado exitosamente en:\n  - {filepath_dir}\n  - {filepath_root}")
