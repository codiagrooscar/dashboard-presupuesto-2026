import openpyxl, os, sys
import unicodedata

sys.stdout.reconfigure(encoding='utf-8')

def norm(s):
    if not isinstance(s, str): return ''
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn').upper().strip()

from generar_excel_informe_cambios_q4 import all_modified_lines

print("=================================================================")
print("APLICANDO REVISIONES DE PREVISIONES Q4 A LAS BASES DE DATOS")
print("=================================================================")

# Archivos de Presupuesto a actualizar
budget_files = [
    'Presupuesto_Ventas_2027_Original_Sin_Aplanar.xlsx',
    'Presupuesto_Ventas_2027_Definitivo.xlsx'
]

for bfile in budget_files:
    if not os.path.exists(bfile):
        print(f"Archivo no encontrado: {bfile}")
        continue
        
    print(f"\nProcesando {bfile}...")
    wb = openpyxl.load_workbook(bfile, data_only=False)
    sheet_name = 'Previsión Matriz Horizontal' if 'Previsión Matriz Horizontal' in wb.sheetnames else wb.sheetnames[2]
    ws = wb[sheet_name]
    
    # Mapear filas por cliente y sku
    rows_map = {}
    for r in range(2, ws.max_row + 1):
        com = norm(ws.cell(r, 1).value or '')
        pais = str(ws.cell(r, 2).value or '').strip().upper()
        cli = norm(ws.cell(r, 3).value or '')
        sku = str(ws.cell(r, 5).value or '').strip().upper()
        sku_clean = sku
        if pais != 'ESPAÑA' and len(sku) > 2:
            sku_clean = sku[:-2]
            
        key = (cli, sku_clean)
        rows_map.setdefault(key, []).append(r)
        
    updated_count = 0
    
    for ch in all_modified_lines:
        ch_cli = norm(ch['cliente'])
        ch_sku = str(ch['sku']).strip().upper()
        ch_com = norm(ch['comercial'])
        
        # Buscar en map
        candidates = rows_map.get((ch_cli, ch_sku), [])
        
        # Si no encuentra directo, buscar coincidencias parciales de cliente
        if not candidates:
            for (map_cli, map_sku), r_list in rows_map.items():
                if map_sku == ch_sku and (ch_cli == map_cli or ch_cli in map_cli or map_cli in ch_cli):
                    candidates = r_list
                    break
                    
        if not candidates:
            # Caso especial Marcoser sustitución SKU
            if 'MARCOSER' in ch_cli:
                for (map_cli, map_sku), r_list in rows_map.items():
                    if 'MARCOSER' in map_cli and 'AGXKAE' in map_sku:
                        candidates = r_list
                        break
                        
        if not candidates:
            print(f"  [AVISO] No se encontró fila para {ch['comercial']} | {ch['cliente']} | {ch['sku']}")
            continue
            
        # Si hay varios candidatos, elegir el activo (el que tenía valores no nulos en Oct/Nov/Dic o mayor Total)
        target_row = candidates[0]
        if len(candidates) > 1:
            best_r = candidates[0]
            max_v = -1
            for r in candidates:
                # Comprobar si coincide comercial
                r_com = norm(ws.cell(r, 1).value or '')
                if ch_com in r_com or r_com in ch_com:
                    # Suma de Q4 actual
                    v_q4 = (float(ws.cell(r, 28).value or 0) + 
                            float(ws.cell(r, 30).value or 0) + 
                            float(ws.cell(r, 32).value or 0))
                    if v_q4 > max_v:
                        max_v = v_q4
                        best_r = r
            target_row = best_r
            
        # Precio medio para valorar €
        pm_val = ws.cell(target_row, 8).value
        p_medio = float(pm_val) if isinstance(pm_val, (int, float)) and pm_val > 0 else ch['precio']
        
        # Actualizar Oct-26
        ws.cell(target_row, 28).value = ch['oct_rev']
        ws.cell(target_row, 29).value = round(ch['oct_rev'] * p_medio, 2)
        
        # Actualizar Nov-26
        ws.cell(target_row, 30).value = ch['nov_rev']
        ws.cell(target_row, 31).value = round(ch['nov_rev'] * p_medio, 2)
        
        # Actualizar Dic-26
        ws.cell(target_row, 32).value = ch['dic_rev']
        ws.cell(target_row, 33).value = round(ch['dic_rev'] * p_medio, 2)
        
        # Recalcular Total 2026 (u) y Total 2026 (€)
        # Meses u están en columnas: 10, 12, 14, 16, 18, 20, 22, 24, 26, 28, 30, 32
        tot_u = sum(float(ws.cell(target_row, c).value or 0) for c in [10, 12, 14, 16, 18, 20, 22, 24, 26, 28, 30, 32])
        tot_eur = sum(float(ws.cell(target_row, c).value or 0) for c in [11, 13, 15, 17, 19, 21, 23, 25, 27, 29, 31, 33])
        
        ws.cell(target_row, 34).value = tot_u
        ws.cell(target_row, 35).value = round(tot_eur, 2)
        
        updated_count += 1
        
    wb.save(bfile)
    print(f"  -> Guardado {bfile} con éxito ({updated_count} líneas actualizadas en matriz presupuestaria)")

print("\n2. Actualizando Revision_Previsiones_Q4_2026_CONSOLIDADO.xlsx...")
con_file = 'Revision_Previsiones_Q4_2026/Revision_Previsiones_Q4_2026_CONSOLIDADO.xlsx'
if os.path.exists(con_file):
    wb_con = openpyxl.load_workbook(con_file, data_only=False)
    
    for com in ['Alfonso', 'Garcia', 'Javier', 'Pedro', 'Ricardo']:
        copy_com_path = f'Revision_Previsiones_Q4_2026/Copia de Revision_Previsiones_Q4_2026_{com}.xlsx'
        if not os.path.exists(copy_com_path): continue
        
        wb_c_single = openpyxl.load_workbook(copy_com_path, data_only=False)
        ws_src = wb_c_single.active
        
        # Buscar hoja en consolidado
        target_sname = None
        for sn in wb_con.sheetnames:
            if norm(sn) == norm(com):
                target_sname = sn
                break
                
        if target_sname:
            ws_dst = wb_con[target_sname]
            # Copiar datos fila a fila (columnas J a T)
            for r in range(10, ws_src.max_row + 1):
                # Col 10: Comentario Sep
                ws_dst.cell(r, 10).value = ws_src.cell(r, 10).value
                # Col 12: Oct Rev
                ws_dst.cell(r, 12).value = ws_src.cell(r, 12).value
                # Col 14: Nov Rev
                ws_dst.cell(r, 14).value = ws_src.cell(r, 14).value
                # Col 16: Dic Rev
                ws_dst.cell(r, 16).value = ws_src.cell(r, 16).value
                # Col 20: Comentario Q4
                ws_dst.cell(r, 20).value = ws_src.cell(r, 20).value
            print(f"  -> Hoja {target_sname} sincronizada en consolidado")
            
    wb_con.save(con_file)
    print("  -> Consolidado guardado con éxito")

print("\n¡Bases de datos de previsiones actualizadas con éxito!")
