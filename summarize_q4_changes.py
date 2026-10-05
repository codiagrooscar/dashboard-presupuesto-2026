import openpyxl, os, sys
sys.stdout.reconfigure(encoding='utf-8')

comerciales = ['Alfonso', 'Garcia', 'Javier', 'Pedro', 'Ricardo']
folder = 'Revision_Previsiones_Q4_2026'

print(f"{'Comercial':10s} | {'Modif Q4':8s} | {'Coment':8s} | {'Oct Ant':10s} | {'Oct Rev':10s} | {'Nov Ant':10s} | {'Nov Rev':10s} | {'Dic Ant':10s} | {'Dic Rev':10s} | {'Total Q4 Ant':12s} | {'Total Q4 Rev':12s} | {'Dif Q4':12s}")
print('-'*135)

total_oct_ant = 0
total_oct_rev = 0
total_nov_ant = 0
total_nov_rev = 0
total_dic_ant = 0
total_dic_rev = 0

all_detailed_changes = []

for com in comerciales:
    f_orig = os.path.join(folder, f'Revision_Previsiones_Q4_2026_{com}.xlsx')
    f_copy = os.path.join(folder, f'Copia de Revision_Previsiones_Q4_2026_{com}.xlsx')
    
    wb_orig = openpyxl.load_workbook(f_orig, data_only=True)
    wb_copy = openpyxl.load_workbook(f_copy, data_only=True)
    
    ws_orig = wb_orig.active
    ws_copy = wb_copy.active
    
    cnt_q4_changes = 0
    cnt_comments = 0
    
    com_oct_ant = 0
    com_oct_rev = 0
    com_nov_ant = 0
    com_nov_rev = 0
    com_dic_ant = 0
    com_dic_rev = 0
    
    for r in range(10, ws_copy.max_row + 1):
        c_pais = ws_copy.cell(r, 2).value
        c_name = ws_copy.cell(r, 3).value
        sku = ws_copy.cell(r, 4).value
        desc = ws_copy.cell(r, 5).value
        
        # If it's a total row
        if not sku or 'TOTAL' in str(c_name).upper() or 'TOTAL' in str(sku).upper():
            continue
            
        oct_ant = float(ws_copy.cell(r, 11).value or 0)
        oct_rev = float(ws_copy.cell(r, 12).value or 0)
        nov_ant = float(ws_copy.cell(r, 13).value or 0)
        nov_rev = float(ws_copy.cell(r, 14).value or 0)
        dic_ant = float(ws_copy.cell(r, 15).value or 0)
        dic_rev = float(ws_copy.cell(r, 16).value or 0)
        
        com_sep = ws_copy.cell(r, 10).value
        com_q4 = ws_copy.cell(r, 20).value
        
        com_sep_o = ws_orig.cell(r, 10).value
        com_q4_o = ws_orig.cell(r, 20).value
        
        has_q4_change = (oct_ant != oct_rev) or (nov_ant != nov_rev) or (dic_ant != dic_rev)
        has_com_change = (com_sep and com_sep != com_sep_o) or (com_q4 and com_q4 != com_q4_o)
        
        if has_q4_change:
            cnt_q4_changes += 1
            
        if has_com_change:
            cnt_comments += 1
            
        if has_q4_change or has_com_change:
            all_detailed_changes.append({
                'comercial': com,
                'fila': r,
                'pais': c_pais,
                'cliente': c_name,
                'sku': sku,
                'desc': desc,
                'sep_prev': ws_copy.cell(r, 6).value or 0,
                'sep_real': ws_copy.cell(r, 7).value or 0,
                'com_sep': com_sep,
                'oct_ant': oct_ant, 'oct_rev': oct_rev, 'oct_dif': oct_rev - oct_ant,
                'nov_ant': nov_ant, 'nov_rev': nov_rev, 'nov_dif': nov_rev - nov_ant,
                'dic_ant': dic_ant, 'dic_rev': dic_rev, 'dic_dif': dic_rev - dic_ant,
                'q4_ant': oct_ant + nov_ant + dic_ant,
                'q4_rev': oct_rev + nov_rev + dic_rev,
                'q4_dif': (oct_rev + nov_rev + dic_rev) - (oct_ant + nov_ant + dic_ant),
                'com_q4': com_q4
            })
            
        com_oct_ant += oct_ant
        com_oct_rev += oct_rev
        com_nov_ant += nov_ant
        com_nov_rev += nov_rev
        com_dic_ant += dic_ant
        com_dic_rev += dic_rev
        
    tot_ant = com_oct_ant + com_nov_ant + com_dic_ant
    tot_rev = com_oct_rev + com_nov_rev + com_dic_rev
    diff = tot_rev - tot_ant
    
    total_oct_ant += com_oct_ant
    total_oct_rev += com_oct_rev
    total_nov_ant += com_nov_ant
    total_nov_rev += com_nov_rev
    total_dic_ant += com_dic_ant
    total_dic_rev += com_dic_rev
    
    print(f"{com:10s} | {cnt_q4_changes:8d} | {cnt_comments:8d} | {com_oct_ant:10,.0f} | {com_oct_rev:10,.0f} | {com_nov_ant:10,.0f} | {com_nov_rev:10,.0f} | {com_dic_ant:10,.0f} | {com_dic_rev:10,.0f} | {tot_ant:12,.0f} | {tot_rev:12,.0f} | {diff:+12,.0f}")

print('-'*135)
tot_all_ant = total_oct_ant + total_nov_ant + total_dic_ant
tot_all_rev = total_oct_rev + total_nov_rev + total_dic_rev
tot_all_diff = tot_all_rev - tot_all_ant
print(f"{'TOTAL':10s} | {len([c for c in all_detailed_changes if c['q4_dif'] != 0]):8d} | {len([c for c in all_detailed_changes if c['com_q4'] or c['com_sep']]):8d} | {total_oct_ant:10,.0f} | {total_oct_rev:10,.0f} | {total_nov_ant:10,.0f} | {total_nov_rev:10,.0f} | {total_dic_ant:10,.0f} | {total_dic_rev:10,.0f} | {tot_all_ant:12,.0f} | {tot_all_rev:12,.0f} | {tot_all_diff:+12,.0f}")

print(f"\nTotal líneas registradas en detalle: {len(all_detailed_changes)}")
