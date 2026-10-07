# -*- coding: utf-8 -*-
import openpyxl, datetime, difflib, unicodedata
from collections import defaultdict

def normalize(text):
    if not text:
        return ""
    text = unicodedata.normalize('NFKD', text)
    text = "".join(c for c in text if not unicodedata.combining(c))
    for char in [",", ".", "-", "'", "´", "`", "\"", "/"]:
        text = text.replace(char, " ")
    return " ".join(text.upper().split())

wb = openpyxl.load_workbook('temp_cf_diplomats.xlsx', data_only=True)
ws = wb['Clients']

client_terms = {}
for r in range(3, ws.max_row + 1):
    cli = str(ws.cell(r, 2).value).strip() if ws.cell(r, 2).value else None
    if not cli or cli == 'None':
        continue
    canal = ws.cell(r, 3).value
    pais = ws.cell(r, 4).value
    dias = ws.cell(r, 10).value
    iva = ws.cell(r, 8).value
    fact = ws.cell(r, 7).value
    rate_iva = 0.10 if (iva and fact and float(fact) > 0 and float(iva)/float(fact) > 0.05) else 0.0
    if cli not in client_terms:
        client_terms[cli] = {
            'canal': canal or 'Cobros por canal Clientes nacionales',
            'pais': pais or 'ES',
            'dias': int(dias) if dias is not None else 60,
            'iva_rate': rate_iva
        }

sept_sales = [
    ('430000001', 'COMERCIAL TÉCNICA FITOCARTHAGO, S.L.', 5240, 14014),
    ('430000003', 'AGROQUIMICAS DRAGO S.L.', 8400, 21544),
    ('430000004', 'COOP.AGRIC. SAN ISIDRO S.C.A.', 0, 0),
    ('430000005', 'ASESORAMIENTO TECNICO DE CULTIVOS, S.L.', 1920, 6125),
    ('430000006', 'FITOMURCIA S.L.', 1400, 6607),
    ('430000008', 'DOTRA CHEMICALS', 21000, 81270),
    ('430000009', 'DIPLOMATS ORCHARDS FOR TRADE AGENCIES', -6, -37),
    ('430000011', 'ECUAQUIMICA', 8800, 14916),
    ('430000015', 'CABRERA MORALES FRANCISCO', 6855, 14348),
    ('430000018', 'TRANSPORTES CAMPILLO, S.A.', 80, 216),
    ('430000019', 'CAMPOEJIDO S.C.A.', 12464, 52104),
    ('430000021', 'FITOXENT S.C.L.V.', 0, 0),
    ('430000023', 'MAJ AGROQUIMICOS S.L.', 60, 108),
    ('430000024', 'FITOSANITARIOS DRAGO S.L.U', 14784, 40585),
    ('430000028', 'SC MARCOSER S.R.L.', 7932, 43438),
    ('430000033', 'CODIMED EXPORT, S.L.', 1460, 7774),
    ('430000037', 'EUROFERTIL, S.A.', 20050, 58212),
    ('430000058', 'AGROTÉCNICA DEL NORTE, S.L', 1200, 4852),
    ('430000064', 'AIFAR S.P.A', 3000, 12000),
    ('430000065', 'AGRICOLA DEL JUCAR DE CUENCA, S.L.', 300, 1255),
    ('430000075', 'AGROVEYCA,S.L', 720, 2702),
    ('430000080', 'SAFIAGRO, S.A.', 640, 1837),
    ('430000089', 'T.J AGRICULTURA S.L.', 1100, 3131),
    ('430000105', 'TRATAMIENTOS AGRICOLAS BRENES, S.L.', 1100, 2035),
    ('430000143', 'INIESTA DEL PINO DISTRIBUCIONES S.L.U', 0, 0),
    ('430000190', 'SUCA, S.C.A.', 640, 3323),
    ('430000192', 'SERVICIOS AGROPECUARIOS EL LLANO, S. A.', 12872, 48446),
    ('430000205', 'AGROCAMP CANET S.L.', 0, 0),
    ('430000237', 'AGRUPACION DE LABRADORES DEL POZUELO, SA', 600, 3112),
    ('430000243', 'DOGA TARIM TIC.SAN.VE TUR. LTD. STI', 0, 0),
    ('430000269', 'SERVICIOS AGRÍCOLAS CAMPO RIO, S.L.', 0, 0),
    ('430000292', 'AGROSER BERNARDINO, S.L.', 1100, 1903),
    ('430000295', 'GUIVARTO AGRICOLA, S.L.', 10370, 32009),
    ('430000326', 'FINCA DOÑA ANA C.B.', 1120, 4213),
    ('430000332', 'SERVIAGRI´ 97, S.L.', 3716, 13311),
    ('430000335', 'AGROGIMEDEL, S.L.U.', 0, 0),
    ('430000344', 'CAMPO ABIERTO, S.A.S.', 0, 0),
    ('430000350', 'BARAHONA HERNAEZ, DAVID', 0, 0),
    ('430000372', 'SUMINISTROS AGRÍCOLAS SALAMPACK, S.L.', 0, 0),
    ('430000379', 'ALCADA EXPORT, S.L.', 0, 0),
    ('430000380', 'LEKKERBIO B.V.', 0, 0),
    ('430000381', 'BARDENAS COMERCIAL, S.L.', 600, 2070),
    ('430000382', 'STE BBMAGRI, SARL', -2600, -1212),
    ('430000384', 'AGRONEW FITO, S.L.', 3480, 8176),
    ('430000385', 'MUNDOHUERTO-CARMEN SOLAZ BLANCO', 164, 575),
    ('430000386', 'VILLALTA SATIVUM, S.L.', 3460, 6193),
    ('430000387', 'EL PINAR BERRIES, S.L.', 1100, 2409),
    ('430000389', 'ASOC. HORTICULTORES DE VILLA DEL PRADO', 210, 1013),
    ('430000394', 'COMERCIAL NORTE AGROCAMPO, S.L.', 220, 682),
    ('430000396', 'ZTB AQRO MMC', 15, 0),
    ('430000405', 'AGRITEC MEDINA, S.L.U.', 0, 0),
    ('430000414', 'BOSCO ABASCAL, S.L.', 0, 0),
    ('430000417', 'KUWAIT FARM COMPANY', 0, 0),
    ('430000418', 'AUDITORIA Y CERTIFICACION EN GESTION DE', 0, 0),
    ('430000419', 'AGROREBOLLO, S.L.', 0, 0),
    ('430000420', 'SALES AMPOSTA, JAUME', 0, 0),
    ('430000421', 'ACTIV. PRODUC. Y COMERC. LA VEGUILLA, SL', 2560, 6771),
    ('430000425', 'DOTRA GULF TRADING COMPANY', 20284, 88497),
    ('430000429', 'FITOSANITARIOS CARCAIXENT, S.L.', 8700, 26197),
    ('430000430', 'ALMENDRALIA IBÉRICA, S.L.U.', 1100, 1958)
]

# Process September real entries
base_date = datetime.date(2026, 9, 15)
sept_vencimientos = []

for code, name, lit, eur in sept_sales:
    if eur == 0 and lit == 0:
        continue
    # find client terms
    n_name = normalize(name)
    matched_term = None
    for k, v in client_terms.items():
        if normalize(k) == n_name:
            matched_term = (k, v)
            break
    if not matched_term:
        for k, v in client_terms.items():
            if n_name in normalize(k) or normalize(k) in n_name:
                matched_term = (k, v)
                break
    if not matched_term:
        # specific manual fallbacks
        if 'DOTRA GULF' in n_name:
            matched_term = ('DOTRA GULF TRADING COMPANY', {'canal': 'Cobros por canal Clientes internac.', 'pais': 'AE', 'dias': 25, 'iva_rate': 0.0})
        elif 'CAMPILLO' in n_name:
            matched_term = ('TRANSPORTES CAMPILLO, S.A.', {'canal': 'Cobros por canal Clientes nacionales', 'pais': 'ES', 'dias': 60, 'iva_rate': 0.10})
        elif 'ALMENDRALIA' in n_name:
            matched_term = ('ALMENDRALIA IBÉRICA, S.L.U.', {'canal': 'Cobros por canal Clientes nacionales', 'pais': 'ES', 'dias': 60, 'iva_rate': 0.10})
        elif 'FINCA DONA ANA' in n_name:
            matched_term = ('FINCA DOÑA ANA C.B.', {'canal': 'Cobros por canal Clientes nacionales', 'pais': 'ES', 'dias': 60, 'iva_rate': 0.10})
        elif 'SUCA' in n_name:
            matched_term = ('SUCA, S.C.A.', {'canal': 'Cobros por canal Clientes nacionales', 'pais': 'ES', 'dias': 60, 'iva_rate': 0.10})
        else:
            matched_term = (name, {'canal': 'Cobros por canal Clientes nacionales', 'pais': 'ES', 'dias': 60, 'iva_rate': 0.10})

    cli_canonical, terms = matched_term
    dias = terms['dias']
    canal = terms['canal']
    pais = terms['pais']
    iva_rate = terms['iva_rate']

    # For diplomats: -37 EUR
    fact_val = float(eur)
    iva_val = fact_val * iva_rate
    imp_val = fact_val + iva_val

    vto_date = base_date + datetime.timedelta(days=dias)
    mes_vto = f"{vto_date.month}-{vto_date.year}"

    sept_vencimientos.append({
        'codigo': code,
        'cliente': cli_canonical,
        'canal': canal,
        'pais': pais,
        'mes_venta': 'sep',
        'fecha_operacion': base_date.strftime('%d/%m/%Y'),
        'facturacion': fact_val,
        'iva': iva_val,
        'importe': imp_val,
        'dias': dias,
        'vto_fecha': vto_date.strftime('%d/%m/%Y'),
        'mes_vto': mes_vto,
        'tipo_origen': 'Venta Real Sep'
    })

print(f"Total registros septiembre con actividad: {len(sept_vencimientos)}")
q4_by_month = defaultdict(lambda: defaultdict(float))
for sv in sept_vencimientos:
    mv = sv['mes_vto']
    q4_by_month[mv][sv['canal']] += sv['importe']

print("\n--- COBROS GENERADOS POR LAS VENTAS REALES DE SEPTIEMBRE ---")
for m, canals in sorted(q4_by_month.items()):
    tot_m = sum(canals.values())
    print(f"Mes Vencimiento: {m} -> Total: {tot_m:,.2f} €")
    for c, val in sorted(canals.items()):
        print(f"   {c}: {val:,.2f} €")
