import openpyxl
import pandas as pd
import numpy as np
import unicodedata
import shutil
import os

def norm(s):
    if not isinstance(s, str): return ''
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn').upper().strip()

print("1. Cargando 'Garcia 2026 V3 CORREGIDA.xlsx'...")
shutil.copy2('Garcia 2026 V3 CORREGIDA.xlsx', 'temp_garcia_v3.xlsx')
df_v3 = pd.read_excel('temp_garcia_v3.xlsx', sheet_name='GARCIA', skiprows=3)
df_v3 = df_v3.dropna(subset=['CLIENTE']).copy()
df_v3['CLIENTE_NORM'] = df_v3['CLIENTE'].apply(norm)
df_v3['SKU_CLEAN'] = df_v3['SKU (NUEVOARTICULO)'].astype(str).str.strip().str.upper()

for m in ['SEP', 'OCT', 'NOV', 'DIC']:
    df_v3[f'PREV_{m}'] = pd.to_numeric(df_v3[f'PREV {m}-26'], errors='coerce').fillna(0)

# Agrupar V3 por CLIENTE_NORM y SKU_CLEAN (resuelve duplicados como el de EMPHYTON)
v3_agg = df_v3.groupby(['CLIENTE_NORM', 'SKU_CLEAN']).agg(
    cliente=('CLIENTE', 'first'),
    pais=('PAÍS', 'first'),
    sku=('SKU (NUEVOARTICULO)', 'first'),
    descripcion=('DESCRIPCIÓN', 'first'),
    sep=('PREV_SEP', 'sum'),
    oct=('PREV_OCT', 'sum'),
    nov=('PREV_NOV', 'sum'),
    dic=('PREV_DIC', 'sum')
).reset_index()

v3_active = v3_agg[(v3_agg['sep'] > 0) | (v3_agg['oct'] > 0) | (v3_agg['nov'] > 0) | (v3_agg['dic'] > 0)].copy()
print(f"Líneas activas en V3 para García (Sep-Dic 2026): {len(v3_active)}")

# 2. Backup de Presupuesto_Ventas_2027.xlsx
backup_file = 'Presupuesto_Ventas_2027_PRE_GARCIA_V3_BACKUP.xlsx'
if not os.path.exists(backup_file):
    shutil.copy2('Presupuesto_Ventas_2027.xlsx', backup_file)
    print(f"Backup creado: {backup_file}")

# 3. Cargar Presupuesto_Ventas_2027.xlsx con openpyxl para preservar fórmulas y estilos
print("2. Abriendo 'Presupuesto_Ventas_2027.xlsx' con openpyxl...")
wb = openpyxl.load_workbook('Presupuesto_Ventas_2027.xlsx')
ws = wb['Previsión Matriz Horizontal']

# Mapear columnas de la cabecera
headers = {}
for col_idx in range(1, ws.max_column + 1):
    val = ws.cell(1, col_idx).value
    if val:
        headers[str(val).strip()] = col_idx

col_com = headers.get('Comercial', 1)
col_cli = headers.get('Cliente', 3)
col_sku = headers.get('Código Artículo', 5)
col_pr26 = headers.get('Precio Medio 2026 (€/u)', 8)
col_pr27 = headers.get('Precio Medio 2027 (€/u)', 9)

col_sep_u = headers.get('Sep-26 (u)', 26)
col_sep_e = headers.get('Sep-26 (€)', 27)
col_oct_u = headers.get('Oct-26 (u)', 28)
col_oct_e = headers.get('Oct-26 (€)', 29)
col_nov_u = headers.get('Nov-26 (u)', 30)
col_nov_e = headers.get('Nov-26 (€)', 31)
col_dic_u = headers.get('Dic-26 (u)', 32)
col_dic_e = headers.get('Dic-26 (€)', 33)

print("Columnas identificadas:")
print(f"  Sep: u={col_sep_u}, e={col_sep_e} | Oct: u={col_oct_u}, e={col_oct_e}")
print(f"  Nov: u={col_nov_u}, e={col_nov_e} | Dic: u={col_dic_u}, e={col_dic_e}")

# Indexar filas de García en la matriz
garcia_rows = []
for r in range(2, ws.max_row + 1):
    com_val = str(ws.cell(r, col_com).value or '')
    if 'GARC' in norm(com_val):
        cli_val = str(ws.cell(r, col_cli).value or '')
        sku_val = str(ws.cell(r, col_sku).value or '').strip().upper()
        # Verificar si esta fila tenía datos en Sep-Dic previamente
        prev_sep_u = ws.cell(r, col_sep_u).value or 0
        prev_oct_u = ws.cell(r, col_oct_u).value or 0
        prev_nov_u = ws.cell(r, col_nov_u).value or 0
        prev_dic_u = ws.cell(r, col_dic_u).value or 0
        has_late_data = any([isinstance(v, (int, float)) and v > 0 for v in [prev_sep_u, prev_oct_u, prev_nov_u, prev_dic_u]])
        garcia_rows.append({
            'row': r,
            'cli_norm': norm(cli_val),
            'sku_clean': sku_val,
            'has_late_data': has_late_data
        })

print(f"Total filas de García en la matriz: {len(garcia_rows)}")

# Limpiar las 4 columnas Sep-Dic de García en todas sus filas actuales
for g in garcia_rows:
    r = g['row']
    ws.cell(r, col_sep_u, 0.0)
    ws.cell(r, col_sep_e, 0.0)
    ws.cell(r, col_oct_u, 0.0)
    ws.cell(r, col_oct_e, 0.0)
    ws.cell(r, col_nov_u, 0.0)
    ws.cell(r, col_nov_e, 0.0)
    ws.cell(r, col_dic_u, 0.0)
    ws.cell(r, col_dic_e, 0.0)

# Aplicar los nuevos valores de V3
updated_combinations = set()
for _, v_row in v3_active.iterrows():
    c_norm = v_row['CLIENTE_NORM']
    s_clean = v_row['SKU_CLEAN']
    
    # Buscar coincidencia en garcia_rows
    # Priorizar la fila que tenía late_data si hay duplicados
    matching = [g for g in garcia_rows if g['cli_norm'] == c_norm and g['sku_clean'] == s_clean]
    
    target_row = None
    if len(matching) == 1:
        target_row = matching[0]['row']
    elif len(matching) > 1:
        # Priorizar fila con late data
        late_matches = [m for m in matching if m['has_late_data']]
        if late_matches:
            target_row = late_matches[0]['row']
        else:
            target_row = matching[-1]['row']
            
    if target_row:
        pr = ws.cell(target_row, col_pr26).value
        try:
            pr = float(pr) if pr is not None else 0.0
        except (ValueError, TypeError):
            pr = 0.0
            
        u_sep = float(v_row['sep'])
        u_oct = float(v_row['oct'])
        u_nov = float(v_row['nov'])
        u_dic = float(v_row['dic'])
        
        ws.cell(target_row, col_sep_u, u_sep)
        ws.cell(target_row, col_sep_e, round(u_sep * pr, 2))
        ws.cell(target_row, col_oct_u, u_oct)
        ws.cell(target_row, col_oct_e, round(u_oct * pr, 2))
        ws.cell(target_row, col_nov_u, u_nov)
        ws.cell(target_row, col_nov_e, round(u_nov * pr, 2))
        ws.cell(target_row, col_dic_u, u_dic)
        ws.cell(target_row, col_dic_e, round(u_dic * pr, 2))
        
        updated_combinations.add((c_norm, s_clean))
    else:
        # Línea nueva que no estaba en la matriz (por ejemplo BBMAGRI AGN0005MA)
        new_r = ws.max_row + 1
        print(f"Añadiendo nueva fila en línea {new_r}: {v_row['cliente']} | {v_row['sku']}")
        ws.cell(new_r, col_com, 'García')
        ws.cell(new_r, headers.get('País', 2), v_row['pais'])
        ws.cell(new_r, col_cli, v_row['cliente'])
        ws.cell(new_r, headers.get('Clasificación Cliente', 4), 'Consolidado/estratégico')
        ws.cell(new_r, col_sku, v_row['sku'])
        ws.cell(new_r, headers.get('Descripción Artículo', 6), v_row['descripcion'])
        ws.cell(new_r, headers.get('Tamaño Envase', 7), 5.0)
        
        # Precio aproximado
        pr = 3.22
        ws.cell(new_r, col_pr26, pr)
        ws.cell(new_r, col_pr27, pr * 1.04)
        
        u_sep = float(v_row['sep'])
        u_oct = float(v_row['oct'])
        u_nov = float(v_row['nov'])
        u_dic = float(v_row['dic'])
        
        # Poner 0 en meses Ene-Ago
        for col_i in range(10, 26):
            ws.cell(new_r, col_i, 0.0)
            
        ws.cell(new_r, col_sep_u, u_sep)
        ws.cell(new_r, col_sep_e, round(u_sep * pr, 2))
        ws.cell(new_r, col_oct_u, u_oct)
        ws.cell(new_r, col_oct_e, round(u_oct * pr, 2))
        ws.cell(new_r, col_nov_u, u_nov)
        ws.cell(new_r, col_nov_e, round(u_nov * pr, 2))
        ws.cell(new_r, col_dic_u, u_dic)
        ws.cell(new_r, col_dic_e, round(u_dic * pr, 2))
        
        # Total formulas
        ws.cell(new_r, headers.get('Total 2026 (u)', 34), f"=J{new_r}+L{new_r}+N{new_r}+P{new_r}+R{new_r}+T{new_r}+V{new_r}+X{new_r}+Z{new_r}+AB{new_r}+AD{new_r}+AF{new_r}")
        ws.cell(new_r, headers.get('Total 2026 (€)', 35), f"=K{new_r}+M{new_r}+O{new_r}+Q{new_r}+S{new_r}+U{new_r}+W{new_r}+Y{new_r}+AA{new_r}+AC{new_r}+AE{new_r}+AG{new_r}")
        
        updated_combinations.add((c_norm, s_clean))

print(f"Total combinaciones V3 actualizadas con éxito: {len(updated_combinations)} de {len(v3_active)}")

# Guardar Presupuesto_Ventas_2027.xlsx
output_file = 'Presupuesto_Ventas_2027.xlsx'
wb.save(output_file)
print(f"¡{output_file} actualizado y guardado correctamente!")
