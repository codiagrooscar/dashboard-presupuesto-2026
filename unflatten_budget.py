import openpyxl
import pandas as pd
import numpy as np
import unicodedata
import shutil
import os

def norm(s):
    if not isinstance(s, str): return ''
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn').upper().strip()

FACTOR = 0.9479961392983893
INV_FACTOR = 1.0 / FACTOR

print("1. Abriendo Presupuesto_Ventas_2027...")
shutil.copy2('Presupuesto_Ventas_2027.xlsx', 'temp_p2027_unflatten.xlsx')
wb = openpyxl.load_workbook('temp_p2027_unflatten.xlsx')
ws = wb['Previsión Matriz Horizontal']

headers = {}
for c in range(1, ws.max_column + 1):
    val = ws.cell(1, c).value
    if val:
        headers[str(val).strip()] = c

col_com = headers.get('Comercial', 1)
col_pr26 = headers.get('Precio Medio 2026 (€/u)', 8)
months = [
    (headers.get('Sep-26 (u)', 26), headers.get('Sep-26 (€)', 27)),
    (headers.get('Oct-26 (u)', 28), headers.get('Oct-26 (€)', 29)),
    (headers.get('Nov-26 (u)', 30), headers.get('Nov-26 (€)', 31)),
    (headers.get('Dic-26 (u)', 32), headers.get('Dic-26 (€)', 33)),
]

count_cells = 0
for r in range(2, ws.max_row + 1):
    com_val = str(ws.cell(r, col_com).value or '')
    # García ya fue actualizado con números limpios V3
    if 'garc' in norm(com_val).lower():
        continue
    
    pr = ws.cell(r, col_pr26).value
    try:
        pr = float(pr) if pr is not None else 0.0
    except (ValueError, TypeError):
        pr = 0.0

    for col_u, col_e in months:
        val_u = ws.cell(r, col_u).value
        if val_u is not None and isinstance(val_u, (int, float)) and val_u > 0:
            restored_u = round(val_u * INV_FACTOR, 2)
            # Si está muy cerca de un entero (ej 300.0001), redondear a entero
            if abs(restored_u - round(restored_u)) < 0.05:
                restored_u = float(round(restored_u))
            ws.cell(r, col_u, restored_u)
            if col_e and pr > 0:
                ws.cell(r, col_e, round(restored_u * pr, 2))
            count_cells += 1

print(f"Total celdas desaplanadas (restauradas a unidades originales): {count_cells}")

# Guardar
try:
    wb.save('Presupuesto_Ventas_2027.xlsx')
    print("¡Presupuesto_Ventas_2027.xlsx guardado directamente!")
except PermissionError:
    wb.save('Presupuesto_Ventas_2027_Original_Sin_Aplanar.xlsx')
    print("Nota: Presupuesto_Ventas_2027.xlsx estaba abierto en Excel. Guardado como Presupuesto_Ventas_2027_Original_Sin_Aplanar.xlsx")

wb.save('temp_budget_input.xlsx')
print("temp_budget_input.xlsx actualizado con unidades originales.")
