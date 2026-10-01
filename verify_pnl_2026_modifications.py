import openpyxl
import pandas as pd

def clean_v(v):
    if v is None: return 0.0
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).strip().replace(' ', '').replace('-', '0'))
    except: return 0.0

wb = openpyxl.load_workbook(r'Gastos/P&L mensual ene-ago V4.xlsx', data_only=False)
ws = wb['PL Mensual 2026']

# Load data_only version to evaluate numbers
wb_data = openpyxl.load_workbook(r'Gastos/P&L mensual ene-ago V4.xlsx', data_only=True)
ws_data = wb_data['PL Mensual 2026']

rows = []
for r in range(4, 206):
    n3 = ws.cell(r, 6).value
    n4 = ws.cell(r, 7).value
    n5 = ws.cell(r, 8).value
    cta = ws.cell(r, 9).value
    code = str(cta).split(' - ')[0].split(' ')[0]
    
    # Months Jan-Dec
    months = [clean_v(ws_data.cell(r, c).value) for c in range(11, 23)]
    tot_year = sum(months)
    tot_8m = sum(months[:8])
    tot_sd = sum(months[8:])
    crit = ws.cell(r, 24).value
    
    rows.append({
        'row': r,
        'code': code,
        'n3': n3,
        'n4': n4,
        'n5': n5,
        'cta': cta,
        'months': months,
        'tot_8m': tot_8m,
        'sep': months[8],
        'oct': months[9],
        'nov': months[10],
        'dic': months[11],
        'tot_sd': tot_sd,
        'tot_year': tot_year,
        'crit': crit
    })

df = pd.DataFrame(rows)

print('=== 1. VERIFICACION CIFRA DE NEGOCIOS ===')
cifra = df[df['n4'] == '1.- Importe neto de la cifra de negocios']['tot_year'].sum()
print(f'Total Cifra de Negocios 2026: {cifra:14,.2f} €  (Objetivo exacto: 15.180.000,00 €)')

print('\n=== 2. CUENTAS MODIFICADAS ESPECIFICAMENTE ===')
target_ctas = [
    '601000000', '602000001', '602000002', '602000003', '607000000',
    '621000014', '622010200', '622010201', '622010300', '623000002',
    '623100002', '623100006', '623100007', '623100008', '624000001',
    '627000008', '631000004'
]

for code in target_ctas:
    sub = df[df['code'] == code]
    if len(sub) > 0:
        r = sub.iloc[0]
        print(f"Cuenta {code} ({r['cta'][:32]}):")
        print(f"  Ene-Ago: {r['tot_8m']:10,.2f} € | Sep: {r['sep']:10,.2f} € | Oct: {r['oct']:10,.2f} € | Nov: {r['nov']:10,.2f} € | Dic: {r['dic']:10,.2f} € | Total 2026: {r['tot_year']:12,.2f} €")
        print(f"  Comentario: {r['crit']}")

print('\n=== 3. RESUMEN P&L 2026 POR NIVEL 4 ===')
pnl_summary = df.groupby('n4')['tot_year'].sum()
for n4, val in pnl_summary.items():
    print(f"{str(n4):45s}: {val:14,.2f} €")

# Calculate Margen Bruto, EBITDA, EBIT
cifra_neg = df[df['n4'] == '1.- Importe neto de la cifra de negocios']['tot_year'].sum()
otros_ingr = df[df['n4'] == '5.- Otros ingresos de explotación']['tot_year'].sum()
var_exist = df[df['n4'] == '2.- Variación de existencias']['tot_year'].sum()
aprov = df[df['n4'] == '4.- Aprovisionamientos']['tot_year'].sum()
personal = df[df['n4'] == '6.- Gastos de personal']['tot_year'].sum()
opex = df[df['n4'] == '7.- Otros gastos de explotación']['tot_year'].sum()
amort = df[df['n4'] == '8.- Amortización del inmovilizado']['tot_year'].sum()
deterioro = df[df['n4'] == '11.- Deterioro y resultado por enajenaciones del inmovilizado']['tot_year'].sum()
otros_res = df[df['n4'] == '13.- Otros resultados']['tot_year'].sum()

ebitda = cifra_neg + otros_ingr + var_exist + aprov + personal + opex + deterioro + otros_res
ebit = ebitda + amort

print('\n=== 4. INDICADORES CLAVE 2026 ===')
print(f"Cifra Neta de Negocios : {cifra_neg:14,.2f} €")
print(f"Aprovisionamientos     : {aprov:14,.2f} €")
print(f"Margen Bruto           : {cifra_neg + aprov + var_exist:14,.2f} € ({(cifra_neg + aprov + var_exist)/cifra_neg*100:5.2f}%)")
print(f"Gastos de Personal     : {personal:14,.2f} €")
print(f"Otros Gastos Explotac. : {opex:14,.2f} €")
print(f"EBITDA 2026            : {ebitda:14,.2f} € ({ebitda/cifra_neg*100:5.2f}%)")
print(f"EBIT 2026              : {ebit:14,.2f} € ({ebit/cifra_neg*100:5.2f}%)")
