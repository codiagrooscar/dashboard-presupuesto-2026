import openpyxl, os, sys
sys.stdout.reconfigure(encoding='utf-8')

comerciales = ['Alfonso', 'Garcia', 'Javier', 'Pedro', 'Ricardo']
folder = 'Revision_Previsiones_Q4_2026'

all_changes = {}

for com in comerciales:
    f_orig = os.path.join(folder, f'Revision_Previsiones_Q4_2026_{com}.xlsx')
    f_copy = os.path.join(folder, f'Copia de Revision_Previsiones_Q4_2026_{com}.xlsx')
    
    print('======================================================================')
    print(f'COMERCIAL: {com}')
    print('======================================================================')
    
    if not os.path.exists(f_copy):
        print(f'File {f_copy} does not exist!')
        continue
        
    wb_orig = openpyxl.load_workbook(f_orig, data_only=True)
    wb_copy = openpyxl.load_workbook(f_copy, data_only=True)
    
    ws_orig = wb_orig.active
    ws_copy = wb_copy.active
    
    print(f'Original rows: {ws_orig.max_row} | Copy rows: {ws_copy.max_row}')
    
    changes = []
    max_r = max(ws_orig.max_row, ws_copy.max_row)
    
    for r in range(10, max_r + 1):
        cliente_o = ws_orig.cell(r, 3).value
        sku_o = ws_orig.cell(r, 4).value
        desc_o = ws_orig.cell(r, 5).value
        
        cliente_c = ws_copy.cell(r, 3).value
        sku_c = ws_copy.cell(r, 4).value
        desc_c = ws_copy.cell(r, 5).value
        
        oct_ant_c = ws_copy.cell(r, 11).value or 0
        oct_rev_c = ws_copy.cell(r, 12).value or 0
        nov_ant_c = ws_copy.cell(r, 13).value or 0
        nov_rev_c = ws_copy.cell(r, 14).value or 0
        dic_ant_c = ws_copy.cell(r, 15).value or 0
        dic_rev_c = ws_copy.cell(r, 16).value or 0
        
        com_sep_c = ws_copy.cell(r, 10).value
        com_q4_c = ws_copy.cell(r, 20).value
        
        oct_rev_o = ws_orig.cell(r, 12).value or 0
        nov_rev_o = ws_orig.cell(r, 14).value or 0
        dic_rev_o = ws_orig.cell(r, 16).value or 0
        com_sep_o = ws_orig.cell(r, 10).value
        com_q4_o = ws_orig.cell(r, 20).value
        
        diff_vs_ant = (oct_rev_c != oct_ant_c) or (nov_rev_c != nov_ant_c) or (dic_rev_c != dic_ant_c)
        diff_vs_orig = (oct_rev_c != oct_rev_o) or (nov_rev_c != nov_rev_o) or (dic_rev_c != dic_rev_o) or (com_sep_c != com_sep_o) or (com_q4_c != com_q4_o)
        
        if diff_vs_orig or diff_vs_ant:
            changes.append({
                'row': r,
                'cliente': cliente_c or cliente_o,
                'sku': sku_c or sku_o,
                'desc': desc_c or desc_o,
                'oct_ant': oct_ant_c, 'oct_rev': oct_rev_c,
                'nov_ant': nov_ant_c, 'nov_rev': nov_rev_c,
                'dic_ant': dic_ant_c, 'dic_rev': dic_rev_c,
                'com_sep': com_sep_c,
                'com_q4': com_q4_c,
                'diff_vs_orig': diff_vs_orig,
                'diff_vs_ant': diff_vs_ant
            })
            
    print(f'Total detected lines with changes/deviations: {len(changes)}')
    for ch in changes:
        c_name = str(ch['cliente'])[:32]
        sku_name = str(ch['sku'])[:10]
        o_diff = f"{ch['oct_ant']}->{ch['oct_rev']}"
        n_diff = f"{ch['nov_ant']}->{ch['nov_rev']}"
        d_diff = f"{ch['dic_ant']}->{ch['dic_rev']}"
        q4_com = str(ch['com_q4'] or '')[:35]
        sep_com = str(ch['com_sep'] or '')[:25]
        print(f"  R{ch['row']:2d} | {c_name:32s} | {sku_name:10s} | Oct:{o_diff:12s} | Nov:{n_diff:12s} | Dic:{d_diff:12s} | Q4:{q4_com}")
        if sep_com:
            print(f"       Comentario Sep: {sep_com}")
    all_changes[com] = changes
