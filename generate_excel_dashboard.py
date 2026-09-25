import sys
import os
import shutil
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
import pandas as pd
import numpy as np
import unicodedata
from datetime import datetime

base_dir = r"c:\Users\oscar.ocampo\OneDrive - Agroquímica Codiagro S.L\Escritorio\Presupuesto 2026\Calsif Comerciales"
os.chdir(base_dir)

def norm(s):
    if not isinstance(s, str): return ''
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn').upper().strip()

print("Cargando y procesando datos...")

def get_latest_pedidos_file():
    candidates = [
        f for f in os.listdir('.')
        if f.startswith('Pedidos') and f.endswith('.xlsx') and not f.startswith(('temp_', '~$'))
    ]
    if not candidates:
        return 'Pedidos 23.09 budget.xlsx'
    def sort_key(f):
        m = re.search(r'(\d{1,2})\.(\d{1,2})', f)
        if m:
            return (int(m.group(2)), int(m.group(1)), os.path.getmtime(f))
        return (0, 0, os.path.getmtime(f))
    candidates.sort(key=sort_key, reverse=True)
    return candidates[0]

pedidos_file = get_latest_pedidos_file()
print(f"Excel Dashboard usando pedidos: {pedidos_file}")
shutil.copy2(pedidos_file, 'temp_pedidos_input.xlsx')
df_ped = pd.read_excel('temp_pedidos_input.xlsx')

m_fc = re.search(r'(\d{1,2})\.(\d{1,2})', pedidos_file)
if m_fc:
    fecha_corte = f"{m_fc.group(1).zfill(2)}/{m_fc.group(2).zfill(2)}/2026"
elif 'FechaPedido' in df_ped.columns and df_ped['FechaPedido'].notna().any():
    max_dt = pd.to_datetime(df_ped['FechaPedido']).max()
    fecha_corte = max_dt.strftime('%d/%m/%Y')
else:
    fecha_corte = datetime.now().strftime('%d/%m/%Y')

def map_client_name(c):
    if not isinstance(c, str): return c
    cn = norm(c)
    if 'DOGA TARIM' in cn:
        return 'ÇITAR ÇIÇEK TARIM TIC. LTD. STI'
    return c

df_ped['Cliente'] = df_ped['Cliente'].apply(map_client_name)
df_ped['CLIENTE_NORM'] = df_ped['Cliente'].apply(norm)
df_ped['SKU_CLEAN'] = df_ped['CodigoArticulo'].astype(str).str.strip().str.upper()

# Excluir Sustainable Agro Solutions
df_ped = df_ped[~df_ped['CLIENTE_NORM'].str.contains('SUSTAINABLE', case=False, na=False)].copy()

# 2. Cargar Budget de forma segura
src_budget = 'Presupuesto_Ventas_2027_Original_Sin_Aplanar.xlsx' if os.path.exists('Presupuesto_Ventas_2027_Original_Sin_Aplanar.xlsx') else 'Presupuesto_Ventas_2027.xlsx'
shutil.copy2(src_budget, 'temp_budget_input.xlsx')
df_bud = pd.read_excel('temp_budget_input.xlsx', sheet_name='Previsión Matriz Horizontal')
df_bud = df_bud[df_bud['Comercial'].notna() & (~df_bud['Comercial'].astype(str).str.contains('TOTAL', case=False))].copy()

FACTOR = 0.9479961392983893
INV_FACTOR = 1.0 / FACTOR
for idx, row in df_bud.iterrows():
    com = str(row['Comercial']).strip()
    if 'garc' not in norm(com).lower():
        v = row['Sep-26 (u)']
        if pd.notna(v) and v > 0 and abs(v - round(v)) > 0.04:
            restored = round(v * INV_FACTOR, 2)
            if abs(restored - round(restored)) < 0.05:
                restored = float(round(restored))
            df_bud.at[idx, 'Sep-26 (u)'] = restored
df_bud = df_bud[df_bud['Comercial'].notna() & (~df_bud['Comercial'].astype(str).str.contains('TOTAL', case=False))].copy()
df_bud['Cliente'] = df_bud['Cliente'].apply(map_client_name)
df_bud['CLIENTE_NORM'] = df_bud['Cliente'].apply(norm)

def clean_budget_sku(row):
    sku = str(row['Código Artículo']).strip().upper()
    pais = str(row['País']).strip().upper()
    if pais != 'ESPAÑA' and len(sku) > 2:
        return sku[:-2]
    return sku

df_bud['SKU_CLEAN'] = df_bud.apply(clean_budget_sku, axis=1)

# Diccionario Cliente -> Comercial
client_map = df_bud.groupby('CLIENTE_NORM')['Comercial'].first().to_dict()
# Asignaciones explícitas de Javier
client_map[norm('FINCA DOÑA ANA C.B.')] = 'Javier'
client_map[norm('FITOSANITARIOS CARCAIXENT, S.L.')] = 'Javier'
client_map[norm('ALMENDRALIA IBÉRICA, S.L.U.')] = 'Javier'

df_ped['Comercial'] = df_ped['CLIENTE_NORM'].map(client_map).fillna('Sin Asignar')

comerciales = ['Alfonso', 'García', 'Irene', 'Javier', 'Mehmet', 'Pedro', 'Ricardo']

# Estilos Openpyxl
font_title = Font(name='Segoe UI', size=16, bold=True, color='FFFFFF')
font_subtitle = Font(name='Segoe UI', size=10, italic=True, color='94A3B8')
font_card_lbl = Font(name='Segoe UI', size=9, bold=True, color='64748B')
font_header = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
font_bold = Font(name='Segoe UI', size=10, bold=True, color='0F172A')
font_regular = Font(name='Segoe UI', size=10, color='1E293B')

fill_dark = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
fill_header = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
fill_subtotal = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
fill_card_kpi = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
fill_card_accent = PatternFill(start_color='ECFDF5', end_color='ECFDF5', fill_type='solid')
fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

fill_success = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')
font_success = Font(name='Segoe UI', size=9, bold=True, color='166534')

fill_warning = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
font_warning = Font(name='Segoe UI', size=9, bold=True, color='92400E')

fill_danger = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
font_danger = Font(name='Segoe UI', size=9, bold=True, color='991B1B')

fill_neutral = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
font_neutral = Font(name='Segoe UI', size=9, color='64748B')

border_thin = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)
border_double = Border(
    top=Side(style='thin', color='94A3B8'),
    bottom=Side(style='double', color='0F172A'),
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1')
)

align_center = Alignment(horizontal='center', vertical='center')
align_left = Alignment(horizontal='left', vertical='center')
align_right = Alignment(horizontal='right', vertical='center')

wb = openpyxl.Workbook()
wb.remove(wb.active)

# ==============================================================================
# 1. HOJA RESUMEN GENERAL (SÓLO UNIDADES)
# ==============================================================================
ws_res = wb.create_sheet(title='Resumen General')
ws_res.views.sheetView[0].showGridLines = True

# Banner Cabecera
ws_res.merge_cells('B2:J2')
ws_res['B2'] = "CODIAGRO - SEGUIMIENTO DIARIO DE VENTAS (SEPTIEMBRE 2026)"
ws_res['B2'].font = font_title
ws_res['B2'].fill = fill_dark
ws_res['B2'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
ws_res.row_dimensions[2].height = 40

ws_res.merge_cells('B3:J3')
ws_res['B3'] = f"Comparativa Estimación Septiembre (Unidades) vs Pedidos Reales a Fecha {fecha_corte} | Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
ws_res['B3'].font = font_subtitle
ws_res['B3'].fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
ws_res['B3'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
ws_res.row_dimensions[3].height = 20

# KPI Cards (SÓLO UNIDADES)
tot_bud = df_bud['Sep-26 (u)'].sum()
tot_ped = df_ped['Pedidas'].sum()
tot_serv = df_ped['Servidas'].sum()
tot_pend = df_ped['Pendientes'].sum()
tot_gap = sum(max(0.0, df_bud[df_bud['Comercial']==c]['Sep-26 (u)'].sum() - df_ped[df_ped['Comercial']==c]['Pedidas'].sum()) for c in comerciales)
tot_pct = (tot_ped / tot_bud * 100) if tot_bud > 0 else 0.0

kpi_data = [
    ('B', 'C', 'ESTIMACIÓN SEP (UDS)', f"{tot_bud:,.0f} u", fill_card_kpi, '0F172A'),
    ('D', 'E', 'PEDIDOS REALES (UDS)', f"{tot_ped:,.0f} u", fill_card_accent, '047857'),
    ('F', 'G', '% CONSECUCIÓN GLOBAL', f"{tot_pct:.1f}%", fill_card_accent if tot_pct>=100 else fill_card_kpi, '047857' if tot_pct>=100 else 'B45309'),
    ('H', 'I', 'FALTA CONSEGUIR (GAP)', f"{tot_gap:,.0f} u", fill_warning if tot_gap>0 else fill_card_accent, 'B45309' if tot_gap>0 else '047857'),
    ('J', 'J', 'PENDIENTES (U)', f"{tot_pend:,.0f} u", fill_card_kpi, '0F172A')
]

for col1, col2, title, val, bg_fill, text_color in kpi_data:
    if col1 != col2:
        ws_res.merge_cells(f"{col1}5:{col2}5")
        ws_res.merge_cells(f"{col1}6:{col2}6")
    c_t = ws_res[f"{col1}5"]
    c_v = ws_res[f"{col1}6"]
    c_t.value = title
    c_t.font = font_card_lbl
    c_t.fill = bg_fill
    c_t.alignment = align_center
    c_v.value = val
    c_v.font = Font(name='Segoe UI', size=14, bold=True, color=text_color)
    c_v.fill = bg_fill
    c_v.alignment = align_center
    
    cols = [col1] if col1 == col2 else [col1, col2]
    for col_letter in cols:
        for r_idx in [5, 6]:
            ws_res[f"{col_letter}{r_idx}"].border = border_thin

ws_res.row_dimensions[5].height = 18
ws_res.row_dimensions[6].height = 28

# Tabla Resumen por Comercial (SÓLO UNIDADES)
ws_res['B8'] = "RESUMEN CONSOLIDADO POR COMERCIAL (UNIDADES)"
ws_res['B8'].font = Font(name='Segoe UI', size=12, bold=True, color='0F172A')

headers_res = [
    'Comercial', 'Clientes Activos', 'Estimación Sep (u)', 'Pedidos Sep (u)',
    'Servidas (u)', 'Pendientes (u)', 'Falta Conseguir (u)', 'Desviación (u)',
    '% Consecución'
]

start_row_res = 9
for c_idx, h in enumerate(headers_res, start=2):
    cell = ws_res.cell(start_row_res, c_idx)
    cell.value = h
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = border_thin
ws_res.row_dimensions[start_row_res].height = 25

cur_row = start_row_res + 1

for c in comerciales:
    sub_ped = df_ped[df_ped['Comercial'] == c]
    sub_bud = df_bud[df_bud['Comercial'] == c]
    
    b_uds = sub_bud['Sep-26 (u)'].sum()
    p_uds = sub_ped['Pedidas'].sum()
    s_uds = sub_ped['Servidas'].sum()
    pend_uds = sub_ped['Pendientes'].sum()
    cli_cnt = sub_ped['Cliente'].nunique()
    
    ws_res.cell(cur_row, 2, c).font = font_bold
    ws_res.cell(cur_row, 3, cli_cnt).font = font_regular
    ws_res.cell(cur_row, 4, b_uds).font = font_regular
    ws_res.cell(cur_row, 5, p_uds).font = font_regular
    ws_res.cell(cur_row, 6, s_uds).font = font_regular
    ws_res.cell(cur_row, 7, pend_uds).font = font_regular
    ws_res.cell(cur_row, 8, f"=MAX(0, D{cur_row}-E{cur_row})").font = font_regular
    ws_res.cell(cur_row, 9, f"=E{cur_row}-D{cur_row}").font = font_regular
    ws_res.cell(cur_row, 10, f"=IF(D{cur_row}>0, E{cur_row}/D{cur_row}, IF(E{cur_row}>0, 1, 0))").font = font_bold
    
    ws_res.cell(cur_row, 2).alignment = align_left
    ws_res.cell(cur_row, 3).alignment = align_center
    for col_i in range(4, 10):
        c_cell = ws_res.cell(cur_row, col_i)
        c_cell.alignment = align_right
        c_cell.number_format = '#,##0.00'
    
    ws_res.cell(cur_row, 10).alignment = align_center
    ws_res.cell(cur_row, 10).number_format = '0.0%'
    
    if cur_row % 2 == 1:
        for c_i in range(2, 11):
            ws_res.cell(cur_row, c_i).fill = fill_zebra
            
    for c_i in range(2, 11):
        ws_res.cell(cur_row, c_i).border = border_thin
        
    cur_row += 1

# Total Row
ws_res.cell(cur_row, 2, "TOTAL GENERAL").font = font_bold
ws_res.cell(cur_row, 3, f"=SUM(C10:C{cur_row-1})").font = font_bold
ws_res.cell(cur_row, 4, f"=SUM(D10:D{cur_row-1})").font = font_bold
ws_res.cell(cur_row, 5, f"=SUM(E10:E{cur_row-1})").font = font_bold
ws_res.cell(cur_row, 6, f"=SUM(F10:F{cur_row-1})").font = font_bold
ws_res.cell(cur_row, 7, f"=SUM(G10:G{cur_row-1})").font = font_bold
ws_res.cell(cur_row, 8, f"=SUM(H10:H{cur_row-1})").font = font_bold
ws_res.cell(cur_row, 9, f"=SUM(I10:I{cur_row-1})").font = font_bold
ws_res.cell(cur_row, 10, f"=IF(D{cur_row}>0, E{cur_row}/D{cur_row}, 0)").font = font_bold

ws_res.cell(cur_row, 2).alignment = align_left
ws_res.cell(cur_row, 3).alignment = align_center
for col_i in range(4, 10):
    ws_res.cell(cur_row, col_i).alignment = align_right
    ws_res.cell(cur_row, col_i).number_format = '#,##0.00'
ws_res.cell(cur_row, 10).alignment = align_center
ws_res.cell(cur_row, 10).number_format = '0.0%'

for c_i in range(2, 11):
    ws_res.cell(cur_row, c_i).fill = fill_subtotal
    ws_res.cell(cur_row, c_i).border = border_double
ws_res.row_dimensions[cur_row].height = 24

# Añadir Gráfico Comparativo de Barras
chart1 = BarChart()
chart1.type = "col"
chart1.style = 10
chart1.title = "Estimación Septiembre vs Pedidos Reales por Comercial (Unidades)"
chart1.y_axis.title = "Unidades"
chart1.x_axis.title = "Comercial"
chart1.width = 18
chart1.height = 11

data_ref = Reference(ws_res, min_col=4, min_row=9, max_col=5, max_row=16)
cats_ref = Reference(ws_res, min_col=2, min_row=10, max_row=16)
chart1.add_data(data_ref, titles_from_data=True)
chart1.set_categories(cats_ref)
ws_res.add_chart(chart1, "B19")

col_widths_res = {'B': 18, 'C': 16, 'D': 16, 'E': 16, 'F': 15, 'G': 15, 'H': 18, 'I': 16, 'J': 15}
for c_l, w in col_widths_res.items():
    ws_res.column_dimensions[c_l].width = w
ws_res.column_dimensions['A'].width = 3

# ==============================================================================
# 2. HOJAS INDIVIDUALES POR COMERCIAL (SÓLO UNIDADES)
# ==============================================================================
for c in comerciales:
    ws = wb.create_sheet(title=c)
    ws.views.sheetView[0].showGridLines = True
    
    sub_ped = df_ped[df_ped['Comercial'] == c]
    sub_bud = df_bud[df_bud['Comercial'] == c]
    
    com_bud_tot = sub_bud['Sep-26 (u)'].sum()
    com_ped_tot = sub_ped['Pedidas'].sum()
    com_serv_tot = sub_ped['Servidas'].sum()
    com_pend_tot = sub_ped['Pendientes'].sum()
    com_gap_tot = max(0.0, com_bud_tot - com_ped_tot)
    com_pct_tot = (com_ped_tot / com_bud_tot * 100) if com_bud_tot > 0 else (100.0 if com_ped_tot > 0 else 0.0)
    
    # Header Banner
    ws.merge_cells('B2:L2')
    ws['B2'] = f"CODIAGRO - SEGUIMIENTO COMERCIAL: {c.upper()} (UNIDADES)"
    ws['B2'].font = font_title
    ws['B2'].fill = fill_dark
    ws['B2'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[2].height = 36
    
    ws.merge_cells('B3:L3')
    ws['B3'] = f"Seguimiento diario de Previsión Septiembre 2026 vs Pedidos a Fecha {fecha_corte}"
    ws['B3'].font = font_subtitle
    ws['B3'].fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    ws['B3'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[3].height = 18
    
    # KPI Cards para el Comercial (SÓLO UNIDADES)
    cards = [
        ('B', 'C', 'ESTIMACIÓN SEP (U)', f"{com_bud_tot:,.0f} u", fill_card_kpi, '0F172A'),
        ('D', 'E', 'PEDIDOS REALES (U)', f"{com_ped_tot:,.0f} u", fill_card_accent if com_pct_tot>=100 else fill_card_kpi, '047857' if com_pct_tot>=100 else '0F172A'),
        ('F', 'G', '% CONSECUCIÓN', f"{com_pct_tot:.1f}%", fill_card_accent if com_pct_tot>=100 else fill_warning, '047857' if com_pct_tot>=100 else 'B45309'),
        ('H', 'I', 'FALTA CONSEGUIR (GAP)', f"{com_gap_tot:,.0f} u", fill_warning if com_gap_tot>0 else fill_card_accent, 'B45309' if com_gap_tot>0 else '047857'),
        ('J', 'K', 'PENDIENTES SERVIR (U)', f"{com_pend_tot:,.0f} u", fill_card_kpi, '0F172A')
    ]
    for c1, c2, title, val, bg_f, col_t in cards:
        ws.merge_cells(f"{c1}5:{c2}5")
        ws.merge_cells(f"{c1}6:{c2}6")
        ws[f"{c1}5"].value = title
        ws[f"{c1}5"].font = font_card_lbl
        ws[f"{c1}5"].fill = bg_f
        ws[f"{c1}5"].alignment = align_center
        ws[f"{c1}6"].value = val
        ws[f"{c1}6"].font = Font(name='Segoe UI', size=13, bold=True, color=col_t)
        ws[f"{c1}6"].fill = bg_f
        ws[f"{c1}6"].alignment = align_center
        for cl in [c1, c2]:
            ws[f"{cl}5"].border = border_thin
            ws[f"{cl}6"].border = border_thin
            
    ws.row_dimensions[5].height = 16
    ws.row_dimensions[6].height = 25
    
    # Detalle Cliente - Artículo (SÓLO UNIDADES)
    ws['B8'] = f"DETALLE COMPARATIVO POR CLIENTE Y ARTÍCULO - {c.upper()} (UNIDADES)"
    ws['B8'].font = Font(name='Segoe UI', size=11, bold=True, color='0F172A')
    
    det_headers = [
        'Cliente', 'Cód. Artículo', 'Descripción Artículo', 'Estimación Sep (u)',
        'Pedidos Sep (u)', 'Servidas (u)', 'Pendientes (u)', 'Falta (u)',
        'Desv. (u)', '% Consec.', 'Estado'
    ]
    start_r = 9
    for c_idx, h in enumerate(det_headers, start=2):
        cell = ws.cell(start_r, c_idx)
        cell.value = h
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_thin
    ws.row_dimensions[start_r].height = 25
    
    keys_set = set()
    for _, r in sub_bud.iterrows():
        keys_set.add((r['Cliente'], r['SKU_CLEAN']))
    for _, r in sub_ped.iterrows():
        keys_set.add((r['Cliente'], r['SKU_CLEAN']))
        
    keys_sorted = sorted(list(keys_set), key=lambda x: (x[0], x[1]))
    
    row_num = start_r + 1
    for cli, sku in keys_sorted:
        b_match = sub_bud[(sub_bud['Cliente'] == cli) & (sub_bud['SKU_CLEAN'] == sku)]
        p_match = sub_ped[(sub_ped['Cliente'] == cli) & (sub_ped['SKU_CLEAN'] == sku)]
        
        desc = ""
        if len(b_match) > 0 and pd.notna(b_match['Descripción Artículo'].iloc[0]):
            desc = str(b_match['Descripción Artículo'].iloc[0])
        elif len(p_match) > 0 and pd.notna(p_match['Articulo'].iloc[0]):
            desc = str(p_match['Articulo'].iloc[0])
            
        b_u = b_match['Sep-26 (u)'].sum() if len(b_match) > 0 else 0.0
        p_u = p_match['Pedidas'].sum() if len(p_match) > 0 else 0.0
        s_u = p_match['Servidas'].sum() if len(p_match) > 0 else 0.0
        pend_u = p_match['Pendientes'].sum() if len(p_match) > 0 else 0.0
        
        # Excluir líneas con todo 0 (sin budget y sin pedidos en Septiembre)
        if b_u == 0 and p_u == 0:
            continue
        
        ws.cell(row_num, 2, cli).font = font_regular
        ws.cell(row_num, 3, sku).font = font_bold
        ws.cell(row_num, 4, desc).font = font_regular
        ws.cell(row_num, 5, b_u).font = font_regular
        ws.cell(row_num, 6, p_u).font = font_regular
        ws.cell(row_num, 7, s_u).font = font_regular
        ws.cell(row_num, 8, pend_u).font = font_regular
        ws.cell(row_num, 9, f"=MAX(0, E{row_num}-F{row_num})").font = font_regular
        ws.cell(row_num, 10, f"=F{row_num}-E{row_num}").font = font_regular
        ws.cell(row_num, 11, f"=IF(E{row_num}>0, F{row_num}/E{row_num}, IF(F{row_num}>0, 1, 0))").font = font_bold
        
        pct_val = (p_u / b_u * 100) if b_u > 0 else (100.0 if p_u > 0 else 0.0)
        status_cell = ws.cell(row_num, 12)
        if b_u == 0 and p_u > 0:
            status_cell.value = "Extra Estimación"
            status_cell.fill = fill_success
            status_cell.font = font_success
        elif pct_val >= 100:
            status_cell.value = "Superado"
            status_cell.fill = fill_success
            status_cell.font = font_success
        elif pct_val >= 50:
            status_cell.value = "En Curso"
            status_cell.fill = fill_warning
            status_cell.font = font_warning
        elif p_u > 0:
            status_cell.value = "Rezagado"
            status_cell.fill = fill_danger
            status_cell.font = font_danger
        else:
            status_cell.value = "Sin Pedido"
            status_cell.fill = fill_neutral
            status_cell.font = font_neutral
            
        status_cell.alignment = align_center
        
        ws.cell(row_num, 2).alignment = align_left
        ws.cell(row_num, 3).alignment = align_center
        ws.cell(row_num, 4).alignment = align_left
        for ci in range(5, 11):
            ws.cell(row_num, ci).alignment = align_right
            ws.cell(row_num, ci).number_format = '#,##0.00'
        ws.cell(row_num, 11).alignment = align_center
        ws.cell(row_num, 11).number_format = '0.0%'
        
        if row_num % 2 == 1:
            for ci in range(2, 12):
                ws.cell(row_num, ci).fill = fill_zebra
                
        for ci in range(2, 13):
            ws.cell(row_num, ci).border = border_thin
            
        row_num += 1
        
    # Fila de Totales del Comercial
    ws.cell(row_num, 2, f"TOTAL {c.upper()}").font = font_bold
    ws.cell(row_num, 3, f"{len(keys_sorted)} líneas").font = font_subtitle
    ws.cell(row_num, 4, "").font = font_regular
    ws.cell(row_num, 5, f"=SUM(E10:E{row_num-1})").font = font_bold
    ws.cell(row_num, 6, f"=SUM(F10:F{row_num-1})").font = font_bold
    ws.cell(row_num, 7, f"=SUM(G10:G{row_num-1})").font = font_bold
    ws.cell(row_num, 8, f"=SUM(H10:H{row_num-1})").font = font_bold
    ws.cell(row_num, 9, f"=SUM(I10:I{row_num-1})").font = font_bold
    ws.cell(row_num, 10, f"=SUM(J10:J{row_num-1})").font = font_bold
    ws.cell(row_num, 11, f"=IF(E{row_num}>0, F{row_num}/E{row_num}, 0)").font = font_bold
    ws.cell(row_num, 12, f"=IF(K{row_num}>=1, \"SUPERADO\", \"EN CURSO\")").font = font_bold
    
    ws.cell(row_num, 2).alignment = align_left
    ws.cell(row_num, 3).alignment = align_center
    for ci in range(5, 11):
        ws.cell(row_num, ci).alignment = align_right
        ws.cell(row_num, ci).number_format = '#,##0.00'
    ws.cell(row_num, 11).alignment = align_center
    ws.cell(row_num, 11).number_format = '0.0%'
    ws.cell(row_num, 12).alignment = align_center
    
    for ci in range(2, 13):
        ws.cell(row_num, ci).fill = fill_subtotal
        ws.cell(row_num, ci).border = border_double
    ws.row_dimensions[row_num].height = 24
    
    col_w_com = {
        'B': 30, 'C': 14, 'D': 28, 'E': 14, 'F': 14, 'G': 13,
        'H': 13, 'I': 13, 'J': 13, 'K': 11, 'L': 13
    }
    for col_l, w in col_w_com.items():
        ws.column_dimensions[col_l].width = w
    ws.column_dimensions['A'].width = 3

output_excel = "Seguimiento_Presupuesto_Sep_2026.xlsx"
try:
    wb.save(output_excel)
    print(f"¡Archivo Excel regenerado (sólo unidades): {output_excel}!")
except PermissionError:
    output_fallback = "Seguimiento_Presupuesto_Sep_2026_Unidades.xlsx"
    wb.save(output_fallback)
    print(f"Nota: {output_excel} estaba abierto. Guardado como: {output_fallback}")
