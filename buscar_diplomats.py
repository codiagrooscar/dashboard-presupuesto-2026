import openpyxl, subprocess, os, json

f1 = r'C:\Users\oscar.ocampo\OneDrive - Agroquímica Codiagro S.L\Escritorio\REPORTING SIMPLIFICADO\SEPTIEMBRE\Cashflow\Cashflow previsión cierre 2026.xlsx'
subprocess.run(['powershell', '-Command', f'Copy-Item -LiteralPath "{f1}" -Destination "temp_cf_diplomats.xlsx" -Force'], check=True)
wb = openpyxl.load_workbook('temp_cf_diplomats.xlsx', data_only=True)

for sname in wb.sheetnames:
    ws = wb[sname]
    for r_idx, r in enumerate(ws.iter_rows(values_only=True), 1):
        txt = ' '.join(str(x) for x in r if x is not None).upper()
        if 'DIPLOMAT' in txt:
            print(f'[{sname}] Row {r_idx}: {r}')

with open('web_dashboard/treasury_data.json', 'r', encoding='utf-8') as f:
    t_data = json.load(f)

for v in t_data.get('client_vencimientos', []):
    if 'DIPLOMAT' in str(v.get('cliente', '')).upper():
        print('JSON Vencimiento:', v)

for p in t_data.get('payment_terms', []):
    if 'DIPLOMAT' in str(p.get('cliente', '')).upper():
        print('JSON Payment Term:', p)
