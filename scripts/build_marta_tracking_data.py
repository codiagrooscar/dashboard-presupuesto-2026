try:
    import sys
    sys.path.insert(0, os.path.dirname(__file__))
    from output_staging import get_dest
except Exception:
    def get_dest(p): return p

import os
import json
import subprocess
import openpyxl
from collections import defaultdict


def round_conserving_sum(items, field, target_total, decimals=2):
    """
    Política explícita de redondeo por Mayor Resto (Largest Remainder Method):
    conserva exactamente la suma `target_total` entre los detalles redondeados.
    """
    if not items:
        return
    factor = 10 ** decimals
    target_cents = int(round(target_total * factor))
    exact_vals = [float(item.get(field, 0.0)) * factor for item in items]
    floored = [int(x // 1) for x in exact_vals]
    remainders = [exact_vals[i] - floored[i] for i in range(len(items))]
    diff = target_cents - sum(floored)
    if diff > 0:
        order = sorted(range(len(items)), key=lambda i: remainders[i], reverse=True)
        for idx in order[:diff]:
            floored[idx] += 1
    elif diff < 0:
        order = sorted(range(len(items)), key=lambda i: remainders[i])
        for idx in order[:abs(diff)]:
            floored[idx] -= 1
    for i, item in enumerate(items):
        item[field] = round(floored[i] / factor, decimals)


def build_data():
    print("=== Procesando fichero de seguimiento de Cierre 2026 (UNIDADES & EUROS) ===")

    # 1. Cargar ERP DATOS (Ventas Septiembre 2026.xlsx)
    marta_sales_file = r"C:\Users\oscar.ocampo\OneDrive - Agroquímica Codiagro S.L\Escritorio\REPORTING SIMPLIFICADO\SEPTIEMBRE\Marta y Pere\Ventas Septiembre 2026.xlsx"
    wb_erp = openpyxl.load_workbook(marta_sales_file, read_only=True, data_only=True)
    ws_datos = wb_erp['DATOS']

    erp_u = defaultdict(lambda: {'2025_tot': 0.0, '2025_sep': 0.0, '2025_oct': 0.0, '2026_sep': 0.0})
    sas_u_2026 = 0.0
    sas_eur_2026 = 0.0

    for i, r in enumerate(ws_datos.iter_rows(values_only=True)):
        if i == 0:
            continue
        ej = r[6] # Ejercicio
        p = r[7]  # Periodo
        cod_cl = str(r[16] or '').strip()
        cl_nom = str(r[35] or '').strip().upper()
        try:
            u = float(r[20] or 0) # Unidades
        except:
            u = 0.0
        try:
            eur = float(r[22] or 0) # ImporteNeto
        except:
            eur = 0.0

        if cod_cl == '430000427' or 'SUSTAINABLE' in cl_nom:
            if ej == 2026 and p and int(p) <= 9:
                sas_u_2026 += u
                sas_eur_2026 += eur
            continue

        if ej == 2025:
            erp_u[cl_nom]['2025_tot'] += u
            if p and int(p) <= 9:
                erp_u[cl_nom]['2025_sep'] += u
            else:
                erp_u[cl_nom]['2025_oct'] += u
        elif ej == 2026 and p and int(p) <= 9:
            erp_u[cl_nom]['2026_sep'] += u

    print(f"Total clientes ERP cargados: {len(erp_u)}")
    print(f"Unidades ERP 2025: {sum(d['2025_tot'] for d in erp_u.values()):,.2f} u")
    print(f"Unidades ERP 2026 Ene-Sep: {sum(d['2026_sep'] for d in erp_u.values()):,.2f} u")
    print(f"SAS 2026: {sas_u_2026:,.2f} u | {sas_eur_2026:,.2f} €")

    # 2. Cargar Presupuesto 2026 en Unidades (Presupuesto_Ventas_2027_Definitivo.xlsx)
    src_ppto = 'Presupuesto_Ventas_2027_Definitivo.xlsx'
    try:
        wb_ppto = openpyxl.load_workbook(src_ppto, data_only=True)
    except Exception:
        tmp_ppto = 'temp_safe_read_marta_ppto.xlsx'
        subprocess.run(['powershell', '-Command', f'Copy-Item "{src_ppto}" "{tmp_ppto}" -Force'], check=True)
        wb_ppto = openpyxl.load_workbook(tmp_ppto, data_only=True)
    ws_rc = wb_ppto['Resumen por Cliente']
    h_rc = [c.value for c in ws_rc[1]]
    idx_cl = h_rc.index('Cliente')
    idx_ea_u = h_rc.index('Real Ene-Ago 26 (u)')
    idx_sd_u = h_rc.index('Prev Sep-Dic 26 (u)')

    ppto_u = {}
    for r in range(2, ws_rc.max_row + 1):
        cl = str(ws_rc.cell(r, idx_cl + 1).value or '').strip().upper()
        ea_u = float(ws_rc.cell(r, idx_ea_u + 1).value or 0)
        sd_u = float(ws_rc.cell(r, idx_sd_u + 1).value or 0)
        ppto_u[cl] = {
            'ea_u': ea_u,
            'sd_u': sd_u,
            'tot_u': ea_u + sd_u
        }

    # 3. Cargar plantilla base de Marta (Ventas_2025_vs_2026_vs_Estimacion_por_Cliente MArta.xlsx)
    src_file = "Ventas_2025_vs_2026_vs_Estimacion_por_Cliente MArta.xlsx"
    if not os.path.exists(src_file) and os.path.exists(os.path.join("antiguos", src_file)):
        src_file = os.path.join("antiguos", src_file)
    temp_file = "temp_marta_export.xlsx"
    subprocess.run(['powershell', '-Command', f'Copy-Item "{src_file}" "{temp_file}" -Force'], check=True)
    wb_m = openpyxl.load_workbook(temp_file, data_only=True)
    ws_mc = wb_m['Comparativa Clientes']

    # --- A) GENERAR DATASET UNIDADES ---
    clientes_u = []
    by_com_u = defaultdict(lambda: {
        'comercial': '', 'num_clientes': 0,
        'v25_u_tot': 0.0, 'v25_u_sep': 0.0, 'v25_u_oct': 0.0,
        'v26_u_sep': 0.0, 'est26_u': 0.0
    })
    conteos_val_u = {}

    # --- B) GENERAR DATASET EUROS ---
    clientes_eur = []
    by_com_eur = defaultdict(lambda: {
        'comercial': '', 'num_clientes': 0,
        'v25_eur_tot': 0.0, 'v25_eur_sep': 0.0, 'v25_eur_oct': 0.0,
        'v26_eur_sep': 0.0, 'est26_eur': 0.0
    })
    conteos_val_eur = {}

    for r in range(8, ws_mc.max_row + 1):
        com = str(ws_mc.cell(r, 1).value or '').strip()
        pais = str(ws_mc.cell(r, 2).value or '').strip()
        cl_raw = str(ws_mc.cell(r, 3).value or '').strip()
        clasif = str(ws_mc.cell(r, 4).value or '').strip()
        if not cl_raw:
            continue

        cl_up = cl_raw.upper()
        if 'HORTIFRUT' in cl_up:
            cl_erp = 'HORTIFRUT - PERU, S.A.C.'
        elif 'POULTSIDIS' in cl_up:
            cl_erp = 'GEORGIOS POULTSIDIS'
        else:
            cl_erp = cl_up

        # --- Datos Unidades ---
        d_erp = erp_u.get(cl_erp, {'2025_tot': 0.0, '2025_sep': 0.0, '2025_oct': 0.0, '2026_sep': 0.0})
        p_info = ppto_u.get(cl_up, {'ea_u': 0.0, 'sd_u': 0.0, 'tot_u': 0.0})

        v25_u_tot = d_erp['2025_tot']
        v25_u_sep = d_erp['2025_sep']
        v25_u_oct = d_erp['2025_oct']
        v26_u_sep = d_erp['2026_sep']
        est26_u = p_info['tot_u']

        pend_u = est26_u - v26_u_sep
        pct_vend_u = (v26_u_sep / est26_u * 100.0) if est26_u > 0 else None
        crec_real_sep_u = ((v26_u_sep / v25_u_sep) - 1.0) * 100.0 if v25_u_sep > 0 else None
        crec_nec_oct_u = ((pend_u / v25_u_oct) - 1.0) * 100.0 if v25_u_oct > 0 else None

        if est26_u == 0:
            val_u = "Sin estimación"
        elif pend_u <= 0:
            val_u = "Ya cumplido"
        elif v25_u_oct == 0:
            val_u = "Sin base 2025"
        elif crec_nec_oct_u is not None:
            if crec_nec_oct_u <= 15.0:
                val_u = "Alcanzable"
            elif crec_nec_oct_u <= 40.0:
                val_u = "Exigente"
            else:
                val_u = "Poco realista"
        else:
            val_u = "Sin base 2025"

        conteos_val_u[val_u] = conteos_val_u.get(val_u, 0) + 1

        clientes_u.append({
            "comercial": com,
            "pais": pais,
            "cliente": cl_raw,
            "clasificacion": clasif,
            "ventas_2025_tot": v25_u_tot,
            "ventas_2025_sep": v25_u_sep,
            "ventas_2025_oct": v25_u_oct,
            "ventas_2026_sep": v26_u_sep,
            "estimacion_2026": est26_u,
            "pendiente_oct_dic": pend_u,
            "pct_vendido": round(pct_vend_u, 1) if pct_vend_u is not None else None,
            "crec_real_sep_pct": round(crec_real_sep_u, 1) if crec_real_sep_u is not None else None,
            "crec_nec_oct_pct": round(crec_nec_oct_u, 1) if crec_nec_oct_u is not None else None,
            "valoracion": val_u
        })

        by_com_u[com]['comercial'] = com
        by_com_u[com]['num_clientes'] += 1
        by_com_u[com]['v25_u_tot'] += v25_u_tot
        by_com_u[com]['v25_u_sep'] += v25_u_sep
        by_com_u[com]['v25_u_oct'] += v25_u_oct
        by_com_u[com]['v26_u_sep'] += v26_u_sep
        by_com_u[com]['est26_u'] += est26_u

        # --- Datos Euros (Directos de la hoja de Marta) ---
        v25_eur_tot = float(ws_mc.cell(r, 5).value or 0)
        v25_eur_sep = float(ws_mc.cell(r, 6).value or 0)
        v25_eur_oct = float(ws_mc.cell(r, 7).value or 0)
        v26_eur_sep = float(ws_mc.cell(r, 8).value or 0)
        est26_eur = float(ws_mc.cell(r, 9).value or 0)
        pend_eur = est26_eur - v26_eur_sep

        pct_raw = ws_mc.cell(r, 14).value
        if pct_raw not in [None, 'n/a', '']:
            try: pct_vend_eur = float(pct_raw) * 100.0 if float(pct_raw) <= 5.0 else float(pct_raw)
            except: pct_vend_eur = (v26_eur_sep / est26_eur * 100.0) if est26_eur > 0 else None
        else:
            pct_vend_eur = (v26_eur_sep / est26_eur * 100.0) if est26_eur > 0 else None

        crec_sep_raw = ws_mc.cell(r, 10).value
        if crec_sep_raw not in [None, 'n/a', '']:
            try: crec_real_sep_eur = float(crec_sep_raw) * 100.0 if abs(float(crec_sep_raw)) <= 10.0 else float(crec_sep_raw)
            except: crec_real_sep_eur = ((v26_eur_sep / v25_eur_sep) - 1.0) * 100.0 if v25_eur_sep > 0 else None
        else:
            crec_real_sep_eur = ((v26_eur_sep / v25_eur_sep) - 1.0) * 100.0 if v25_eur_sep > 0 else None

        crec_oct_raw = ws_mc.cell(r, 13).value
        if crec_oct_raw not in [None, 'n/a', '']:
            try: crec_nec_oct_eur = float(crec_oct_raw) * 100.0 if abs(float(crec_oct_raw)) <= 20.0 else float(crec_oct_raw)
            except: crec_nec_oct_eur = ((pend_eur / v25_eur_oct) - 1.0) * 100.0 if v25_eur_oct > 0 else None
        else:
            crec_nec_oct_eur = ((pend_eur / v25_eur_oct) - 1.0) * 100.0 if v25_eur_oct > 0 else None

        val_eur = str(ws_mc.cell(r, 15).value or 'Sin estimación').strip()
        conteos_val_eur[val_eur] = conteos_val_eur.get(val_eur, 0) + 1

        clientes_eur.append({
            "comercial": com,
            "pais": pais,
            "cliente": cl_raw,
            "clasificacion": clasif,
            "ventas_2025_tot": v25_eur_tot,
            "ventas_2025_sep": v25_eur_sep,
            "ventas_2025_oct": v25_eur_oct,
            "ventas_2026_sep": v26_eur_sep,
            "estimacion_2026": est26_eur,
            "pendiente_oct_dic": pend_eur,
            "pct_vendido": round(pct_vend_eur, 1) if pct_vend_eur is not None else None,
            "crec_real_sep_pct": round(crec_real_sep_eur, 1) if crec_real_sep_eur is not None else None,
            "crec_nec_oct_pct": round(crec_nec_oct_eur, 1) if crec_nec_oct_eur is not None else None,
            "valoracion": val_eur
        })

        by_com_eur[com]['comercial'] = com
        by_com_eur[com]['num_clientes'] += 1
        by_com_eur[com]['v25_eur_tot'] += v25_eur_tot
        by_com_eur[com]['v25_eur_sep'] += v25_eur_sep
        by_com_eur[com]['v25_eur_oct'] += v25_eur_oct
        by_com_eur[com]['v26_eur_sep'] += v26_eur_sep
        by_com_eur[com]['est26_eur'] += est26_eur

    # 4. Totales y Comerciales para UNIDADES (con política de redondeo por Mayor Resto)
    comerciales_u = []
    tot_v25_u = round(sum(d['v25_u_tot'] for d in by_com_u.values()), 2)
    tot_v25_sep_u = round(sum(d['v25_u_sep'] for d in by_com_u.values()), 2)
    tot_v25_oct_u = round(sum(d['v25_u_oct'] for d in by_com_u.values()), 2)
    tot_v26_sep_u = round(sum(d['v26_u_sep'] for d in by_com_u.values()), 2)
    tot_est26_u = round(sum(d['est26_u'] for d in by_com_u.values()), 2)
    tot_pend_u = round(tot_est26_u - tot_v26_sep_u, 2)
    tot_pct_vend_u = (tot_v26_sep_u / tot_est26_u * 100.0) if tot_est26_u > 0 else 0.0
    tot_crec_oct_u = ((tot_pend_u / tot_v25_oct_u) - 1.0) * 100.0 if tot_v25_oct_u > 0 else 0.0

    for k_f, tgt_v in [
        ("ventas_2025_tot", tot_v25_u),
        ("ventas_2025_sep", tot_v25_sep_u),
        ("ventas_2025_oct", tot_v25_oct_u),
        ("ventas_2026_sep", tot_v26_sep_u),
        ("estimacion_2026", tot_est26_u),
    ]:
        round_conserving_sum(clientes_u, k_f, tgt_v, 2)
    for c_item in clientes_u:
        c_item["pendiente_oct_dic"] = round(c_item["estimacion_2026"] - c_item["ventas_2026_sep"], 2)

    for com in sorted(by_com_u.keys()):
        c_rows = [c for c in clientes_u if c["comercial"] == com]
        v25_t = round(sum(c["ventas_2025_tot"] for c in c_rows), 2)
        v25_s = round(sum(c["ventas_2025_sep"] for c in c_rows), 2)
        v25_o = round(sum(c["ventas_2025_oct"] for c in c_rows), 2)
        v26_s = round(sum(c["ventas_2026_sep"] for c in c_rows), 2)
        est_26 = round(sum(c["estimacion_2026"] for c in c_rows), 2)
        p_u = round(est_26 - v26_s, 2)
        pct_u = (v26_s / est_26 * 100.0) if est_26 > 0 else None
        c_sep_u = ((v26_s / v25_s) - 1.0) * 100.0 if v25_s > 0 else None
        c_oct_u = ((p_u / v25_o) - 1.0) * 100.0 if v25_o > 0 else None
        comerciales_u.append({
            "comercial": com,
            "num_clientes": len(c_rows),
            "ventas_2025_tot": v25_t,
            "ventas_2025_sep": v25_s,
            "ventas_2025_oct": v25_o,
            "ventas_2026_sep": v26_s,
            "estimacion_2026": est_26,
            "pendiente_oct_dic": p_u,
            "pct_vendido": round(pct_u, 1) if pct_u is not None else None,
            "crec_real_sep_pct": round(c_sep_u, 1) if c_sep_u is not None else None,
            "crec_nec_oct_pct": round(c_oct_u, 1) if c_oct_u is not None else None
        })

    data_unidades = {
        "metadata": {
            "titulo": "Cierre 2026 & Termómetro Q4 (Unidades)",
            "unidad_medida": "u",
            "unidad_simbolo": "u",
            "fecha_corte": "30/09/2026",
            "totales": {
                "num_clientes": len(clientes_u),
                "ventas_2025_tot": tot_v25_u,
                "ventas_2025_sep": tot_v25_sep_u,
                "ventas_2025_oct": tot_v25_oct_u,
                "ventas_2026_sep": tot_v26_sep_u,
                "estimacion_2026": tot_est26_u,
                "pendiente_oct_dic": tot_pend_u,
                "pct_vendido": round(tot_pct_vend_u, 1),
                "crec_nec_oct_pct": round(tot_crec_oct_u, 1)
            },
            "conteos_valoracion": conteos_val_u,
            "sas_info": {
                "cliente": "SUSTAINABLE AGRO SOLUTIONS SAU",
                "codigo": "430000427",
                "unidades_2026": round(sas_u_2026, 2),
                "descripcion": "Servicios de formulación y envasado de maquila SAS en unidades"
            }
        },
        "comerciales": comerciales_u,
        "clientes": clientes_u
    }

    # 5. Totales y Comerciales para EUROS (con política de redondeo por Mayor Resto)
    comerciales_eur = []
    tot_v25_eur = round(sum(d['v25_eur_tot'] for d in by_com_eur.values()), 2)
    tot_v25_sep_eur = round(sum(d['v25_eur_sep'] for d in by_com_eur.values()), 2)
    tot_v25_oct_eur = round(sum(d['v25_eur_oct'] for d in by_com_eur.values()), 2)
    tot_v26_sep_eur = round(sum(d['v26_eur_sep'] for d in by_com_eur.values()), 2)
    tot_est26_eur = round(sum(d['est26_eur'] for d in by_com_eur.values()), 2)
    tot_pend_eur = round(tot_est26_eur - tot_v26_sep_eur, 2)
    tot_pct_vend_eur = (tot_v26_sep_eur / tot_est26_eur * 100.0) if tot_est26_eur > 0 else 0.0
    tot_crec_oct_eur = ((tot_pend_eur / tot_v25_oct_eur) - 1.0) * 100.0 if tot_v25_oct_eur > 0 else 0.0

    for k_f, tgt_v in [
        ("ventas_2025_tot", tot_v25_eur),
        ("ventas_2025_sep", tot_v25_sep_eur),
        ("ventas_2025_oct", tot_v25_oct_eur),
        ("ventas_2026_sep", tot_v26_sep_eur),
        ("estimacion_2026", tot_est26_eur),
    ]:
        round_conserving_sum(clientes_eur, k_f, tgt_v, 2)
    for c_item in clientes_eur:
        c_item["pendiente_oct_dic"] = round(c_item["estimacion_2026"] - c_item["ventas_2026_sep"], 2)

    for com in sorted(by_com_eur.keys()):
        c_rows = [c for c in clientes_eur if c["comercial"] == com]
        v25_t = round(sum(c["ventas_2025_tot"] for c in c_rows), 2)
        v25_s = round(sum(c["ventas_2025_sep"] for c in c_rows), 2)
        v25_o = round(sum(c["ventas_2025_oct"] for c in c_rows), 2)
        v26_s = round(sum(c["ventas_2026_sep"] for c in c_rows), 2)
        est_26 = round(sum(c["estimacion_2026"] for c in c_rows), 2)
        p_eur = round(est_26 - v26_s, 2)
        pct_eur = (v26_s / est_26 * 100.0) if est_26 > 0 else None
        c_sep_eur = ((v26_s / v25_s) - 1.0) * 100.0 if v25_s > 0 else None
        c_oct_eur = ((p_eur / v25_o) - 1.0) * 100.0 if v25_o > 0 else None
        comerciales_eur.append({
            "comercial": com,
            "num_clientes": len(c_rows),
            "ventas_2025_tot": v25_t,
            "ventas_2025_sep": v25_s,
            "ventas_2025_oct": v25_o,
            "ventas_2026_sep": v26_s,
            "estimacion_2026": est_26,
            "pendiente_oct_dic": p_eur,
            "pct_vendido": round(pct_eur, 1) if pct_eur is not None else None,
            "crec_real_sep_pct": round(c_sep_eur, 1) if c_sep_eur is not None else None,
            "crec_nec_oct_pct": round(c_oct_eur, 1) if c_oct_eur is not None else None
        })

    data_euros = {
        "metadata": {
            "titulo": "Cierre 2026 & Termómetro Q4 (Euros)",
            "unidad_medida": "€",
            "unidad_simbolo": "€",
            "fecha_corte": "30/09/2026",
            "totales": {
                "num_clientes": len(clientes_eur),
                "ventas_2025_tot": round(tot_v25_eur, 2),
                "ventas_2025_sep": round(tot_v25_sep_eur, 2),
                "ventas_2025_oct": round(tot_v25_oct_eur, 2),
                "ventas_2026_sep": round(tot_v26_sep_eur, 2),
                "estimacion_2026": round(tot_est26_eur, 2),
                "pendiente_oct_dic": round(tot_pend_eur, 2),
                "pct_vendido": round(tot_pct_vend_eur, 1),
                "crec_nec_oct_pct": round(tot_crec_oct_eur, 1)
            },
            "conteos_valoracion": conteos_val_eur,
            "sas_info": {
                "cliente": "SUSTAINABLE AGRO SOLUTIONS SAU",
                "codigo": "430000427",
                "euros_2026": round(sas_eur_2026, 2),
                "descripcion": "Servicios de formulación y envasado de maquila SAS en euros"
            }
        },
        "comerciales": comerciales_eur,
        "clientes": clientes_eur
    }

    # 6. Salida combinada
    output_combined = {
        "active_mode": "unidades",
        "unidades": data_unidades,
        "euros": data_euros,
        # Fallbacks directos para retrocompatibilidad
        "metadata": data_unidades["metadata"],
        "comerciales": data_unidades["comerciales"],
        "clientes": data_unidades["clientes"]
    }

    # Guardar en web_dashboard_direccion y en web_dashboard
    target_dirs = ["web_dashboard_direccion", "web_dashboard"]
    for out_dir in target_dirs:
        os.makedirs(out_dir, exist_ok=True)
        json_path = os.path.join(out_dir, "marta_tracking_data.json")
        js_path = os.path.join(out_dir, "marta_tracking_data.js")

        with open(json_path, 'w', encoding='utf-8') as f:
            json.dump(output_combined, f, ensure_ascii=False, indent=2)

        with open(js_path, 'w', encoding='utf-8') as f:
            f.write("// Datos generados en UNIDADES y EUROS para Cierre 2026 & Termómetro Q4\n")
            f.write("window.MARTA_TRACKING_DATA = ")
            json.dump(output_combined, f, ensure_ascii=False, indent=2)
            f.write(";\n")

    print("\nGeneración completada con éxito en UNIDADES y EUROS:")
    print(f"  * UNIDADES: Est {tot_est26_u:,.0f} u | Real Sep {tot_v26_sep_u:,.0f} u ({tot_pct_vend_u:.1f}%) | SAS: {sas_u_2026:,.0f} u")
    print(f"  * EUROS:    Est {tot_est26_eur:,.2f} € | Real Sep {tot_v26_sep_eur:,.2f} € ({tot_pct_vend_eur:.1f}%) | SAS: {sas_eur_2026:,.2f} €")

if __name__ == '__main__':
    build_data()
