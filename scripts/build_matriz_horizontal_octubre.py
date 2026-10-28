import os
import shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_horizontal_matrix_octubre():
    print("=== Generando Prevision Matriz Horizontal con datos de Octubre ===")

    # 1. Copia segura de la matriz original
    src_file = os.path.join("antiguos", "Presupuesto_Ventas_2027_Definitivo_BACKUP_PRE_17180K.xlsx")
    if not os.path.exists(src_file):
        src_file = os.path.join("antiguos", "Presupuesto_Ventas_2027_Original_Sin_Aplanar.xlsx")
    
    temp_src = os.path.join("scratch", "temp_matriz_h_src.xlsx")
    os.makedirs("scratch", exist_ok=True)
    shutil.copy2(src_file, temp_src)

    wb_src = openpyxl.load_workbook(temp_src, data_only=True)
    ws_src = [s for s in wb_src if 'Horizontal' in s.title][0]

    TARGET_EUR = 910333.55
    TARGET_U = 311317.0

    # Leer todas las filas 2 a 1266 de la matriz
    rows_data = []
    tot_orig_eur = 0.0
    tot_orig_u = 0.0

    for r in range(2, 1267):
        com = ws_src.cell(r, 1).value
        pais = ws_src.cell(r, 2).value
        cl = ws_src.cell(r, 3).value
        clasif = ws_src.cell(r, 4).value
        sku = ws_src.cell(r, 5).value
        desc = ws_src.cell(r, 6).value
        env = ws_src.cell(r, 7).value
        pm26 = ws_src.cell(r, 8).value
        pm27 = ws_src.cell(r, 9).value

        u = float(ws_src.cell(r, 28).value or 0)
        eur = float(ws_src.cell(r, 29).value or 0)

        tot_orig_u += u
        tot_orig_eur += eur

        rows_data.append({
            'row_idx': r,
            'comercial': com,
            'pais': pais,
            'cliente': cl,
            'clasificacion': clasif,
            'sku': sku,
            'descripcion': desc,
            'envase': env,
            'pm26': pm26,
            'pm27': pm27,
            'unidades': u,
            'eur_orig': eur
        })

    wb_src.close()

    factor = TARGET_EUR / tot_orig_eur
    print(f"Total original matriz: {tot_orig_u:,.0f} u | {tot_orig_eur:,.2f} €")
    print(f"Factor ajuste P&L: {factor:.10f}")

    for item in rows_data:
        item['eur_adj'] = round(item['eur_orig'] * factor, 2)

    diff_cents = int(round((TARGET_EUR - sum(item['eur_adj'] for item in rows_data)) * 100))
    print(f"Ajuste de céntimos para cuadre exacto: {diff_cents}")

    # Distribuir céntimos de redondeo en las líneas con mayor importe
    sorted_indices = sorted(range(len(rows_data)), key=lambda i: rows_data[i]['eur_orig'], reverse=True)
    if diff_cents > 0:
        for i in range(diff_cents):
            rows_data[sorted_indices[i]]['eur_adj'] = round(rows_data[sorted_indices[i]]['eur_adj'] + 0.01, 2)
    elif diff_cents < 0:
        for i in range(-diff_cents):
            rows_data[sorted_indices[i]]['eur_adj'] = round(rows_data[sorted_indices[i]]['eur_adj'] - 0.01, 2)

    final_sum_eur = sum(item['eur_adj'] for item in rows_data)
    final_sum_u = sum(item['unidades'] for item in rows_data)
    print(f"Comprobación Suma Ajustada: {final_sum_u:,.0f} u | {final_sum_eur:,.2f} €")

    # -------------------------------------------------------------
    # CONSTRUIR LIBRO EXCEL (ÚNICA HOJA: Previsión Matriz Horizontal)
    # -------------------------------------------------------------
    out_file = "Desglose_Ventas_Octubre_2026_PL_910k.xlsx"
    wb_out = openpyxl.Workbook()
    ws = wb_out.active
    ws.title = "Previsión Matriz Horizontal"
    ws.views.sheetView[0].showGridLines = True

    # Estilos exactos de la matriz original
    NAVY_BLUE = "1F497D"       # Azul corporativo exacto de cabecera
    TOTAL_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    HEADER_FILL = PatternFill(start_color=NAVY_BLUE, end_color=NAVY_BLUE, fill_type="solid")

    font_header = Font(name="Segoe UI", size=8.5, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=8.5, color="000000")
    font_total = Font(name="Segoe UI", size=8.5, bold=True, color="000000")

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Side(border_style="thin", color="D9D9D9")
    cell_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
    double_bottom = Border(bottom=Side(border_style="double", color="1F497D"),
                           top=Side(border_style="thin", color="1F497D"),
                           left=thin_border, right=thin_border)

    fmt_currency = '#,##0.00 €;[Red]-#,##0.00 €;"0,00 €"'
    fmt_units = '#,##0;[Red]-#,##0;"0"'
    fmt_envase = '#,##0'

    # Fila 1: Cabeceras exactas
    headers = [
        ("Comercial", align_left, 14),
        ("País", align_left, 14),
        ("Cliente", align_left, 32),
        ("Clasificación Cliente", align_center, 22),
        ("Código Artículo", align_center, 16),
        ("Descripción Artículo", align_left, 35),
        ("Tamaño Envase", align_center, 14),
        ("Precio Medio 2026 (€/u)", align_right, 15),
        ("Precio Medio 2027 (€/u)", align_right, 15),
        ("Oct-26 (u)", align_right, 14),
        ("Oct-26 (€)", align_right, 16)
    ]

    ws.row_dimensions[1].height = 28

    for col_i, (h_name, h_align, width) in enumerate(headers, 1):
        c = ws.cell(1, col_i, h_name)
        c.font = font_header
        c.fill = HEADER_FILL
        c.alignment = align_header
        c.border = cell_border
        col_letter = get_column_letter(col_i)
        ws.column_dimensions[col_letter].width = width

    # Filas 2 a 1266: Los 1.265 registros en su orden exacto original
    for i, row in enumerate(rows_data):
        r_idx = 2 + i
        ws.row_dimensions[r_idx].height = 18

        ws.cell(r_idx, 1, row['comercial']).alignment = align_left
        ws.cell(r_idx, 2, row['pais']).alignment = align_left
        ws.cell(r_idx, 3, row['cliente']).alignment = align_left
        ws.cell(r_idx, 4, row['clasificacion']).alignment = align_center
        ws.cell(r_idx, 5, row['sku']).alignment = align_center
        ws.cell(r_idx, 6, row['descripcion']).alignment = align_left

        c_env = ws.cell(r_idx, 7, row['envase'])
        c_env.alignment = align_center
        c_env.number_format = fmt_envase

        c_pm26 = ws.cell(r_idx, 8, row['pm26'])
        c_pm26.alignment = align_right
        c_pm26.number_format = fmt_currency

        c_pm27 = ws.cell(r_idx, 9, row['pm27'])
        c_pm27.alignment = align_right
        c_pm27.number_format = fmt_currency

        c_u = ws.cell(r_idx, 10, row['unidades'])
        c_u.alignment = align_right
        c_u.number_format = fmt_units

        c_eur = ws.cell(r_idx, 11, row['eur_adj'])
        c_eur.alignment = align_right
        c_eur.number_format = fmt_currency

        for col_i in range(1, 12):
            cell = ws.cell(r_idx, col_i)
            cell.font = font_data
            cell.border = cell_border

    # Fila 1267: TOTAL GENERAL con fórmulas =SUM(...)
    tot_row = 1267
    ws.row_dimensions[tot_row].height = 22

    ws.cell(tot_row, 1, "TOTAL GENERAL").alignment = align_left
    for col_i in range(2, 10):
        ws.cell(tot_row, col_i, "").alignment = align_center

    t_u = ws.cell(tot_row, 10, f"=SUM(J2:J1266)")
    t_u.alignment = align_right
    t_u.number_format = fmt_units

    t_eur = ws.cell(tot_row, 11, f"=SUM(K2:K1266)")
    t_eur.alignment = align_right
    t_eur.number_format = fmt_currency

    for col_i in range(1, 12):
        cell = ws.cell(tot_row, col_i)
        cell.font = font_total
        cell.fill = TOTAL_FILL
        cell.border = double_bottom

    # Filtros automáticos en fila 1
    ws.auto_filter.ref = f"A1:K{tot_row}"

    # Inmovilizar paneles en D2 (Comercial, País y Cliente siempre visibles al desplazarse)
    ws.freeze_panes = "D2"

    wb_out.save(out_file)
    print(f"[OK] Generado exitosamente: {os.path.abspath(out_file)}")

if __name__ == '__main__':
    build_horizontal_matrix_octubre()
