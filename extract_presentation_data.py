import openpyxl, sys, collections, json

wb = openpyxl.load_workbook('Presupuesto_Ventas_2027_Definitivo.xlsx', data_only=True)
ws = wb['Previsión Matriz Horizontal']
header = [ws.cell(1, c).value for c in range(1, ws.max_column+1)]
col_26 = [i for i, h in enumerate(header, 1) if h and 'Total 2026 (€)' in str(h)][0]
col_27 = [i for i, h in enumerate(header, 1) if h and 'Total 2027 (€)' in str(h)][0]
col_u26 = [i for i, h in enumerate(header, 1) if h and 'Total 2026 (u)' in str(h)][0]
col_u27 = [i for i, h in enumerate(header, 1) if h and 'Total 2027 (u)' in str(h)][0]

data = collections.defaultdict(lambda: {'v26': 0.0, 'v27': 0.0, 'u26': 0.0, 'u27': 0.0})
for r in range(2, ws.max_row+1):
    c = ws.cell(r, 1).value
    if c and not 'TOTAL' in str(c).upper():
        com = str(c).strip()
        data[com]['v26'] += float(ws.cell(r, col_26).value or 0)
        data[com]['v27'] += float(ws.cell(r, col_27).value or 0)
        data[com]['u26'] += float(ws.cell(r, col_u26).value or 0)
        data[com]['u27'] += float(ws.cell(r, col_u27).value or 0)

print("--- VENTAS POR COMERCIAL ---")
res_com = []
for com, vals in sorted(data.items(), key=lambda x: -x[1]['v27']):
    pct = ((vals['v27']/vals['v26'])-1)*100 if vals['v26'] else 0
    pct_u = ((vals['u27']/vals['u26'])-1)*100 if vals['u26'] else 0
    print(f"{com:12s} | 2026: {vals['v26']:11,.0f} € ({vals['u26']:8,.0f} u) | 2027: {vals['v27']:11,.0f} € ({vals['u27']:8,.0f} u) | Var: {pct:+5.1f}% EUR, {pct_u:+5.1f}% UDS")
    res_com.append({
        'comercial': com,
        'v26': vals['v26'],
        'v27': vals['v27'],
        'u26': vals['u26'],
        'u27': vals['u27'],
        'var_pct_eur': pct,
        'var_pct_uds': pct_u
    })

# Check P&L summary from 2027-Budget-Codiagro P&L.xlsx
wb_pnl = openpyxl.load_workbook('Gastos/2027-Budget-Codiagro P&L.xlsx', data_only=True)
print("\n--- P&L 2027 SUMMARY ---")
ws_pnl = wb_pnl['2027']
pnl_summary = {}
for r in range(1, 35):
    concept = ws_pnl.cell(r, 2).value
    val = ws_pnl.cell(r, ws_pnl.max_column).value
    if concept:
        print(f"R{r:02d}: {str(concept):30s} = {val}")
        pnl_summary[str(concept).strip()] = val

# Also check 2026-Forecast-Codiagro P&L.xlsx
wb_fc = openpyxl.load_workbook('Gastos/2026-Forecast-Codiagro P&L.xlsx', data_only=True)
print("\n--- P&L 2026 FORECAST SUMMARY ---")
ws_fc = wb_fc['2026']
fc_summary = {}
for r in range(1, 35):
    concept = ws_fc.cell(r, 2).value
    val = ws_fc.cell(r, ws_fc.max_column).value
    if concept:
        print(f"R{r:02d}: {str(concept):30s} = {val}")
        fc_summary[str(concept).strip()] = val
