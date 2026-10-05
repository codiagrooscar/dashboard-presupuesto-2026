import openpyxl, os, sys
import unicodedata

def norm(s):
    if not isinstance(s, str): return ''
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn').upper().strip()

# Import all_modified_lines from generar_excel_informe_cambios_q4
from generar_excel_informe_cambios_q4 import all_modified_lines

print(f"Total líneas modificadas a aplicar: {len(all_modified_lines)}")

# Test matching against Presupuesto_Ventas_2027_Original_Sin_Aplanar.xlsx
wb = openpyxl.load_workbook('Presupuesto_Ventas_2027_Original_Sin_Aplanar.xlsx', data_only=True)
ws = wb['Previsión Matriz Horizontal'] if 'Previsión Matriz Horizontal' in wb.sheetnames else wb[wb.sheetnames[2]]

budget_rows = []
for r in range(2, ws.max_row+1):
    com = str(ws.cell(r, 1).value or '')
    pais = str(ws.cell(r, 2).value or '').strip().upper()
    cli = str(ws.cell(r, 3).value or '')
    sku = str(ws.cell(r, 5).value or '').strip().upper()
    
    sku_clean = sku
    if pais != 'ESPAÑA' and len(sku) > 2:
        sku_clean = sku[:-2]
        
    budget_rows.append({
        'row': r,
        'com': com,
        'com_norm': norm(com),
        'cli': cli,
        'cli_norm': norm(cli),
        'sku': sku,
        'sku_clean': sku_clean,
        'p_2026': ws.cell(r, 8).value,
        'oct_u': ws.cell(r, 28).value,
        'nov_u': ws.cell(r, 30).value,
        'dic_u': ws.cell(r, 32).value
    })

print(f"Filas indexadas en matriz presupuestaria: {len(budget_rows)}")

matched_count = 0
unmatched = []

for ch in all_modified_lines:
    ch_com = norm(ch['comercial'])
    ch_cli = norm(ch['cliente'])
    ch_sku = str(ch['sku']).strip().upper()
    
    found = []
    for b in budget_rows:
        if ch_cli == b['cli_norm'] or ch_cli in b['cli_norm'] or b['cli_norm'] in ch_cli:
            if ch_sku == b['sku'] or ch_sku == b['sku_clean']:
                found.append(b)
                
    if len(found) == 1:
        matched_count += 1
    elif len(found) > 1:
        f_com = [b for b in found if ch_com in b['com_norm'] or b['com_norm'] in ch_com]
        if len(f_com) == 1:
            matched_count += 1
        else:
            print("Ambigüedad:", ch['comercial'], ch['cliente'], ch_sku, len(found))
    else:
        unmatched.append(ch)

print(f"Resultado del emparejamiento: {matched_count} de {len(all_modified_lines)}")
if unmatched:
    print("No emparejados:")
    for u in unmatched:
        print(" ", u['comercial'], "|", u['cliente'], "|", u['sku'])
