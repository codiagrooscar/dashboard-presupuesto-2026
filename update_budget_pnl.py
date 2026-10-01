import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def clean_v(v):
    if v is None: return 0.0
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).strip().replace(' ', '').replace('-', '0'))
    except: return 0.0

def update_pnl():
    file_path = r'Gastos/P&L mensual ene-ago V3.xlsx'
    wb = openpyxl.load_workbook(file_path, data_only=False)
    ws = wb['PL Mensual 2026']

    num_fmt = '#,##0.00;[Red]-#,##0.00;"-"'

    # Update Column X Header
    ws.cell(3, 24).value = 'Criterio Previsión (Sep-Dic 2026 - Modificaciones)'

    # Track modifications for reporting
    modifications = []

    # Map of target account codes
    # 1. 601000000: Compras de materias primas (quitar 4.5% inflacion)
    # 2. 602000001: COMPRA DE ENVASES Y EMBALAJES (quitar 4.5% inflacion)
    # 3. 602000002: COMPRA DE ETIQUETAS (quitar 4.5% inflacion)
    # 4. 602000003: COMPRA DE PALETS (quitar 4.5% inflacion)
    # 5. 607000000: Trabajos por otras empresas (sep-nov = 0, dic = -7970)
    # 6. 621000014: ALQUILER MURANO (sep-dic = 0)
    # 7. (622): Mantenimiento (quitar * 1.045, dejar avg_active)
    # 8. 623000002: SERVICIOS CONTABLES (a la mitad: -1400/mes)
    # 9. 623100002: GASTOS MAURICIO ALISTE (mantener, no subir en 2027)
    # 10. 623100006: GASTOS Y SERV. AGRICOLA DEL JUCAR (mantener, no subir en 2027)
    # 11. 623100007: GASTOS Y SUELDO MEHMET (mantener, no subir en 2027)
    # 12. 623100008: COMISIONES KAMAL KEHAL (sep-dic = 0)
    # 13. 624000001: TRANSPORTES NACIONAL (dejar por ahora)
    # 14. 627000008: MARKETING (sep-dic = 0)
    # 15. 631000004: IBIB (sep-dic = 0)

    for r in range(4, 206):
        cta = str(ws.cell(r, 9).value or '')
        n5 = str(ws.cell(r, 8).value or '')
        code = cta.split(' - ')[0].split(' ')[0]

        # Historical Jan-Aug
        jan_aug = [clean_v(ws.cell(r, c).value) for c in range(11, 19)]
        act = [v for v in jan_aug if abs(v) > 0.01]
        avg_act = (sum(act) / len(act)) if act else 0.0

        old_sep = clean_v(ws.cell(r, 19).value)
        old_oct = clean_v(ws.cell(r, 20).value)
        old_nov = clean_v(ws.cell(r, 21).value)
        old_dic = clean_v(ws.cell(r, 22).value)
        old_sd = old_sep + old_oct + old_nov + old_dic
        old_crit = str(ws.cell(r, 24).value or '')

        changed = False
        new_sep, new_oct, new_nov, new_dic = old_sep, old_oct, old_nov, old_dic
        new_crit = old_crit

        if code == '601000000':
            # Quitar 4.5% inflacion de 2026
            new_sep = round(old_sep / 1.045, 2)
            new_oct = round(old_oct / 1.045, 2)
            new_nov = round(old_nov / 1.045, 2)
            new_dic = round(old_dic / 1.045, 2)
            new_crit = 'Coste variable directo: volumen según presupuesto ajustado (ratio histórico 25,85% s/ventas sin incremento de inflación en 2026; la inflación del +4,5% se aplica solo en 2027).'
            changed = True
        elif code == '602000001':
            new_sep = round(old_sep / 1.045, 2)
            new_oct = round(old_oct / 1.045, 2)
            new_nov = round(old_nov / 1.045, 2)
            new_dic = round(old_dic / 1.045, 2)
            new_crit = 'Coste variable por envases: ratio histórico s/ventas ajustadas sin incremento de inflación en 2026; la inflación del +4,5% se reserva exclusivamente para 2027.'
            changed = True
        elif code == '602000002':
            new_sep = round(old_sep / 1.045, 2)
            new_oct = round(old_oct / 1.045, 2)
            new_nov = round(old_nov / 1.045, 2)
            new_dic = round(old_dic / 1.045, 2)
            new_crit = 'Coste variable de etiquetas: ratio s/ventas ajustadas sin ajuste de precios de imprenta por inflación en 2026 (solo en 2027).'
            changed = True
        elif code == '602000003':
            new_sep = round(old_sep / 1.045, 2)
            new_oct = round(old_oct / 1.045, 2)
            new_nov = round(old_nov / 1.045, 2)
            new_dic = round(old_dic / 1.045, 2)
            new_crit = 'Coste variable de palets: ratio s/ventas ajustadas sin inflación de suministros logísticos en 2026 (aplicable en 2027).'
            changed = True
        elif code == '607000000':
            new_sep = 0.0
            new_oct = 0.0
            new_nov = 0.0
            new_dic = -7970.00
            new_crit = 'Coste variable de subcontratación fabril: eliminadas previsiones de sep a nov, dejando únicamente diciembre (-7.970,00 €) por refuerzo a cierre de ejercicio según indicación.'
            changed = True
        elif code == '621000014':
            new_sep = 0.0
            new_oct = 0.0
            new_nov = 0.0
            new_dic = 0.0
            new_crit = 'Arrendamiento Murano: no hay más cuotas en 2026 (ya está todo pagado, eliminadas sep-dic). En 2027 poner la misma cantidad anual ajustada a la inflación (+4,5%).'
            changed = True
        elif '(622)' in n5:
            # Quitar inflacion 4.5% de mantenimiento
            val = round(avg_act, 2)
            new_sep = val
            new_oct = val
            new_nov = val
            new_dic = val
            new_crit = 'Coste semi-variable: mantenimiento preventivo y correctivo de planta según promedio mensual histórico (sin incremento de inflación en 2026; la inflación del +4,5% se aplica solo en 2027).'
            changed = True
        elif code == '623000002':
            # Bajar a la mitad: -1400/mes
            new_sep = -1400.00
            new_oct = -1400.00
            new_nov = -1400.00
            new_dic = -1400.00
            new_crit = 'Servicios contables y financieros externos: importe mensual reducido a la mitad (-1.400,00 €/mes de sep a dic) según indicación de controlling.'
            changed = True
        elif code == '623100002':
            new_crit = 'Honorarios y gastos comerciales de zona según plan presupuestado + 5.000,00 € adicionales en diciembre según instrucción comercial. Para 2027 no hay que subirles nada (mismo importe).'
            changed = True
        elif code == '623100006':
            new_crit = 'Dotación mensual de 500,00 € de sep a dic + 30.000,00 € extraordinarios en diciembre según instrucción comercial. Para 2027 no hay que subirles nada (mismo importe).'
            changed = True
        elif code == '623100007':
            new_crit = 'Dotación mensual de 6.000,00 € de sep a dic + 100.000,00 € extraordinarios en diciembre según instrucción comercial. Para 2027 no hay que subirles nada (mismo importe).'
            changed = True
        elif code == '623100008':
            new_sep = 0.0
            new_oct = 0.0
            new_nov = 0.0
            new_dic = 0.0
            new_crit = 'Comisiones Kamal Kehal: eliminados los importes de sep a dic del 2026. Para 2027 poner la misma cantidad anual real del 2026 (-3.398,55 €).'
            changed = True
        elif code == '624000001':
            new_crit = 'Coste variable de transporte nacional: mantenida temporalmente la previsión según volumen de expediciones presupuestadas; pendiente de recibir excel con previsiones logísticas.'
            changed = True
        elif code == '627000008':
            new_sep = 0.0
            new_oct = 0.0
            new_nov = 0.0
            new_dic = 0.0
            new_crit = 'Gastos de marketing: eliminadas las partidas de sep a dic del 2026. Para 2027 pendiente de adjuntar el plan de marketing.'
            changed = True
        elif code == '631000004':
            new_sep = 0.0
            new_oct = 0.0
            new_nov = 0.0
            new_dic = 0.0
            new_crit = 'Tasas e impuestos municipales (IBI): eliminadas las partidas de sep a dic del 2026 por estar ya pagado en el ejercicio. En el 2027 poner la misma cantidad que ya se ha pagado en el 2026 (-7.396,59 €).'
            changed = True

        if changed:
            ws.cell(r, 19).value = new_sep
            ws.cell(r, 19).number_format = num_fmt

            ws.cell(r, 20).value = new_oct
            ws.cell(r, 20).number_format = num_fmt

            ws.cell(r, 21).value = new_nov
            ws.cell(r, 21).number_format = num_fmt

            ws.cell(r, 22).value = new_dic
            ws.cell(r, 22).number_format = num_fmt

            ws.cell(r, 23).value = f'=SUM(K{r}:V{r})'
            ws.cell(r, 23).number_format = num_fmt

            ws.cell(r, 24).value = new_crit

            new_sd = new_sep + new_oct + new_nov + new_dic
            diff = new_sd - old_sd
            modifications.append({
                'row': r,
                'code': code,
                'cta': cta,
                'old_sd': old_sd,
                'new_sd': new_sd,
                'diff': diff,
                'sep': new_sep,
                'oct': new_oct,
                'nov': new_nov,
                'dic': new_dic,
                'criterio': new_crit
            })

    # Save to V3 and also create V4
    out_v3 = r'Gastos/P&L mensual ene-ago V3.xlsx'
    out_v4 = r'Gastos/P&L mensual ene-ago V4.xlsx'
    wb.save(out_v3)
    wb.save(out_v4)
    print(f'Saved updated workbook to {out_v3} and {out_v4}')
    print(f'Total modified accounts: {len(modifications)}')

    total_diff = sum(m['diff'] for m in modifications)
    print(f'Total impact on Sep-Dic (positive = cost reduction): {total_diff:12,.2f} €')

    return modifications

if __name__ == '__main__':
    update_pnl()
