"""
Generador Maestro de Revisiones de Previsiones Q4 2026 (Corte 27/10/2026)
Genera los archivos Excel individuales para todos los comerciales:
  - Alfonso, García, Irene, Javier, Mehmet, Pedro, Ricardo
Y el libro CONSOLIDADO con hoja de Resumen de Dirección y pestañas por comercial.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from collections import defaultdict
import os
import shutil
import unicodedata

# --- Configuración y Paletas de Estilo ---
FONT_FAMILY = "Segoe UI"

font_title = Font(name=FONT_FAMILY, size=15, bold=True, color="FFFFFF")
font_subtitle = Font(name=FONT_FAMILY, size=10, italic=True, color="E2E8F0")
font_instruct = Font(name=FONT_FAMILY, size=9, bold=True, color="1E293B")

font_kpi_title = Font(name=FONT_FAMILY, size=9, bold=True)
font_kpi_val = Font(name=FONT_FAMILY, size=15, bold=True)

font_group_hdr = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
font_col_hdr = Font(name=FONT_FAMILY, size=9, bold=True, color="FFFFFF")

font_regular = Font(name=FONT_FAMILY, size=9, color="000000")
font_sku = Font(name="Consolas", size=9, bold=True, color="1E293B")
font_editable = Font(name=FONT_FAMILY, size=9, bold=True, color="064E3B")
font_total = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")

fill_title = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
fill_instruct = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")

fill_grp_info = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
fill_grp_oct = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
fill_grp_q4 = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
fill_grp_tot = PatternFill(start_color="475569", end_color="475569", fill_type="solid")

fill_col_hdr = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
fill_col_oct = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
fill_col_q4 = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
fill_col_tot = PatternFill(start_color="64748B", end_color="64748B", fill_type="solid")

fill_zebra_0 = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
fill_zebra_1 = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
fill_prev_hist = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
fill_editable = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
fill_comment_oct = PatternFill(start_color="FEFCE8", end_color="FEFCE8", fill_type="solid")
fill_extra_row = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")
fill_total = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

border_thin_val = Side(style='thin', color='CBD5E1')
border_thin = Border(left=border_thin_val, right=border_thin_val, top=border_thin_val, bottom=border_thin_val)
border_editable_val = Side(style='thin', color='10B981')
border_editable = Border(left=border_editable_val, right=border_editable_val, top=border_editable_val, bottom=border_editable_val)
border_tot_top = Side(style='medium', color='38BDF8')
border_tot_bottom = Side(style='double', color='38BDF8')
border_total = Border(left=border_thin_val, right=border_thin_val, top=border_tot_top, bottom=border_tot_bottom)

align_center = Alignment(horizontal='center', vertical='center')
align_left = Alignment(horizontal='left', vertical='center')
align_right = Alignment(horizontal='right', vertical='center')
align_col_hdr = Alignment(horizontal='center', vertical='center', wrap_text=True)

# Comerciales configurados
COMERCIALES = [
    {
        'key': 'Alfonso',
        'display_name': 'ALFONSO',
        'sheet_name': 'Alfonso',
        'individual_sheet': 'Previsiones_Alfonso',
        'src_file': 'antiguos/Revision_Previsiones_Q4_2026/Copia de Revision_Previsiones_Q4_2026_Alfonso.xlsx',
        'out_file': 'Revision_Previsiones_Q4_2026_Alfonso.xlsx'
    },
    {
        'key': 'Garcia',
        'display_name': 'GARCÍA',
        'sheet_name': 'García',
        'individual_sheet': 'Previsiones_García',
        'src_file': 'antiguos/Revision_Previsiones_Q4_2026/Copia de Revision_Previsiones_Q4_2026_Garcia.xlsx',
        'out_file': 'Revision_Previsiones_Q4_2026_Garcia.xlsx'
    },
    {
        'key': 'Irene',
        'display_name': 'IRENE',
        'sheet_name': 'Irene',
        'individual_sheet': 'Previsiones_Irene',
        'src_file': 'antiguos/Revision_Previsiones_Q4_2026/Copia de Revision_Previsiones_Q4_2026_Irene.xlsx',
        'out_file': 'Revision_Previsiones_Q4_2026_Irene.xlsx'
    },
    {
        'key': 'Javier',
        'display_name': 'JAVIER',
        'sheet_name': 'Javier',
        'individual_sheet': 'Previsiones_Javier',
        'src_file': 'antiguos/Revision_Previsiones_Q4_2026/Copia de Revision_Previsiones_Q4_2026_Javier.xlsx',
        'out_file': 'Revision_Previsiones_Q4_2026_Javier.xlsx'
    },
    {
        'key': 'Mehmet',
        'display_name': 'MEHMET',
        'sheet_name': 'Mehmet',
        'individual_sheet': 'Previsiones_Mehmet',
        'src_file': 'antiguos/Revision_Previsiones_Q4_2026/Revision_Previsiones_Q4_2026_Mehmet.xlsx',
        'out_file': 'Revision_Previsiones_Q4_2026_Mehmet.xlsx'
    },
    {
        'key': 'Pedro',
        'display_name': 'PEDRO',
        'sheet_name': 'Pedro',
        'individual_sheet': 'Previsiones_Pedro',
        'src_file': 'antiguos/Revision_Previsiones_Q4_2026/Copia de Revision_Previsiones_Q4_2026_Pedro.xlsx',
        'out_file': 'Revision_Previsiones_Q4_2026_Pedro.xlsx'
    },
    {
        'key': 'Ricardo',
        'display_name': 'RICARDO',
        'sheet_name': 'Ricardo',
        'individual_sheet': 'Previsiones_Ricardo',
        'src_file': 'antiguos/Revision_Previsiones_Q4_2026/Copia de Revision_Previsiones_Q4_2026_Ricardo.xlsx',
        'out_file': 'Revision_Previsiones_Q4_2026_Ricardo.xlsx'
    }
]

def norm(s):
    if not isinstance(s, str): return ''
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn').upper().strip()

def load_client_commercial_map():
    """Carga la asignación oficial de clientes a comerciales desde el presupuesto definitivo."""
    client_map = {}
    if os.path.exists('Presupuesto_Ventas_2027_Definitivo.xlsx'):
        wb_b = openpyxl.load_workbook('Presupuesto_Ventas_2027_Definitivo.xlsx', data_only=True)
        if 'Previsión Matriz Horizontal' in wb_b.sheetnames:
            ws_b = wb_b['Previsión Matriz Horizontal']
            com_col, cli_col, pais_col = 3, 6, 2
            for c in range(1, 15):
                val = str(ws_b.cell(1, c).value or '').strip()
                if val == 'Comercial': com_col = c
                elif val == 'Cliente': cli_col = c
                elif val == 'País': pais_col = c
            for r in range(2, ws_b.max_row + 1):
                com = str(ws_b.cell(r, com_col).value or '').strip()
                cli = str(ws_b.cell(r, cli_col).value or '').strip()
                pais = str(ws_b.cell(r, pais_col).value or '').strip()
                if com and cli and 'TOTAL' not in com.upper():
                    client_map[norm(cli)] = {
                        'comercial': com,
                        'pais': pais or 'ESPAÑA',
                        'cliente_orig': cli
                    }
    # Asignaciones fijas validadas
    client_map[norm('FINCA DOÑA ANA C.B.')] = {'comercial': 'Javier', 'pais': 'ESPAÑA', 'cliente_orig': 'FINCA DOÑA ANA C.B.'}
    client_map[norm('FITOSANITARIOS CARCAIXENT, S.L.')] = {'comercial': 'Javier', 'pais': 'ESPAÑA', 'cliente_orig': 'FITOSANITARIOS CARCAIXENT, S.L.'}
    client_map[norm('ALMENDRALIA IBÉRICA, S.L.U.')] = {'comercial': 'Javier', 'pais': 'ESPAÑA', 'cliente_orig': 'ALMENDRALIA IBÉRICA, S.L.U.'}
    client_map[norm('TÉCNICAS AGRÍCOLAS, S.A.')] = {'comercial': 'Ricardo', 'pais': 'CHILE', 'cliente_orig': 'TÉCNICAS AGRÍCOLAS, S.A.'}
    return client_map

def load_real_orders():
    """Carga los pedidos reales al 27/10 agrupados por (Cliente, SKU)."""
    wb_ped = openpyxl.load_workbook('Pedidos27.10 comerciales.xlsx', data_only=True)
    ws_ped = wb_ped.active
    
    pedidos = defaultdict(lambda: {'qty': 0.0, 'desc': '', 'cli': '', 'sku': ''})
    for r in range(2, ws_ped.max_row + 1):
        c_name = ws_ped.cell(r, 6).value
        sku = ws_ped.cell(r, 7).value
        desc = ws_ped.cell(r, 8).value
        qty = ws_ped.cell(r, 9).value or 0
        if c_name and sku and float(qty) > 0:
            cli_clean = str(c_name).strip()
            sku_clean = str(sku).strip()
            key = (cli_clean, sku_clean)
            pedidos[key]['qty'] += float(qty)
            pedidos[key]['desc'] = str(desc).strip() if desc else ''
            pedidos[key]['cli'] = cli_clean
            pedidos[key]['sku'] = sku_clean
    return pedidos

def load_commercial_data(com_info, pedidos_dict, client_map):
    """Carga las líneas de previsión vigentes del comercial deduplicando filas y agregando pedidos reales y clientes extras."""
    tmp_path = f"temp_{com_info['key']}_src.xlsx"
    shutil.copy2(com_info['src_file'], tmp_path)
    wb_src = openpyxl.load_workbook(tmp_path, data_only=True)
    ws_src = wb_src.active
    
    # Encontrar fila total
    tot_r = None
    for r in range(ws_src.max_row, 5, -1):
        v = ws_src.cell(r, 2).value
        if v and 'TOTAL' in str(v).upper():
            tot_r = r
            break
            
    raw_items = {}
    client_countries = {}
    com_clients = set()
    com_norm = norm(com_info['key'])
    
    for r in range(10, tot_r):
        pais = str(ws_src.cell(r, 2).value or 'ESPAÑA').strip()
        cli = str(ws_src.cell(r, 3).value or '').strip()
        sku = str(ws_src.cell(r, 4).value or '').strip()
        desc = str(ws_src.cell(r, 5).value or '').strip()
        
        # Col 12: Oct Revisado anterior, Col 14: Nov Revisado, Col 16: Dic Revisado
        oct_prev = float(ws_src.cell(r, 12).value or 0)
        nov_prev = float(ws_src.cell(r, 14).value or 0)
        dic_prev = float(ws_src.cell(r, 16).value or 0)
        
        if cli and sku:
            key = (cli, sku)
            client_countries[cli] = pais
            com_clients.add(cli)
            
            if key not in raw_items:
                raw_items[key] = {
                    'pais': pais,
                    'cliente': cli,
                    'sku': sku,
                    'desc': desc,
                    'oct_prev_vigente': oct_prev,
                    'nov_prev_vigente': nov_prev,
                    'dic_prev_vigente': dic_prev,
                    'es_extra': False
                }
            else:
                # Fusión de fila duplicada de la plantilla para no duplicar pedidos
                raw_items[key]['oct_prev_vigente'] += oct_prev
                raw_items[key]['nov_prev_vigente'] += nov_prev
                raw_items[key]['dic_prev_vigente'] += dic_prev
                if not raw_items[key]['desc'] and desc:
                    raw_items[key]['desc'] = desc
                    
    # Agregar clientes asignados a este comercial desde client_map
    for c_norm, c_info in client_map.items():
        if norm(c_info['comercial']) == com_norm:
            com_clients.add(c_info['cliente_orig'])
            if c_info['cliente_orig'] not in client_countries:
                client_countries[c_info['cliente_orig']] = c_info['pais']

    # Asignar pedidos reales exactamente una vez por (Cliente, SKU)
    items = []
    for key, it in raw_items.items():
        cli, sku = key
        real_data = pedidos_dict.get((cli, sku), {'qty': 0.0})
        if real_data['qty'] == 0:
            for (p_cli, p_sku), p_data in pedidos_dict.items():
                if norm(p_cli) == norm(cli) and p_sku.upper() == sku.upper():
                    real_data = p_data
                    break
        it['oct_real'] = real_data['qty']
        items.append(it)
            
    # Agregar pedidos extras de clientes del comercial que no tenían SKU presupuestado
    for (cli, sku), data in pedidos_dict.items():
        is_client_of_com = (cli in com_clients) or any(norm(cli) == norm(c) for c in com_clients)
        if not is_client_of_com:
            c_info = client_map.get(norm(cli))
            if c_info and norm(c_info['comercial']) == com_norm:
                is_client_of_com = True
                if cli not in client_countries:
                    client_countries[cli] = c_info['pais']

        key_exists = (cli, sku) in raw_items or any(
            norm(k_cli) == norm(cli) and k_sku.upper() == sku.upper()
            for (k_cli, k_sku) in raw_items.keys()
        )
        
        if is_client_of_com and not key_exists:
            pais = client_countries.get(cli, 'ESPAÑA')
            items.append({
                'pais': pais,
                'cliente': cli,
                'sku': sku,
                'desc': data['desc'],
                'oct_prev_vigente': 0.0,
                'oct_real': data['qty'],
                'nov_prev_vigente': 0.0,
                'dic_prev_vigente': 0.0,
                'es_extra': True
            })
            
    # Ordenar por cliente y sku
    items.sort(key=lambda x: (x['cliente'], x['sku']))
    
    if os.path.exists(tmp_path):
        os.remove(tmp_path)
        
    return items

def build_commercial_sheet(ws, com_info, items):
    """Construye y formatea la hoja de trabajo para un comercial."""
    ws.views.sheetView[0].showGridLines = True
    
    # 1. Título y Subtítulo
    ws.merge_cells('B2:R2')
    ws['B2'] = "CODIAGRO · ACTUALIZACIÓN DE PREVISIONES COMERCIALES Q4 2026 (OCTUBRE - DICIEMBRE)"
    ws['B2'].font = font_title
    ws['B2'].fill = fill_title
    ws['B2'].alignment = align_center
    ws.row_dimensions[2].height = 30
    
    ws.merge_cells('B3:R3')
    ws['B3'] = f"DELEGADO COMERCIAL: {com_info['display_name']} | Real Octubre a Fecha de Hoy (27/10/2026) vs Previsiones a Final de Año | Generado: 27/10/2026"
    ws['B3'].font = font_subtitle
    ws['B3'].fill = fill_title
    ws['B3'].alignment = align_center
    ws.row_dimensions[3].height = 20
    
    # 2. Instrucciones
    ws.merge_cells('B4:R4')
    ws['B4'] = "📌 INSTRUCCIONES: 1) Revise el bloque de OCTUBRE (Previsión Vigente, Real a 27/10 y Desviación) e indique en col. J el comentario correspondiente. 2) En las columnas verdes (L y N), revise y actualice sus cantidades estimadas de NOVIEMBRE y DICIEMBRE. 3) Comente en col. R cualquier justificación relevante sobre la variación total del Q4."
    ws['B4'].font = font_instruct
    ws['B4'].fill = fill_instruct
    ws['B4'].alignment = align_left
    ws.row_dimensions[4].height = 24
    
    # Fila de Totales final
    tot_row = len(items) + 10
    
    # 3. Tarjetas KPI
    kpis = [
        ("B", "D", "PREVISIÓN OCTUBRE VIGENTE (U)", f"=F{tot_row}", "2563EB", "DBEAFE"),
        ("E", "G", "REAL OCTUBRE A 27/10 (U)", f"=G{tot_row}", "047857", "D1FAE5"),
        ("H", "I", "% CONSECUCIÓN A FECHA DE HOY", f"=I{tot_row}", "0284C7", "E0F2FE"),
        ("J", "L", "PREVISIÓN VIGENTE Q4 (U)", f"=O{tot_row}", "475569", "F1F5F9"),
        ("M", "O", "PREVISIÓN REVISADA Q4 (U)", f"=P{tot_row}", "047857", "ECFDF5"),
        ("P", "R", "VARIACIÓN NETA Q4 (U)", f"=Q{tot_row}", "D97706", "FEF3C7"),
    ]
    
    for c_start, c_end, title, formula, color_hex, bg_hex in kpis:
        # Fila 5: Título
        ws.merge_cells(f"{c_start}5:{c_end}5")
        cell_t = ws[f"{c_start}5"]
        cell_t.value = title
        cell_t.font = Font(name=FONT_FAMILY, size=9, bold=True, color="334155")
        cell_t.fill = PatternFill(start_color=bg_hex, end_color=bg_hex, fill_type="solid")
        cell_t.alignment = align_center
        
        # Fila 6: Valor / Fórmula
        ws.merge_cells(f"{c_start}6:{c_end}6")
        cell_v = ws[f"{c_start}6"]
        cell_v.value = formula
        cell_v.font = Font(name=FONT_FAMILY, size=15, bold=True, color=color_hex)
        cell_v.fill = PatternFill(start_color=bg_hex, end_color=bg_hex, fill_type="solid")
        cell_v.alignment = align_center
        
        if "%" in title:
            cell_v.number_format = '0.0%'
        elif "VARIACIÓN" in title:
            cell_v.number_format = '+#,##0 u;-#,##0 u;0 u'
        else:
            cell_v.number_format = '#,##0" u"'
            
    ws.row_dimensions[5].height = 18
    ws.row_dimensions[6].height = 26
    
    # 4. Cabeceras Agrupadas (Fila 8)
    ws.merge_cells('B8:E8')
    ws['B8'] = "DATOS CLIENTE Y ARTÍCULO"
    ws['B8'].font = font_group_hdr
    ws['B8'].fill = fill_grp_info
    ws['B8'].alignment = align_center
    
    ws.merge_cells('F8:J8')
    ws['F8'] = "OCTUBRE 2026 (EN CURSO - REAL AL 27/10/2026)"
    ws['F8'].font = font_group_hdr
    ws['F8'].fill = fill_grp_oct
    ws['F8'].alignment = align_center
    
    ws.merge_cells('K8:N8')
    ws['K8'] = "REVISIÓN PREVISIONES (NOVIEMBRE Y DICIEMBRE 2026)"
    ws['K8'].font = font_group_hdr
    ws['K8'].fill = fill_grp_q4
    ws['K8'].alignment = align_center
    
    ws.merge_cells('O8:R8')
    ws['O8'] = "IMPACTO TOTAL Q4 Y JUSTIFICACIÓN"
    ws['O8'].font = font_group_hdr
    ws['O8'].fill = fill_grp_tot
    ws['O8'].alignment = align_center
    ws.row_dimensions[8].height = 22
    
    # 5. Cabeceras Columnas (Fila 9)
    col_headers = [
        ('B', 'País', fill_col_hdr),
        ('C', 'Cliente', fill_col_hdr),
        ('D', 'Cód. SKU', fill_col_hdr),
        ('E', 'Descripción Artículo', fill_col_hdr),
        ('F', 'Oct Prev. Vigente (u)', fill_col_oct),
        ('G', 'Oct Real a 27/10 (u)', fill_col_oct),
        ('H', 'Pendiente / Desv. (u)', fill_col_oct),
        ('I', '% Consec. a 27/10', fill_col_oct),
        ('J', 'Comentario Octubre (¿Cierre previsto / Atraso / Extra?)', fill_col_oct),
        ('K', 'Nov Prev. Anterior (u)', fill_col_q4),
        ('L', 'Nov Revisado (u)', fill_col_q4),
        ('M', 'Dic Prev. Anterior (u)', fill_col_q4),
        ('N', 'Dic Revisado (u)', fill_col_q4),
        ('O', 'Total Q4 Ant (u)', fill_col_tot),
        ('P', 'Total Q4 Rev (u)', fill_col_tot),
        ('Q', 'Dif Q4 (u)', fill_col_tot),
        ('R', 'Comentarios / Justificación Q4', fill_col_tot),
    ]
    for col_letter, text, fill in col_headers:
        cell = ws[f"{col_letter}9"]
        cell.value = text
        cell.font = font_col_hdr
        cell.fill = fill
        cell.alignment = align_col_hdr
        cell.border = border_thin
    ws.row_dimensions[9].height = 28
    
    # 6. Filas de Datos (10 a tot_row-1)
    dv_coment = DataValidation(
        type="list",
        formula1='"Cierre según previsto,Atraso a Noviembre,Atraso a Diciembre,Adelanto de campaña,Pedido extra no previsto,Caída de venta"',
        allow_blank=True
    )
    ws.add_data_validation(dv_coment)
    
    for idx, item in enumerate(items):
        r = 10 + idx
        r_str = str(r)
        ws.row_dimensions[r].height = 19
        fill_base = fill_zebra_1 if idx % 2 == 1 else fill_zebra_0
        
        # B: País
        ws[f"B{r_str}"] = item['pais']
        ws[f"B{r_str}"].font = font_regular
        ws[f"B{r_str}"].fill = fill_base
        ws[f"B{r_str}"].alignment = align_center
        ws[f"B{r_str}"].border = border_thin
        
        # C: Cliente
        ws[f"C{r_str}"] = item['cliente']
        ws[f"C{r_str}"].font = font_regular
        ws[f"C{r_str}"].fill = fill_base
        ws[f"C{r_str}"].alignment = align_left
        ws[f"C{r_str}"].border = border_thin
        
        # D: SKU
        ws[f"D{r_str}"] = item['sku']
        ws[f"D{r_str}"].font = font_sku
        ws[f"D{r_str}"].fill = fill_base
        ws[f"D{r_str}"].alignment = align_center
        ws[f"D{r_str}"].border = border_thin
        
        # E: Descripción
        ws[f"E{r_str}"] = item['desc']
        ws[f"E{r_str}"].font = font_regular
        ws[f"E{r_str}"].fill = fill_base
        ws[f"E{r_str}"].alignment = align_left
        ws[f"E{r_str}"].border = border_thin
        
        # F: Oct Previsto Vigente (u)
        ws[f"F{r_str}"] = item['oct_prev_vigente']
        ws[f"F{r_str}"].font = font_regular
        ws[f"F{r_str}"].fill = fill_prev_hist
        ws[f"F{r_str}"].alignment = align_right
        ws[f"F{r_str}"].number_format = '#,##0'
        ws[f"F{r_str}"].border = border_thin
        
        # G: Oct Real a 27/10 (u)
        ws[f"G{r_str}"] = item['oct_real']
        ws[f"G{r_str}"].font = Font(name=FONT_FAMILY, size=9, bold=True, color="1E3A8A")
        ws[f"G{r_str}"].fill = fill_base
        ws[f"G{r_str}"].alignment = align_right
        ws[f"G{r_str}"].number_format = '#,##0'
        ws[f"G{r_str}"].border = border_thin
        
        # H: Pendiente / Desviación Octubre = |Real - Previsto|
        ws[f"H{r_str}"] = f"=ABS(G{r_str}-F{r_str})"
        ws[f"H{r_str}"].font = font_regular
        ws[f"H{r_str}"].alignment = align_right
        ws[f"H{r_str}"].number_format = '#,##0'
        ws[f"H{r_str}"].border = border_thin
        
        # I: % Consecución a fecha de hoy = Real / Previsto
        ws[f"I{r_str}"] = f'=IF(F{r_str}>0, G{r_str}/F{r_str}, IF(G{r_str}>0, 1.0, 0.0))'
        ws[f"I{r_str}"].font = font_regular
        ws[f"I{r_str}"].alignment = align_right
        ws[f"I{r_str}"].number_format = '0.0%'
        ws[f"I{r_str}"].border = border_thin
        
        # J: Comentario Octubre
        default_coment = ""
        if item['es_extra']:
            default_coment = "Pedido extra no previsto"
        elif item['oct_prev_vigente'] > 0 and item['oct_real'] >= item['oct_prev_vigente']:
            default_coment = "Cierre según previsto"
        ws[f"J{r_str}"] = default_coment
        ws[f"J{r_str}"].font = font_regular
        ws[f"J{r_str}"].fill = fill_comment_oct
        ws[f"J{r_str}"].alignment = align_left
        ws[f"J{r_str}"].border = border_thin
        dv_coment.add(ws[f"J{r_str}"])
        
        # K: Nov Prev. Anterior (u)
        ws[f"K{r_str}"] = item['nov_prev_vigente']
        ws[f"K{r_str}"].font = font_regular
        ws[f"K{r_str}"].fill = fill_prev_hist
        ws[f"K{r_str}"].alignment = align_right
        ws[f"K{r_str}"].number_format = '#,##0'
        ws[f"K{r_str}"].border = border_thin
        
        # L: Nov Revisado (u) (Editable verde: de partida igual a Nov Anterior)
        ws[f"L{r_str}"] = item['nov_prev_vigente']
        ws[f"L{r_str}"].font = font_editable
        ws[f"L{r_str}"].fill = fill_editable
        ws[f"L{r_str}"].alignment = align_right
        ws[f"L{r_str}"].number_format = '#,##0'
        ws[f"L{r_str}"].border = border_editable
        
        # M: Dic Prev. Anterior (u)
        ws[f"M{r_str}"] = item['dic_prev_vigente']
        ws[f"M{r_str}"].font = font_regular
        ws[f"M{r_str}"].fill = fill_prev_hist
        ws[f"M{r_str}"].alignment = align_right
        ws[f"M{r_str}"].number_format = '#,##0'
        ws[f"M{r_str}"].border = border_thin
        
        # N: Dic Revisado (u) (Editable verde: de partida igual a Dic Anterior)
        ws[f"N{r_str}"] = item['dic_prev_vigente']
        ws[f"N{r_str}"].font = font_editable
        ws[f"N{r_str}"].fill = fill_editable
        ws[f"N{r_str}"].alignment = align_right
        ws[f"N{r_str}"].number_format = '#,##0'
        ws[f"N{r_str}"].border = border_editable
        
        # O: Total Q4 Ant (u) = F (Oct Vigente) + K (Nov Ant) + M (Dic Ant)
        ws[f"O{r_str}"] = f"=F{r_str}+K{r_str}+M{r_str}"
        ws[f"O{r_str}"].font = font_regular
        ws[f"O{r_str}"].fill = fill_prev_hist
        ws[f"O{r_str}"].alignment = align_right
        ws[f"O{r_str}"].number_format = '#,##0'
        ws[f"O{r_str}"].border = border_thin
        
        # P: Total Q4 Rev (u) = F (Oct Vigente) + L (Nov Rev) + N (Dic Rev)
        ws[f"P{r_str}"] = f"=F{r_str}+L{r_str}+N{r_str}"
        ws[f"P{r_str}"].font = Font(name=FONT_FAMILY, size=9, bold=True, color="064E3B")
        ws[f"P{r_str}"].fill = fill_editable
        ws[f"P{r_str}"].alignment = align_right
        ws[f"P{r_str}"].number_format = '#,##0'
        ws[f"P{r_str}"].border = border_editable
        
        # Q: Dif Q4 (u) = P - O
        ws[f"Q{r_str}"] = f"=P{r_str}-O{r_str}"
        ws[f"Q{r_str}"].font = Font(name=FONT_FAMILY, size=9, bold=True)
        ws[f"Q{r_str}"].alignment = align_right
        ws[f"Q{r_str}"].number_format = '+#,##0;-#,##0;0'
        ws[f"Q{r_str}"].border = border_thin
        
        # R: Comentarios Q4
        ws[f"R{r_str}"] = ""
        ws[f"R{r_str}"].font = font_regular
        ws[f"R{r_str}"].alignment = align_left
        ws[f"R{r_str}"].border = border_thin
        
    # 7. Fila de Totales
    tot_r_str = str(tot_row)
    ws.row_dimensions[tot_row].height = 24
    
    ws.merge_cells(f"B{tot_r_str}:E{tot_r_str}")
    ws[f"B{tot_r_str}"] = f"TOTAL CONSOLIDADO - {com_info['display_name']} (UNIDADES)"
    ws[f"B{tot_r_str}"].font = font_total
    ws[f"B{tot_r_str}"].fill = fill_total
    ws[f"B{tot_r_str}"].alignment = align_center
    
    for col_idx in range(2, 6):
        c_let = get_column_letter(col_idx)
        ws[f"{c_let}{tot_r_str}"].border = border_total
        ws[f"{c_let}{tot_r_str}"].fill = fill_total
        
    sum_cols = [
        ('F', f"=SUM(F10:F{tot_row-1})", '#,##0'),
        ('G', f"=SUM(G10:G{tot_row-1})", '#,##0'),
        ('H', f"=SUM(H10:H{tot_row-1})", '#,##0'),
        ('I', f"=IF(F{tot_r_str}>0, G{tot_r_str}/F{tot_r_str}, 1.0)", '0.0%'),
        ('J', None, None),
        ('K', f"=SUM(K10:K{tot_row-1})", '#,##0'),
        ('L', f"=SUM(L10:L{tot_row-1})", '#,##0'),
        ('M', f"=SUM(M10:M{tot_row-1})", '#,##0'),
        ('N', f"=SUM(N10:N{tot_row-1})", '#,##0'),
        ('O', f"=SUM(O10:O{tot_row-1})", '#,##0'),
        ('P', f"=SUM(P10:P{tot_row-1})", '#,##0'),
        ('Q', f"=SUM(Q10:Q{tot_row-1})", '+#,##0;-#,##0;0'),
        ('R', None, None),
    ]
    
    for c_let, fmla, num_fmt in sum_cols:
        cell = ws[f"{c_let}{tot_r_str}"]
        cell.fill = fill_total
        cell.font = font_total
        cell.border = border_total
        cell.alignment = align_right
        if fmla:
            cell.value = fmla
            cell.number_format = num_fmt
            
    # Ajustar anchos de columnas
    col_widths = {
        'A': 3,
        'B': 14,
        'C': 34,
        'D': 15,
        'E': 32,
        'F': 16,
        'G': 16,
        'H': 16,
        'I': 14,
        'J': 26,
        'K': 16,
        'L': 16,
        'M': 16,
        'N': 16,
        'O': 16,
        'P': 16,
        'Q': 14,
        'R': 30
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    return tot_row

def build_summary_sheet(ws_res, commercial_data_meta):
    """Construye la pestaña de Resumen Consolidado Dirección en el libro consolidado."""
    ws_res.views.sheetView[0].showGridLines = True
    
    # Título
    ws_res.merge_cells('B2:M2')
    ws_res['B2'] = "CODIAGRO · RESUMEN DIRECCIÓN: REVISIÓN DE PREVISIONES Q4 2026"
    ws_res['B2'].font = font_title
    ws_res['B2'].fill = fill_title
    ws_res['B2'].alignment = align_center
    ws_res.row_dimensions[2].height = 28
    
    ws_res.merge_cells('B3:M3')
    ws_res['B3'] = "Consolidado de Octubre en Curso (Real a 27/10/2026) y Previsiones a Final de Año (Octubre a Diciembre) | Codiagro 2026"
    ws_res['B3'].font = font_subtitle
    ws_res['B3'].fill = fill_title
    ws_res['B3'].alignment = align_center
    ws_res.row_dimensions[3].height = 20
    
    # Cabeceras Agrupadas (Fila 5)
    ws_res.merge_cells('B5:C5')
    ws_res['B5'] = "INFORMACIÓN COMERCIAL"
    ws_res['B5'].font = font_group_hdr
    ws_res['B5'].fill = fill_grp_info
    ws_res['B5'].alignment = align_center
    
    ws_res.merge_cells('D5:G5')
    ws_res['D5'] = "OCTUBRE 2026 EN CURSO (REAL AL 27/10/2026)"
    ws_res['D5'].font = font_group_hdr
    ws_res['D5'].fill = fill_grp_oct
    ws_res['D5'].alignment = align_center
    
    ws_res.merge_cells('H5:I5')
    ws_res['H5'] = "REVISIÓN PREVISIONES (NOV Y DIC)"
    ws_res['H5'].font = font_group_hdr
    ws_res['H5'].fill = fill_grp_q4
    ws_res['H5'].alignment = align_center
    
    ws_res.merge_cells('J5:L5')
    ws_res['J5'] = "IMPACTO TOTAL Q4 2026"
    ws_res['J5'].font = font_group_hdr
    ws_res['J5'].fill = fill_grp_tot
    ws_res['J5'].alignment = align_center
    
    ws_res['M5'] = "ESTADO"
    ws_res['M5'].font = font_group_hdr
    ws_res['M5'].fill = fill_grp_tot
    ws_res['M5'].alignment = align_center
    ws_res.row_dimensions[5].height = 22
    
    # Cabeceras de Columnas (Fila 6)
    headers = [
        ('B', 'Comercial', fill_col_hdr),
        ('C', 'Clientes Activos', fill_col_hdr),
        ('D', 'Oct Previsto Vigente (u)', fill_col_oct),
        ('E', 'Oct Real a 27/10 (u)', fill_col_oct),
        ('F', 'Pendiente / Desv. (u)', fill_col_oct),
        ('G', '% Consec. a 27/10', fill_col_oct),
        ('H', 'Nov Revisado (u)', fill_col_q4),
        ('I', 'Dic Revisado (u)', fill_col_q4),
        ('J', 'Total Q4 Vigente (u)', fill_col_tot),
        ('K', 'Total Q4 Revisado (u)', fill_col_tot),
        ('L', 'Variación Neta Q4 (u)', fill_col_tot),
        ('M', 'Estado Octubre', fill_col_tot),
    ]
    for col_let, text, fill in headers:
        cell = ws_res[f"{col_let}6"]
        cell.value = text
        cell.font = font_col_hdr
        cell.fill = fill
        cell.alignment = align_col_hdr
        cell.border = border_thin
    ws_res.row_dimensions[6].height = 26
    
    # Filas 7 a 13: Comerciales
    for idx, c_meta in enumerate(commercial_data_meta):
        r = 7 + idx
        r_str = str(r)
        ws_res.row_dimensions[r].height = 20
        fill_base = fill_zebra_1 if idx % 2 == 1 else fill_zebra_0
        sheet_ref = f"'{c_meta['sheet_name']}'"
        tot_r_ref = str(c_meta['tot_row'])
        
        # B: Comercial
        ws_res[f"B{r_str}"] = c_meta['name']
        ws_res[f"B{r_str}"].font = Font(name=FONT_FAMILY, size=9, bold=True)
        ws_res[f"B{r_str}"].fill = fill_base
        ws_res[f"B{r_str}"].alignment = align_left
        ws_res[f"B{r_str}"].border = border_thin
        
        # C: Clientes Activos
        ws_res[f"C{r_str}"] = c_meta['num_clients']
        ws_res[f"C{r_str}"].font = font_regular
        ws_res[f"C{r_str}"].fill = fill_base
        ws_res[f"C{r_str}"].alignment = align_center
        ws_res[f"C{r_str}"].number_format = '#,##0'
        ws_res[f"C{r_str}"].border = border_thin
        
        # D: Oct Previsto Vigente
        ws_res[f"D{r_str}"] = f"={sheet_ref}!F{tot_r_ref}"
        ws_res[f"D{r_str}"].font = font_regular
        ws_res[f"D{r_str}"].fill = fill_base
        ws_res[f"D{r_str}"].alignment = align_right
        ws_res[f"D{r_str}"].number_format = '#,##0'
        ws_res[f"D{r_str}"].border = border_thin
        
        # E: Oct Real a 27/10
        ws_res[f"E{r_str}"] = f"={sheet_ref}!G{tot_r_ref}"
        ws_res[f"E{r_str}"].font = Font(name=FONT_FAMILY, size=9, bold=True, color="1E3A8A")
        ws_res[f"E{r_str}"].fill = fill_base
        ws_res[f"E{r_str}"].alignment = align_right
        ws_res[f"E{r_str}"].number_format = '#,##0'
        ws_res[f"E{r_str}"].border = border_thin
        
        # F: Pendiente / Desviación
        ws_res[f"F{r_str}"] = f"={sheet_ref}!H{tot_r_ref}"
        ws_res[f"F{r_str}"].font = font_regular
        ws_res[f"F{r_str}"].fill = fill_base
        ws_res[f"F{r_str}"].alignment = align_right
        ws_res[f"F{r_str}"].number_format = '#,##0'
        ws_res[f"F{r_str}"].border = border_thin
        
        # G: % Consecución Octubre
        ws_res[f"G{r_str}"] = f"={sheet_ref}!I{tot_r_ref}"
        ws_res[f"G{r_str}"].font = Font(name=FONT_FAMILY, size=9, bold=True)
        ws_res[f"G{r_str}"].fill = fill_base
        ws_res[f"G{r_str}"].alignment = align_right
        ws_res[f"G{r_str}"].number_format = '0.0%'
        ws_res[f"G{r_str}"].border = border_thin
        
        # H: Nov Revisado
        ws_res[f"H{r_str}"] = f"={sheet_ref}!L{tot_r_ref}"
        ws_res[f"H{r_str}"].font = font_regular
        ws_res[f"H{r_str}"].fill = fill_base
        ws_res[f"H{r_str}"].alignment = align_right
        ws_res[f"H{r_str}"].number_format = '#,##0'
        ws_res[f"H{r_str}"].border = border_thin
        
        # I: Dic Revisado
        ws_res[f"I{r_str}"] = f"={sheet_ref}!N{tot_r_ref}"
        ws_res[f"I{r_str}"].font = font_regular
        ws_res[f"I{r_str}"].fill = fill_base
        ws_res[f"I{r_str}"].alignment = align_right
        ws_res[f"I{r_str}"].number_format = '#,##0'
        ws_res[f"I{r_str}"].border = border_thin
        
        # J: Total Q4 Vigente
        ws_res[f"J{r_str}"] = f"={sheet_ref}!O{tot_r_ref}"
        ws_res[f"J{r_str}"].font = font_regular
        ws_res[f"J{r_str}"].fill = fill_base
        ws_res[f"J{r_str}"].alignment = align_right
        ws_res[f"J{r_str}"].number_format = '#,##0'
        ws_res[f"J{r_str}"].border = border_thin
        
        # K: Total Q4 Revisado
        ws_res[f"K{r_str}"] = f"={sheet_ref}!P{tot_r_ref}"
        ws_res[f"K{r_str}"].font = Font(name=FONT_FAMILY, size=9, bold=True, color="064E3B")
        ws_res[f"K{r_str}"].fill = fill_base
        ws_res[f"K{r_str}"].alignment = align_right
        ws_res[f"K{r_str}"].number_format = '#,##0'
        ws_res[f"K{r_str}"].border = border_thin
        
        # L: Variación Neta Q4
        ws_res[f"L{r_str}"] = f"={sheet_ref}!Q{tot_r_ref}"
        ws_res[f"L{r_str}"].font = Font(name=FONT_FAMILY, size=9, bold=True)
        ws_res[f"L{r_str}"].fill = fill_base
        ws_res[f"L{r_str}"].alignment = align_right
        ws_res[f"L{r_str}"].number_format = '+#,##0;-#,##0;0'
        ws_res[f"L{r_str}"].border = border_thin
        
        # M: Estado Octubre
        ws_res[f"M{r_str}"] = f'=IF(G{r_str}>=1.0, "Superado (+)", IF(G{r_str}>=0.7, "En Rango", "Revisar"))'
        ws_res[f"M{r_str}"].font = font_regular
        ws_res[f"M{r_str}"].fill = fill_base
        ws_res[f"M{r_str}"].alignment = align_center
        ws_res[f"M{r_str}"].border = border_thin
        
    # Fila 14: Total Codiagro
    ws_res.row_dimensions[14].height = 24
    ws_res['B14'] = "TOTAL CODIAGRO"
    ws_res['B14'].font = font_total
    ws_res['B14'].fill = fill_total
    ws_res['B14'].alignment = align_center
    ws_res['B14'].border = border_total
    
    tot_formulas = [
        ('C', "=SUM(C7:C13)", '#,##0'),
        ('D', "=SUM(D7:D13)", '#,##0'),
        ('E', "=SUM(E7:E13)", '#,##0'),
        ('F', "=SUM(F7:F13)", '#,##0'),
        ('G', "=E14/D14", '0.0%'),
        ('H', "=SUM(H7:H13)", '#,##0'),
        ('I', "=SUM(I7:I13)", '#,##0'),
        ('J', "=SUM(J7:J13)", '#,##0'),
        ('K', "=SUM(K7:K13)", '#,##0'),
        ('L', "=SUM(L7:L13)", '+#,##0;-#,##0;0'),
        ('M', "CONSOLIDADO", None),
    ]
    for c_let, fmla, num_fmt in tot_formulas:
        cell = ws_res[f"{c_let}14"]
        cell.font = font_total
        cell.fill = fill_total
        cell.border = border_total
        cell.alignment = align_center if c_let in ['C', 'M'] else align_right
        cell.value = fmla
        if num_fmt:
            cell.number_format = num_fmt
            
    # Anchos
    widths = {
        'A': 3, 'B': 16, 'C': 16, 'D': 22, 'E': 20, 'F': 20, 'G': 18,
        'H': 18, 'I': 18, 'J': 20, 'K': 20, 'L': 18, 'M': 16
    }
    for col_letter, width in widths.items():
        ws_res.column_dimensions[col_letter].width = width

def safe_save(wb, target_path):
    try:
        wb.save(target_path)
        print(f"  Guardado: {target_path}")
    except PermissionError:
        print(f"  [AVISO] {target_path} está abierto por el usuario en Excel. Se conserva la versión actual.")

def main():
    print("=== Generador Maestro de Revisiones Q4 2026 (Corte 27/10/2026) ===")
    
    # 1. Cargar pedidos y mapa oficial de comerciales
    pedidos_dict = load_real_orders()
    client_map = load_client_commercial_map()
    print(f"Pedidos al 27/10 cargados: {len(pedidos_dict)} combinaciones cliente-artículo.")
    print(f"Mapa comercial-cliente cargado: {len(client_map)} clientes mapeados.")
    
    out_dir = "Revision_Previsiones_Q4_2026"
    os.makedirs(out_dir, exist_ok=True)
    
    # Para el libro consolidado
    wb_cons = openpyxl.Workbook()
    # Primera hoja será el resumen
    ws_resumen = wb_cons.active
    ws_resumen.title = "Resumen Consolidado Dirección"
    
    commercial_meta = []
    
    # 2. Generar libro individual y pestaña consolidada por cada comercial
    for com_info in COMERCIALES:
        print(f"\nProcesando comercial: {com_info['display_name']}...")
        items = load_commercial_data(com_info, pedidos_dict, client_map)
        num_clients = len(set(it['cliente'] for it in items))
        print(f"  Líneas totales: {len(items)} (Clientes: {num_clients})")
        
        # A) Libro Individual
        wb_indiv = openpyxl.Workbook()
        ws_indiv = wb_indiv.active
        ws_indiv.title = com_info['individual_sheet']
        tot_row_indiv = build_commercial_sheet(ws_indiv, com_info, items)
        
        out_indiv_dir = os.path.join(out_dir, com_info['out_file'])
        out_indiv_root = com_info['out_file']
        safe_save(wb_indiv, out_indiv_dir)
        safe_save(wb_indiv, out_indiv_root)
        
        # B) Pestaña en Libro Consolidado
        ws_cons_com = wb_cons.create_sheet(title=com_info['sheet_name'])
        tot_row_cons = build_commercial_sheet(ws_cons_com, com_info, items)
        
        commercial_meta.append({
            'name': com_info['display_name'].capitalize() if com_info['display_name'] != 'GARCÍA' else 'García',
            'sheet_name': com_info['sheet_name'],
            'num_clients': num_clients,
            'tot_row': tot_row_cons
        })
        
    # 3. Construir pestaña de resumen consolidado
    print("\nConstruyendo pestaña 'Resumen Consolidado Dirección'...")
    build_summary_sheet(ws_resumen, commercial_meta)
    
    out_cons_dir = os.path.join(out_dir, "Revision_Previsiones_Q4_2026_CONSOLIDADO.xlsx")
    out_cons_root = "Revision_Previsiones_Q4_2026_CONSOLIDADO.xlsx"
    safe_save(wb_cons, out_cons_dir)
    safe_save(wb_cons, out_cons_root)
    
    print("\n¡Proceso completado con éxito para los 7 comerciales y el consolidado!")

if __name__ == '__main__':
    main()
