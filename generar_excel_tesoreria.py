import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime

print("=== Generando Prevision_Tesoreria_CashFlow_2026.xlsx ===")

with open('web_dashboard/treasury_data.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

wb = openpyxl.Workbook()
wb.remove(wb.active) # Eliminar hoja por defecto

# Estilos
COLOR_HEADER_BG = "0F172A"       # Slate 900
COLOR_GRP_REAL = "1E3A8A"        # Azul Cobalto (Real)
COLOR_GRP_PREV = "065F46"        # Verde Esmeralda (Previsión)
COLOR_ROW_ALT = "F8FAFC"
COLOR_ROW_WHITE = "FFFFFF"
COLOR_ROW_HIGHLIGHT = "FEF3C7"   # Ámbar suave para subtotales

font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
font_sub = Font(name="Calibri", size=9, italic=True, color="94A3B8")
font_head = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
font_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")
font_reg = Font(name="Calibri", size=9.5, color="1E293B")
font_tot = Font(name="Calibri", size=10.5, bold=True, color="0F172A")

fill_header = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
fill_real = PatternFill(start_color=COLOR_GRP_REAL, end_color=COLOR_GRP_REAL, fill_type="solid")
fill_prev = PatternFill(start_color=COLOR_GRP_PREV, end_color=COLOR_GRP_PREV, fill_type="solid")
fill_alt = PatternFill(start_color=COLOR_ROW_ALT, end_color=COLOR_ROW_ALT, fill_type="solid")
fill_white = PatternFill(start_color=COLOR_ROW_WHITE, end_color=COLOR_ROW_WHITE, fill_type="solid")
fill_subtot = PatternFill(start_color=COLOR_ROW_HIGHLIGHT, end_color=COLOR_ROW_HIGHLIGHT, fill_type="solid")

thin_border_side = Side(border_style="thin", color="CBD5E1")
border_thin = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
border_tot = Border(top=Side(border_style="thin", color="0F172A"), bottom=Side(border_style="double", color="0F172A"))

# -------------------------------------------------------------
# HOJA 1: CASH FLOW MENSUAL 2026
# -------------------------------------------------------------
ws_cf = wb.create_sheet(title="Cash Flow 2026")
ws_cf.views.sheetView[0].showGridLines = True

# Título y Subtítulo
ws_cf.merge_cells("A1:Q1")
c1 = ws_cf["A1"]
c1.value = "  CODIAGRO S.L.  |  PREVISIÓN DE TESORERÍA Y CASH FLOW MENSUAL 2026"
c1.font = font_title
c1.fill = fill_header
c1.alignment = Alignment(vertical="center", horizontal="left")
ws_cf.row_dimensions[1].height = 28

ws_cf.merge_cells("A2:Q2")
c2 = ws_cf["A2"]
c2.value = f"  Datos Reales Auditados: Ene - Sep 2026 (Saldo Bancos Sep: 3.203.215 €)  |  Proyección Q4 basada en Vencimientos y Ventas  |  Fecha: {datetime.now().strftime('%d/%m/%Y')}"
c2.font = font_sub
c2.fill = fill_header
c2.alignment = Alignment(vertical="center", horizontal="left")
ws_cf.row_dimensions[2].height = 18

ws_cf.row_dimensions[3].height = 10

# Cabecera Grupos
ws_cf.merge_cells("A4:B4")
ws_cf["A4"].value = "PARTIDAS Y PRESUPUESTO"
ws_cf["A4"].font = font_head
ws_cf["A4"].fill = fill_header
ws_cf["A4"].alignment = Alignment(horizontal="center", vertical="center")

ws_cf.merge_cells("C4:K4")
ws_cf["C4"].value = "EJERCICIO 2026 REAL (AUDITADO)"
ws_cf["C4"].font = font_head
ws_cf["C4"].fill = fill_real
ws_cf["C4"].alignment = Alignment(horizontal="center", vertical="center")

ws_cf.merge_cells("L4:N4")
ws_cf["L4"].value = "PROYECCIÓN Q4 (VENTAS & COBROS)"
ws_cf["L4"].font = font_head
ws_cf["L4"].fill = fill_prev
ws_cf["L4"].alignment = Alignment(horizontal="center", vertical="center")

ws_cf.merge_cells("O4:Q4")
ws_cf["O4"].value = "CIERRE ESTIMADO 2026"
ws_cf["O4"].font = font_head
ws_cf["O4"].fill = fill_header
ws_cf["O4"].alignment = Alignment(horizontal="center", vertical="center")

ws_cf.row_dimensions[4].height = 20

# Encabezados de Columna
headers = [
    ("A", "Concepto / Partida Tesorería", 36, "left"),
    ("B", "Presupuesto 2026", 18, "right"),
    ("C", "Ene (R)", 13, "right"),
    ("D", "Feb (R)", 13, "right"),
    ("E", "Mar (R)", 13, "right"),
    ("F", "Abr (R)", 13, "right"),
    ("G", "May (R)", 13, "right"),
    ("H", "Jun (R)", 13, "right"),
    ("I", "Jul (R)", 13, "right"),
    ("J", "Ago (R)", 13, "right"),
    ("K", "Sep (R)", 13, "right"),
    ("L", "Oct (Prev)", 14, "right"),
    ("M", "Nov (Prev)", 14, "right"),
    ("N", "Dic (Prev)", 14, "right"),
    ("O", "Total Real+Prev", 16, "right"),
    ("P", "Desviación Presup.", 16, "right"),
    ("Q", "% Cumplimiento", 15, "right")
]

for col_l, text, width, align in headers:
    cell = ws_cf[f"{col_l}5"]
    cell.value = text
    cell.font = font_head
    if col_l in ["L", "M", "N"]:
        cell.fill = fill_prev
    elif col_l in ["C", "D", "E", "F", "G", "H", "I", "J", "K"]:
        cell.fill = fill_real
    else:
        cell.fill = fill_header
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border = border_thin
    ws_cf.column_dimensions[col_l].width = width

ws_cf.row_dimensions[5].height = 24

# Filas de datos
cur_row = 6
q4_cobros = data.get("q4_cobros_proyectados", {})
oct_cobros_tot = sum(q4_cobros.get("2026-10", {}).values())
nov_cobros_tot = sum(q4_cobros.get("2026-11", {}).values())
dic_cobros_tot = sum(q4_cobros.get("2026-12", {}).values())

for r in data.get("cf_rows", []):
    label = r.get("label", "")
    budget = r.get("budget_annual", 0) or 0
    reales = r.get("reales", {})
    
    budget_q4 = r.get("budget_q4", {})
    oct_val = budget_q4.get("2026-10", 0.0) or 0.0
    nov_val = budget_q4.get("2026-11", 0.0) or 0.0
    dic_val = budget_q4.get("2026-12", 0.0) or 0.0
    
    if label == "Cobros por canal Clientes nacionales":
        oct_val = q4_cobros.get("2026-10", {}).get("Cobros por canal Clientes nacionales", 0.0)
        nov_val = q4_cobros.get("2026-11", {}).get("Cobros por canal Clientes nacionales", 0.0)
        dic_val = q4_cobros.get("2026-12", {}).get("Cobros por canal Clientes nacionales", 0.0)
    elif label == "Cobros por canal Clientes internac.":
        oct_val = q4_cobros.get("2026-10", {}).get("Cobros por canal Clientes internac.", 0.0)
        nov_val = q4_cobros.get("2026-11", {}).get("Cobros por canal Clientes internac.", 0.0)
        dic_val = q4_cobros.get("2026-12", {}).get("Cobros por canal Clientes internac.", 0.0)
    elif label == "Cobros por canal U.E.":
        oct_val = q4_cobros.get("2026-10", {}).get("Cobros por canal U.E.", 0.0)
        nov_val = q4_cobros.get("2026-11", {}).get("Cobros por canal U.E.", 0.0)
        dic_val = q4_cobros.get("2026-12", {}).get("Cobros por canal U.E.", 0.0)
    elif label in ["Total Cobros", "Total Cobros con ajuste"]:
        oct_val = oct_cobros_tot
        nov_val = nov_cobros_tot
        dic_val = dic_cobros_tot

    sum_real = sum((reales.get(f"2026-{str(i).zfill(2)}", 0) or 0) for i in range(1, 10))
    tot_anual = sum_real + oct_val + nov_val + dic_val
    desv = tot_anual - budget if budget != 0 else 0
    pct = (tot_anual / budget * 100) if budget != 0 else None

    is_highlight = label.startswith("Total") or "Tesorería" in label or "Cash flow" in label
    row_fill = fill_subtot if is_highlight else (fill_alt if cur_row % 2 == 0 else fill_white)
    row_font = font_bold if is_highlight else font_reg

    ws_cf.cell(cur_row, 1, label).font = row_font
    ws_cf.cell(cur_row, 1).alignment = Alignment(horizontal="left", vertical="center")
    
    ws_cf.cell(cur_row, 2, budget).number_format = '#,##0'
    ws_cf.cell(cur_row, 2).font = row_font
    ws_cf.cell(cur_row, 2).alignment = Alignment(horizontal="right", vertical="center")

    for i in range(1, 10):
        val_m = reales.get(f"2026-{str(i).zfill(2)}", 0) or 0
        c_m = ws_cf.cell(cur_row, 2 + i, val_m)
        c_m.number_format = '#,##0'
        c_m.font = row_font
        c_m.alignment = Alignment(horizontal="right", vertical="center")

    c_oct = ws_cf.cell(cur_row, 12, oct_val)
    c_oct.number_format = '#,##0'
    c_oct.font = font_bold if oct_val > 0 else row_font
    c_oct.alignment = Alignment(horizontal="right", vertical="center")

    c_nov = ws_cf.cell(cur_row, 13, nov_val)
    c_nov.number_format = '#,##0'
    c_nov.font = font_bold if nov_val > 0 else row_font
    c_nov.alignment = Alignment(horizontal="right", vertical="center")

    c_dic = ws_cf.cell(cur_row, 14, dic_val)
    c_dic.number_format = '#,##0'
    c_dic.font = font_bold if dic_val > 0 else row_font
    c_dic.alignment = Alignment(horizontal="right", vertical="center")

    c_tot = ws_cf.cell(cur_row, 15, tot_anual)
    c_tot.number_format = '#,##0'
    c_tot.font = font_bold
    c_tot.alignment = Alignment(horizontal="right", vertical="center")

    c_desv = ws_cf.cell(cur_row, 16, desv)
    c_desv.number_format = '#,##0'
    c_desv.font = font_bold
    c_desv.alignment = Alignment(horizontal="right", vertical="center")

    c_pct = ws_cf.cell(cur_row, 17)
    if pct is not None:
        c_pct.value = pct / 100.0
        c_pct.number_format = '0.0%'
    else:
        c_pct.value = "—"
    c_pct.font = font_bold
    c_pct.alignment = Alignment(horizontal="center", vertical="center")

    for col_i in range(1, 18):
        c_cell = ws_cf.cell(cur_row, col_i)
        c_cell.fill = row_fill
        c_cell.border = border_tot if is_highlight else border_thin

    ws_cf.row_dimensions[cur_row].height = 20
    cur_row += 1


# -------------------------------------------------------------
# HOJA 2: VENCIMIENTOS CLIENTES Q4
# -------------------------------------------------------------
ws_cli = wb.create_sheet(title="Vencimientos Clientes Q4")
ws_cli.views.sheetView[0].showGridLines = True

ws_cli.merge_cells("A1:M1")
ws_cli["A1"].value = "  CODIAGRO S.L.  |  DETALLE FACTURACIÓN Y VENCIMIENTOS CLIENTES Q4 2026 (CON RETRASOS REALES)"
ws_cli["A1"].font = font_title
ws_cli["A1"].fill = fill_header
ws_cli["A1"].alignment = Alignment(vertical="center", horizontal="left")
ws_cli.row_dimensions[1].height = 28

cli_cols = [
    ("A", "Origen / Tipo", 18, "center"),
    ("B", "Cliente", 38, "left"),
    ("C", "Comercial", 16, "center"),
    ("D", "Canal Comercial", 26, "center"),
    ("E", "País", 12, "center"),
    ("F", "Plazo Pactado", 15, "center"),
    ("G", "Plazo Real / Efectivo", 20, "center"),
    ("H", "Fecha Operación / Factura", 24, "center"),
    ("I", "Fecha Vencimiento", 18, "center"),
    ("J", "Mes Cobro Efectivo", 18, "center"),
    ("K", "Base Imponible (€)", 20, "right"),
    ("L", "IVA (€)", 16, "right"),
    ("M", "Total a Cobrar (€)", 20, "right")
]

for col_l, text, width, align in cli_cols:
    cell = ws_cli[f"{col_l}3"]
    cell.value = text
    cell.font = font_head
    cell.fill = fill_real
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border = border_thin
    ws_cli.column_dimensions[col_l].width = width

ws_cli.row_dimensions[3].height = 24

cur_cr = 4
venc_q4 = [x for x in data.get("client_vencimientos", []) if x.get("mes_vto") in ["10-2026", "11-2026", "12-2026"]]

for v in venc_q4:
    row_fill = fill_alt if cur_cr % 2 == 0 else fill_white
    
    orig_badge = v.get("tipo_origen", "Presupuesto (BG)")
    ws_cli.cell(cur_cr, 1, orig_badge).font = font_bold
    ws_cli.cell(cur_cr, 1).alignment = Alignment(horizontal="center", vertical="center")

    ws_cli.cell(cur_cr, 2, v.get("cliente")).font = font_bold
    ws_cli.cell(cur_cr, 2).alignment = Alignment(horizontal="left", vertical="center")

    ws_cli.cell(cur_cr, 3, v.get("comercial") or "—").font = font_reg
    ws_cli.cell(cur_cr, 3).alignment = Alignment(horizontal="center", vertical="center")
    
    ws_cli.cell(cur_cr, 4, v.get("canal")).font = font_reg
    ws_cli.cell(cur_cr, 4).alignment = Alignment(horizontal="center", vertical="center")

    ws_cli.cell(cur_cr, 5, v.get("pais")).font = font_reg
    ws_cli.cell(cur_cr, 5).alignment = Alignment(horizontal="center", vertical="center")

    d_pact = v.get("dias_pactados", v.get("dias"))
    ws_cli.cell(cur_cr, 6, f"{d_pact} d").font = font_reg
    ws_cli.cell(cur_cr, 6).alignment = Alignment(horizontal="center", vertical="center")

    d_efect = v.get("dias_efectivos", v.get("dias"))
    ws_cli.cell(cur_cr, 7, f"{d_efect} d").font = font_bold
    ws_cli.cell(cur_cr, 7).alignment = Alignment(horizontal="center", vertical="center")

    f_op = v.get("fecha_operacion") or v.get("fecha_factura") or ""
    ws_cli.cell(cur_cr, 8, f_op).font = font_reg
    ws_cli.cell(cur_cr, 8).alignment = Alignment(horizontal="center", vertical="center")

    ws_cli.cell(cur_cr, 9, v.get("vto_fecha")).font = font_reg
    ws_cli.cell(cur_cr, 9).alignment = Alignment(horizontal="center", vertical="center")

    ws_cli.cell(cur_cr, 10, v.get("mes_vto")).font = font_bold
    ws_cli.cell(cur_cr, 10).alignment = Alignment(horizontal="center", vertical="center")

    c_b = ws_cli.cell(cur_cr, 11, v.get("facturacion", 0))
    c_b.number_format = '#,##0.00'
    c_b.font = font_reg
    c_b.alignment = Alignment(horizontal="right", vertical="center")

    c_i = ws_cli.cell(cur_cr, 12, v.get("iva", 0))
    c_i.number_format = '#,##0.00'
    c_i.font = font_reg
    c_i.alignment = Alignment(horizontal="right", vertical="center")

    c_t = ws_cli.cell(cur_cr, 13, v.get("importe", 0))
    c_t.number_format = '#,##0.00'
    c_t.font = font_bold
    c_t.alignment = Alignment(horizontal="right", vertical="center")

    for col_i in range(1, 14):
        ws_cli.cell(cur_cr, col_i).fill = row_fill
        ws_cli.cell(cur_cr, col_i).border = border_thin

    ws_cli.row_dimensions[cur_cr].height = 20
    cur_cr += 1

# Fila Totales Hoja 2
tot_cr = cur_cr
ws_cli.merge_cells(f"A{tot_cr}:J{tot_cr}")
ws_cli[f"A{tot_cr}"].value = "TOTAL COBROS PREVISTOS Q4 2026"
ws_cli[f"A{tot_cr}"].font = font_tot
ws_cli[f"A{tot_cr}"].alignment = Alignment(horizontal="right", vertical="center")

ws_cli.cell(tot_cr, 11, f"=SUM(K4:K{tot_cr - 1})").number_format = '#,##0.00'
ws_cli.cell(tot_cr, 11).font = font_tot
ws_cli.cell(tot_cr, 11).alignment = Alignment(horizontal="right", vertical="center")

ws_cli.cell(tot_cr, 12, f"=SUM(L4:L{tot_cr - 1})").number_format = '#,##0.00'
ws_cli.cell(tot_cr, 12).font = font_tot
ws_cli.cell(tot_cr, 12).alignment = Alignment(horizontal="right", vertical="center")

ws_cli.cell(tot_cr, 13, f"=SUM(M4:M{tot_cr - 1})").number_format = '#,##0.00'
ws_cli.cell(tot_cr, 13).font = font_tot
ws_cli.cell(tot_cr, 13).alignment = Alignment(horizontal="right", vertical="center")

for col_i in range(1, 14):
    ws_cli.cell(tot_cr, col_i).fill = fill_subtot
    ws_cli.cell(tot_cr, col_i).border = border_tot

ws_cli.row_dimensions[tot_cr].height = 24


# -------------------------------------------------------------
# HOJA 3: FORMAS DE PAGO DE CLIENTES (CATÁLOGO REAL CON RETRASOS)
# -------------------------------------------------------------
ws_cond = wb.create_sheet(title="Condiciones de Pago")
ws_cond.views.sheetView[0].showGridLines = True

ws_cond.merge_cells("A1:K1")
ws_cond["A1"].value = "  CODIAGRO S.L.  |  CATÁLOGO DE CONDICIONES DE PAGO Y PLAZOS REALES DE COBRO"
ws_cond["A1"].font = font_title
ws_cond["A1"].fill = fill_header
ws_cond["A1"].alignment = Alignment(vertical="center", horizontal="left")
ws_cond.row_dimensions[1].height = 28

cond_cols = [
    ("A", "Código", 14, "center"),
    ("B", "Tipo Canal", 18, "center"),
    ("C", "Comercial", 18, "center"),
    ("D", "Razón Social / Cliente", 42, "left"),
    ("E", "Forma de Pago Acordada", 36, "left"),
    ("F", "Plazo Pactado (Días)", 20, "center"),
    ("G", "Plazo Real Observado (Días)", 24, "center"),
    ("H", "Plazo Efectivo Modelo (Días)", 24, "center"),
    ("I", "Retraso Aplicado (Días)", 20, "center"),
    ("J", "Total Facturado 2026 (€)", 22, "right"),
    ("K", "Total Cobrado/Prev. 2026 (€)", 24, "right")
]

for col_l, text, width, align in cond_cols:
    cell = ws_cond[f"{col_l}3"]
    cell.value = text
    cell.font = font_head
    cell.fill = fill_real
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border = border_thin
    ws_cond.column_dimensions[col_l].width = width

ws_cond.row_dimensions[3].height = 24

cur_p = 4
for p in data.get("client_master", []):
    row_fill = fill_alt if cur_p % 2 == 0 else fill_white
    
    ws_cond.cell(cur_p, 1, p.get("codigo", "—")).font = font_reg
    ws_cond.cell(cur_p, 1).alignment = Alignment(horizontal="center", vertical="center")

    ws_cond.cell(cur_p, 2, p.get("tipo", "Nacional")).font = font_reg
    ws_cond.cell(cur_p, 2).alignment = Alignment(horizontal="center", vertical="center")

    ws_cond.cell(cur_p, 3, p.get("comercial") or "—").font = font_reg
    ws_cond.cell(cur_p, 3).alignment = Alignment(horizontal="center", vertical="center")

    ws_cond.cell(cur_p, 4, p.get("cliente")).font = font_bold
    ws_cond.cell(cur_p, 4).alignment = Alignment(horizontal="left", vertical="center")

    ws_cond.cell(cur_p, 5, p.get("forma_pago", "—")).font = font_reg
    ws_cond.cell(cur_p, 5).alignment = Alignment(horizontal="left", vertical="center")

    ws_cond.cell(cur_p, 6, p.get("dias_pactados", p.get("dias", 0))).font = font_reg
    ws_cond.cell(cur_p, 6).alignment = Alignment(horizontal="center", vertical="center")

    d_obs = p.get("dias_reales_obs")
    obs_str = f"{d_obs:.1f} d" if d_obs is not None else "—"
    ws_cond.cell(cur_p, 7, obs_str).font = font_reg
    ws_cond.cell(cur_p, 7).alignment = Alignment(horizontal="center", vertical="center")

    d_ef = p.get("dias_efectivos", p.get("dias", 0))
    ws_cond.cell(cur_p, 8, f"{d_ef} d").font = font_bold
    ws_cond.cell(cur_p, 8).alignment = Alignment(horizontal="center", vertical="center")

    retraso = p.get("retraso_dias", 0)
    ret_str = f"+{retraso} d" if retraso > 0 else (f"{retraso} d" if retraso < 0 else "0 d")
    ws_cond.cell(cur_p, 9, ret_str).font = font_bold
    ws_cond.cell(cur_p, 9).alignment = Alignment(horizontal="center", vertical="center")

    c_fact = ws_cond.cell(cur_p, 10, p.get("total_facturacion", 0))
    c_fact.number_format = '#,##0.00'
    c_fact.font = font_reg
    c_fact.alignment = Alignment(horizontal="right", vertical="center")

    c_imp = ws_cond.cell(cur_p, 11, p.get("total_cobros", 0))
    c_imp.number_format = '#,##0.00'
    c_imp.font = font_bold
    c_imp.alignment = Alignment(horizontal="right", vertical="center")

    for col_i in range(1, 12):
        ws_cond.cell(cur_p, col_i).fill = row_fill
        ws_cond.cell(cur_p, col_i).border = border_thin

    ws_cond.row_dimensions[cur_p].height = 20
    cur_p += 1

out_path = "web_dashboard/Prevision_Tesoreria_CashFlow_2026.xlsx"
try:
    wb.save(out_path)
    print(f"Archivo Excel generado exitosamente en: {out_path}")
except PermissionError:
    alt_path = "web_dashboard/Prevision_Tesoreria_CashFlow_2026_v2.xlsx"
    wb.save(alt_path)
    print(f"[AVISO] {out_path} está abierto en Excel (bloqueado). Guardado como {alt_path}")
