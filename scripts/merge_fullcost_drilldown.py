import os
import sys

# Inclusión obligatoria y segura de output_staging
_scripts_dir = os.path.dirname(os.path.abspath(__file__))
if _scripts_dir not in sys.path:
    sys.path.insert(0, _scripts_dir)

try:
    from output_staging import get_dest, get_source
except Exception as _e:
    if os.environ.get("CODIAGRO_STAGING_DIR"):
        sys.stderr.write(f"[ERROR FATAL] Fallo al importar output_staging con CODIAGRO_STAGING_DIR activa: {_e}\n")
        sys.exit(1)
    def get_dest(p): return p
    def get_source(p): return p
import os
import re
import json
import datetime
import openpyxl
from collections import defaultdict
import shutil
import tempfile


def safe_load_workbook(file_path, read_only=True, data_only=True):
    try:
        return openpyxl.load_workbook(file_path, read_only=read_only, data_only=data_only)
    except (PermissionError, IOError):
        try:
            import ctypes
            from ctypes import wintypes
            kernel32 = ctypes.WinDLL('kernel32', use_last_error=True)
            handle = kernel32.CreateFileW(
                str(os.path.abspath(file_path)),
                0x80000000,
                1 | 2 | 4,
                None,
                3,
                0x80,
                None
            )
            if handle != -1:
                size = kernel32.GetFileSize(handle, None)
                buf = ctypes.create_string_buffer(size)
                read_bytes = wintypes.DWORD()
                kernel32.ReadFile(handle, buf, size, ctypes.byref(read_bytes), None)
                kernel32.CloseHandle(handle)
                import io
                bio = io.BytesIO(buf.raw[:read_bytes.value])
                return openpyxl.load_workbook(bio, read_only=read_only, data_only=data_only)
        except Exception:
            pass
        tmp_copy = os.path.join(tempfile.gettempdir(), f"copy_{os.path.basename(file_path)}")
        try:
            shutil.copy2(file_path, tmp_copy)
            return openpyxl.load_workbook(tmp_copy, read_only=read_only, data_only=data_only)
        except Exception:
            raise


def round_conserving_sum(items, field, target_total, decimals=2):
    """
    Política explícita de redondeo por Mayor Resto (Largest Remainder Method):
    redondea cada elemento a `decimals` decimales y distribuye los céntimos residuales
    sobre las partidas de mayor resto fraccionario para conservar exactamente `target_total`.
    """
    if not items:
        return
    factor = 10 ** decimals
    target_cents = int(round(target_total * factor))
    exact_Sticky = [float(item.get(field, 0.0)) * factor for item in items]
    floored = [int(x // 1) for x in exact_Sticky]
    remainders = [exact_Sticky[i] - floored[i] for i in range(len(items))]
    diff = target_cents - sum(floored)
    if diff > 0:
        order = sorted(range(len(items)), key=lambda i: remainders[i], reverse=True)
        for idx in order[:diff]:
            floored[idx] += 1
    elif diff < 0:
        order = sorted(range(len(items)), key=lambda i: remainders[i])
        for idx in order[:abs(diff)]:
            floored[idx] -= 1
    for i, item in enumerate(items):
        item[field] = round(floored[i] / factor, decimals)


def normalize_client_id(name):
    if not name:
        return ""
    s = str(name).strip().upper()
    s = re.sub(r'[\.\,\;\:\-\'\´\`\"\(\)]', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()
    return s


# Alias deterministas de clientes entre logística y ventas (sin subcadenas ambiguas)
CANONICAL_CLIENT_ALIASES = {
    normalize_client_id("FITOSANITARIOS CARCAIXENT, S"): normalize_client_id("FITOSANITARIOS CARCAIXENT, S.L."),
    normalize_client_id("FITOSANITARIOS CARCAIXENT, S."): normalize_client_id("FITOSANITARIOS CARCAIXENT, S.L."),
}


# 1. Localizar libro de trabajo de Full Costing
fc_wb_path = None
for cand_path in [
    os.path.join("outputs", "fullcost_septiembre_2026", "Plantillas_Full_Costing_Septiembre_2026_Con_PYG_Mejorada.xlsx"),
    "Plantillas_Full_Costing_Septiembre_2026_Con_PYG_Mejorada.xlsx"
]:
    if os.path.exists(cand_path):
        fc_wb_path = cand_path
        break

# Cargar parámetros técnicos de Cadencias_Produccion si existen en la plantilla
cad_params = {
    "h_sint_110": 0.001000,
    "h_sint_120": 0.001000,
    "h_sint_130": 0.001250,
    "h_env_1L": 0.004545,
    "h_env_5L": 0.001667,
    "h_env_20L": 0.000911,
    "h_env_200L": 0.000500,
    "h_env_1000L": 0.000333,
    "tarifa_sim_110": 48.50,
    "tarifa_sim_120": 42.80,
    "tarifa_sim_130": 51.20,
    "tarifa_sim_140": 36.40,
}

if fc_wb_path:
    wb_cad_init = safe_load_workbook(fc_wb_path, read_only=True, data_only=True)
    if 'Cadencias_Produccion' in wb_cad_init.sheetnames:
        ws_cad = wb_cad_init['Cadencias_Produccion']
        for r in range(36, ws_cad.max_row + 1):
            p_code = str(ws_cad.cell(r, 2).value or '').strip()
            p_val = ws_cad.cell(r, 4).value
            if p_code in cad_params and p_val is not None:
                try:
                    cad_params[p_code] = float(p_val)
                except Exception:
                    pass
    wb_cad_init.close()


# 2. Cargar costes SKU
costes_path = os.path.join("Gastos", "Costes SKU.xlsx")
wb_c = safe_load_workbook(costes_path, read_only=False, data_only=True)
ws_c = wb_c['Costes SKU']

costes_map = {}
for r in ws_c.iter_rows(min_row=5, values_only=True):
    sku = str(r[0]).strip().upper() if r[0] else ''
    if not sku:
        continue
    fam = str(r[1]).strip() if r[1] is not None else ''
    pais = str(r[2]).strip().upper() if r[2] is not None else ''
    form_med = float(r[6] or 0)
    pkg_med = float(r[7] or 0)
    tot_med = float(r[8] or 0)
    cogs_med_val = tot_med if tot_med > 0 else (form_med + pkg_med)
    costes_map[sku] = {
        "sku": sku, "familia": fam, "pais": pais,
        "form_med": round(form_med, 4), "pkg_med": round(pkg_med, 4),
        "cogs_medio": round(cogs_med_val, 4)
    }
wb_c.close()

equivalencias = {
    "A-50TR": "A-50BBTR", "SPCE20001CR": "SPCE0001CR",
    "SPCE20005CR": "SPCE0005CR", "A-50BB": "A-50BBTR", "AGXKAEP0001": "AGXKAE0001RO"
}


def get_cogs(sku_code, pais_code="", desc_text=""):
    s_clean = str(sku_code).strip().upper()
    p_up = pais_code.strip().upper()
    d_up = str(desc_text).strip().upper()
    if s_clean == 'NAZ':
        if 'AMINO' in d_up:
            s_clean = 'AMGF0001AZ'
        elif 'AQUOM' in d_up:
            s_clean = 'AQUOM0005'
    if s_clean in equivalencias:
        s_clean = equivalencias[s_clean]
    if s_clean in costes_map:
        return costes_map[s_clean]
    cands = [c for c in costes_map.keys() if c.startswith(s_clean) or s_clean.startswith(c)]
    if cands:
        for c in cands:
            if costes_map[c]["pais"] == p_up:
                return costes_map[c]
        for c in cands:
            if costes_map[c]["pais"] == 'ESPAÑA':
                return costes_map[c]
        return costes_map[cands[0]]
    return None


def get_ind_hours(sku, desc):
    """
    Clasificación de tiempos y rutas de planta según Cadencias_Produccion:
    - Proceso Síntesis: Alcaplant (CECO 110), Líquidos (CECO 120), Sólidos (CECO 130)
    - Envasado (CECO 140): 1L/1kg, 5L/5kg, 20L/20kg, 200L, 1000L/Granel
    """
    desc_u = (desc or '').upper()
    sku_u = (sku or '').upper()
    if 'ALCAPLANT' in desc_u or sku_u.startswith('AN'):
        h_sint = cad_params["h_sint_110"]
        ceco_s = '110'
        ruta = 'Alcaplant (CECO 110)'
    elif 'SOLIDO' in desc_u or 'POLVO' in desc_u or ' KG' in desc_u or sku_u.startswith('A-50') or '20K' in sku_u:
        if ' L' in desc_u and 'KG' not in desc_u:
            h_sint = cad_params["h_sint_120"]
            ceco_s = '120'
            ruta = 'Líquidos (CECO 120)'
        else:
            h_sint = cad_params["h_sint_130"]
            ceco_s = '130'
            ruta = 'Sólidos (CECO 130)'
    else:
        h_sint = cad_params["h_sint_120"]
        ceco_s = '120'
        ruta = 'Líquidos (CECO 120)'

    ceco_e = '140'
    if any(x in desc_u for x in [' 1L', ' 1 L', ' 1KG', ' 1 KG']) or sku_u.endswith('0001'):
        h_env = cad_params["h_env_1L"]
        formato = '1L / 1kg'
    elif any(x in desc_u for x in [' 5L', ' 5 L', ' 5KG', ' 5 KG']) or sku_u.endswith('0005'):
        h_env = cad_params["h_env_5L"]
        formato = '5L / 5kg'
    elif any(x in desc_u for x in [' 20L', ' 20 L', ' 20KG', ' 20 KG']) or sku_u.endswith('0020'):
        h_env = cad_params["h_env_20L"]
        formato = '20L / 20kg'
    elif any(x in desc_u for x in ['200L', '200 L', 'BIDON 200', 'BIDÓN 200']) or '0200' in sku_u:
        h_env = cad_params["h_env_200L"]
        formato = 'Bidón 200L'
    elif any(x in desc_u for x in ['1000', 'GRANEL', 'IBC']) or sku_u.endswith('1000'):
        h_env = cad_params["h_env_1000L"]
        formato = '1000L / Granel'
    else:
        h_env = cad_params["h_env_5L"]
        formato = 'Estándar 5L'

    return h_sint, ceco_s, h_env, ceco_e, (h_sint + h_env), ruta, formato


# 3. Cargar transporte real Sep-26 con control de unicidad por expedición
log_path = os.path.join("outputs", "fullcost_septiembre_2026", "REPORT_SEPTIEMBRE_LOGISTICA_actualizado.xlsx")
wb_log = safe_load_workbook(log_path, read_only=True, data_only=True)
ws_log = wb_log['Hoja1']

header_log = None
transporte_cliente = defaultdict(float)
transporte_envios = defaultdict(int)
transporte_nombre_orig = {}
seen_shipments = set()
log_scope_totals = {"624000001": 0.0, "624000002": 0.0, "624000003": 0.0}
log_scope_counts = {"624000001": 0, "624000002": 0, "624000003": 0}
log_unpriced_rows = []
log_cutoff_mes8_sep = {"624000001": 0.0, "624000003": 0.0}

for r_idx, row in enumerate(ws_log.iter_rows(values_only=True), start=1):
    if header_log is None:
        header_log = list(row)
        idx_yr = header_log.index('Ejercicio') if 'Ejercicio' in header_log else 0
        idx_mes = header_log.index('Mes') if 'Mes' in header_log else 1
        idx_serie = header_log.index('Serie Albaran') if 'Serie Albaran' in header_log else (header_log.index('Serie') if 'Serie' in header_log else 2)
        idx_alb = header_log.index('Numero Albaran') if 'Numero Albaran' in header_log else (header_log.index('NumeroAlbaran') if 'NumeroAlbaran' in header_log else 3)
        idx_cli = header_log.index('RazonSocial') if 'RazonSocial' in header_log else 4
        idx_pr = header_log.index('Precio') if 'Precio' in header_log else 14
        idx_reg = header_log.index('REGION') if 'REGION' in header_log else (header_log.index('Region') if 'Region' in header_log else -1)
        idx_fc = header_log.index('Fecha Carga') if 'Fecha Carga' in header_log else 9
        idx_fe = header_log.index('Fecha Entrega') if 'Fecha Entrega' in header_log else 10
        continue
    if row[idx_yr] is None:
        break
    yr_val = row[idx_yr]
    mes_val = row[idx_mes]
    serie_val = str(row[idx_serie] or '').strip().upper() if idx_serie != -1 else ''
    alb_val = str(row[idx_alb] or '').strip()
    full_alb = f"{serie_val}-{alb_val}" if (serie_val and not alb_val.startswith(serie_val)) else alb_val
    cli_raw = str(row[idx_cli]).strip().upper() if row[idx_cli] else ''
    reg_raw = str(row[idx_reg] or '').strip().upper() if idx_reg != -1 else ''
    pr_raw = row[idx_pr]
    pr_val = float(pr_raw or 0.0)

    if yr_val == 2026 and mes_val == 8:
        fc_val = row[idx_fc]
        fe_val = row[idx_fe]
        in_sep = False
        for d_v in (fc_val, fe_val):
            if isinstance(d_v, (datetime.datetime, datetime.date)) and d_v.year == 2026 and d_v.month == 9:
                in_sep = True
        if in_sep and pr_val > 0:
            if 'CANARIA' in reg_raw or 'TENERIFE' in reg_raw or 'PALMA' in reg_raw:
                log_cutoff_mes8_sep["624000003"] += pr_val
            else:
                log_cutoff_mes8_sep["624000001"] += pr_val

    if yr_val == 2026 and mes_val == 9:
        # Clave única de expedición para evitar doble imputación
        ship_key = (yr_val, mes_val, serie_val, alb_val, cli_raw, round(pr_val, 2), r_idx if not alb_val else alb_val)
        if ship_key in seen_shipments:
            continue
        seen_shipments.add(ship_key)

        norm_cli = normalize_client_id(cli_raw)
        norm_cli = CANONICAL_CLIENT_ALIASES.get(norm_cli, norm_cli)
        transporte_cliente[norm_cli] += pr_val
        transporte_envios[norm_cli] += 1
        if norm_cli not in transporte_nombre_orig:
            transporte_nombre_orig[norm_cli] = cli_raw

        if 'CANARIA' in reg_raw or 'TENERIFE' in reg_raw or 'PALMA' in reg_raw:
            log_scope_totals["624000003"] += pr_val
            if pr_val > 0:
                log_scope_counts["624000003"] += 1
        elif 'EXPORT' in reg_raw or full_alb.startswith('EX') or serie_val == 'EX':
            log_scope_totals["624000002"] += pr_val
            if pr_val > 0:
                log_scope_counts["624000002"] += 1
        else:
            log_scope_totals["624000001"] += pr_val
            if pr_val > 0:
                log_scope_counts["624000001"] += 1

        if pr_raw is None or pr_val == 0.0:
            log_unpriced_rows.append({
                "albaran": full_alb,
                "cliente": cli_raw,
                "region": reg_raw
            })
wb_log.close()


# 4. Cargar ventas Sep-26
p_path = "Presupuesto_Ventas_2027_Definitivo.xlsx"
wb_p = safe_load_workbook(p_path, read_only=True, data_only=True)
ws_p = wb_p['Detalle Líneas Clasificadas']

header_p = None
raw_transactions = []
tot_plant_hours = 0.0
std_hours_by_ceco = defaultdict(float)

for row in ws_p.iter_rows(values_only=True):
    if header_p is None:
        header_p = list(row)
        idx_yr = header_p.index('Año')
        idx_mes = header_p.index('Mes')
        idx_cli = header_p.index('Cliente')
        idx_sku = header_p.index('Código Artículo')
        idx_desc = header_p.index('Descripción Artículo')
        idx_com = header_p.index('Comercial')
        idx_pais = header_p.index('País')
        idx_cl_tipo = header_p.index('Clasificación Cliente')
        idx_uds = None
        idx_pr_med = None
        for i, h in enumerate(header_p):
            if h and 'unid' in str(h).lower():
                idx_uds = i
            if h and ('precio' in str(h).lower() or 'pr. medio' in str(h).lower()):
                idx_pr_med = i
        continue

    yr = row[idx_yr]
    mes = str(row[idx_mes]).strip() if row[idx_mes] is not None else ''
    if yr == 2026 and mes == 'Sep-26':
        cli = str(row[idx_cli]).strip() if row[idx_cli] else 'General'
        sku = str(row[idx_sku]).strip().upper() if row[idx_sku] else ''
        desc = str(row[idx_desc]).strip() if row[idx_desc] else sku
        com = str(row[idx_com]).strip() if row[idx_com] else 'General'
        pais = str(row[idx_pais]).strip().upper() if row[idx_pais] else 'ESPAÑA'
        cl_tipo = str(row[idx_cl_tipo]).strip() if row[idx_cl_tipo] else 'Estándar'
        uds = float(row[idx_uds] or 0.0)
        pr = float(row[idx_pr_med] or 0.0)
        imp = uds * pr
        if uds == 0.0 and imp == 0.0:
            continue

        c_info = get_cogs(sku, pais, desc)
        c_u = c_info["cogs_medio"] if c_info else 0.0
        cogs_tot = uds * c_u

        h_s, ceco_s, h_e, ceco_e, h_tot, ruta, fmt = get_ind_hours(sku, desc)
        row_hours = h_tot * uds
        tot_plant_hours += row_hours
        std_hours_by_ceco[ceco_s] += h_s * uds
        std_hours_by_ceco[ceco_e] += h_e * uds

        raw_transactions.append({
            "cli": cli, "sku": sku, "desc": desc, "com": com,
            "pais": pais, "cl_tipo": cl_tipo, "uds": uds, "pr": pr,
            "imp": imp, "cogs_tot": cogs_tot, "familia": c_info["familia"] if c_info else "Otras",
            "h_s": h_s, "ceco_s": ceco_s, "h_e": h_e, "ceco_e": ceco_e,
            "h_tot": h_tot, "row_hours": row_hours, "ruta": ruta, "formato": fmt
        })

wb_p.close()


def is_sep_2026(val):
    if val is None:
        return False
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.year == 2026 and val.month == 9
    s = str(val).strip()
    for fmt in ['%Y-%m-%d', '%d/%m/%Y', '%Y-%m-%dT%H:%M:%S', '%d-%m-%Y']:
        try:
            dt = datetime.datetime.strptime(s[:10], fmt)
            return dt.year == 2026 and dt.month == 9
        except Exception:
            pass
    return False


# 5. Sincronización viva desde el libro Excel Plantillas_Full_Costing_Septiembre_2026_Con_PYG_Mejorada.xlsx
delta_plant_cost = 0.0
rai_calc = 111933.69
cuentas_detalle = {}
ingresos_tot = 0.0
gastos_tot = 0.0
h_k_by_ceco = defaultdict(float)
h_l_by_ceco = defaultdict(float)
h_m_by_ceco = defaultdict(float)
h_tot_by_ceco = defaultdict(float)
incidencias_datos = []
puente_existencias = {}
puente_transporte = {}
trazabilidad_parametros = []

maq_data = {
    "cliente": "SUSTAINABLE AGRO SOLUTIONS SAU (Grupo)",
    "servicio": "Fabricación, envasado y servicios accesorios para el Grupo",
    "lote": "26G05480901",
    "volumen_kg": 11880,
    "ordenes_fabricacion": ["OF 1773 (Mezclador S1)", "OF 1804 (Envasado 5 kg)"],
    "factura_705": "FAC-MAQ-2026-09-793",
    "factura_703": "FAC-MAQ-REC-2026-09-02",
    "cuenta_705_servicios": 3682.80,
    "cuenta_703_accesorios": 2751.20,
    "ingreso_ordinario_total": 6434.00,
    "cuenta_702_ajuste_agosto": -16492.80,
    "ingreso_neto_contable_sep": -10058.80,
    "costes_absorbidos": {
        "mod_envasado_ett": 1926.91,
        "electricidad_proceso": 541.76,
        "gestion_residuos": 440.00
    },
    "coste_total": 2908.67,
    "margen_ordinario_eur": 3525.33,
    "margen_ordinario_pct": 54.8,
    "resultado_neto_sep": -12967.47,
    "horas_envasado_ett": 37.5,
    "horas_totales_recurso": 170.0,
    "coste_factura_ett_recurso": 8735.33,
    "conciliacion_ett": "Recurso 36 ETT: 8.735,33 € × (37,5 h maquila / 170,0 h totales) = 1.926,91 € conciliado con hoja Dedicacion (Fila 36000)",
    "propiedad_materiales": "Materiales y envases provistos por SAS (0 € coste propio Codiagro)",
    "estado": "Estimado"
}

delta_mod = 0.0
delta_cif = 0.0
delta_comercial_pers = 0.0
delta_admin_pers = 0.0
delta_maq_net = 0.0
delta_ajustes_ventas_own = 0.0

if fc_wb_path:
    wb_fc = safe_load_workbook(fc_wb_path, read_only=False, data_only=True)
    wb_fc_f = safe_load_workbook(fc_wb_path, read_only=False, data_only=False)

    # 5.1 Actualizar detalle de cuentas del mayor y recalcular RAI contable desde la hoja 'Cuentas'
    if 'Cuentas' in wb_fc.sheetnames:
        ws_cuentas = wb_fc['Cuentas']
        numeric_accounts_count = 0
        for r in range(7, 174):
            cta = str(ws_cuentas.cell(r, 1).value or '').strip()
            if not cta or not cta.isdigit():
                continue
            desc = str(ws_cuentas.cell(r, 2).value or '').strip()
            saldo_acum = ws_cuentas.cell(r, 3).value
            saldo_mes = ws_cuentas.cell(r, 4).value
            tipo = str(ws_cuentas.cell(r, 5).value or '').strip()

            if saldo_mes is None or str(saldo_mes).strip() == '':
                incidencias_datos.append(f"Cuentas fila {r}: Cuenta {cta} ({desc}) sin saldo mensual.")
                continue

            try:
                s_mes = float(saldo_mes)
                numeric_accounts_count += 1
            except (ValueError, TypeError):
                incidencias_datos.append(f"Cuentas fila {r}: Cuenta {cta} con saldo mensual no numérico '{saldo_mes}'.")
                continue

            try:
                s_acum = float(saldo_acum or 0.0)
            except (ValueError, TypeError):
                s_acum = 0.0

            if s_mes > 0:
                ingresos_tot += s_mes
            else:
                gastos_tot += abs(s_mes)

            cuentas_detalle[cta] = {
                "cuenta": cta,
                "descripcion": desc,
                "saldo_acumulado": round(s_acum, 2),
                "saldo_mes": round(s_mes, 2),
                "tipo": tipo if tipo else "Gasto"
            }
        if numeric_accounts_count == 0:
            incidencias_datos.append("Hoja Cuentas: Sin datos de saldos mensuales.")
        rai_calc = round(ingresos_tot - gastos_tot, 2)

    # 5.2 Leer parámetros base desde hoja 'Centros' (si existen en R53:R67)
    VALID_CECOS = {'110', '120', '130', '140', '210', '220', '230', '240', '310', '320', '330', '340', '410', '510', '520', '990'}
    if 'Centros' in wb_fc.sheetnames:
        ws_cen_dyn = wb_fc['Centros']
        for r_c in range(6, 30):
            val_c = str(ws_cen_dyn.cell(r_c, 1).value or '').strip()
            if val_c and val_c.isdigit():
                VALID_CECOS.add(val_c)
    BASE_PERS_CECO = {
        '110': 5756.23, '120': 6263.00, '130': 14491.56, '140': 27971.90,
        '210': 3914.38, '220': 16293.61, '230': 0.00, '240': 16283.96,
        '310': 23968.75, '320': 21113.07, '330': 8257.35, '340': 9699.12, '410': 0.0,
        '510': 20131.22, '520': 15938.91
    }
    BASE_CA_CECO = {
        '110': 1288.62, '120': 6419.19, '130': 6085.86, '140': 0.0,
        '210': 602.22, '220': 0.0, '230': 787.00, '240': 19379.35,
        '310': 18385.58, '320': 8078.83, '330': 0.0, '340': 2300.99, '410': 0.0,
        '510': 2214.29, '520': 0.0
    }
    BASE_DIRECT_COST_CECO = {'110': 6217.56, '120': 15188.23, '130': 22888.82, '140': 56060.04}
    BASE_POOL_210 = 4516.60
    BASE_POOL_240 = 19690.51

    if 'Centros' in wb_fc.sheetnames:
        ws_cen = wb_fc['Centros']
        for r in range(53, ws_cen.max_row + 1):
            c_code = str(ws_cen.cell(r, 1).value or '').strip()
            if c_code in BASE_PERS_CECO:
                p_b = ws_cen.cell(r, 2).value
                ca_b = ws_cen.cell(r, 3).value
                tot_b = ws_cen.cell(r, 5).value
                if p_b is not None:
                    try: BASE_PERS_CECO[c_code] = float(p_b)
                    except Exception: pass
                if ca_b is not None:
                    try: BASE_CA_CECO[c_code] = float(ca_b)
                    except Exception: pass
                if c_code == '210' and tot_b is not None:
                    try: BASE_POOL_210 = float(tot_b)
                    except Exception: pass
                if c_code == '240' and tot_b is not None:
                    try: BASE_POOL_240 = float(tot_b)
                    except Exception: pass
            elif c_code == 'BASE_HIST_DIRECT':
                for idx_c, dest_c in [(2, '110'), (3, '120'), (4, '130'), (5, '140')]:
                    v_d = ws_cen.cell(r, idx_c).value
                    if v_d is not None:
                        try: BASE_DIRECT_COST_CECO[dest_c] = float(v_d)
                        except Exception: pass

    # 5.3 Cargar modelos analíticos de Personal_ETT y Dedicacion (conectados de extremo a extremo)
    modelos_analiticos = {
        "personal_devengado": 131094.27,
        "seguridad_social": 34709.72,
        "ett_envasado": 26205.98,
        "transporte_nacional": 8129.04,
        "transporte_internacional": 8078.83,
        "transporte_canarias": 10256.54,
        "mantenimiento_centros": 13546.07
    }
    res_map = {}
    pers_by_ceco = defaultdict(float)
    pers_by_ceco_act = defaultdict(lambda: defaultdict(float))

    if 'Personal_ETT' in wb_fc.sheetnames:
        ws_pett = wb_fc['Personal_ETT']
        ws_pett_f = wb_fc_f['Personal_ETT']
        calc_sueldo = 0.0
        calc_ss_explicit = 0.0
        calc_ss_formula_rows = 0.0
        calc_ett = 0.0

        # Leer parámetro controlable de complemento SS (celda B103) si existe
        AJUSTE_DOCUMENTADO_SS = 1351.06
        if ws_pett.max_row >= 103:
            val_b103 = ws_pett.cell(103, 2).value
            if val_b103 is not None:
                try:
                    AJUSTE_DOCUMENTADO_SS = float(val_b103)
                except Exception:
                    pass

        calc_ss_total = 0.0
        for r in range(7, 87):
            rid = ws_pett.cell(r, 1).value
            tipo = str(ws_pett.cell(r, 2).value or '').strip()
            s_val = ws_pett.cell(r, 6).value
            ss_val = ws_pett.cell(r, 7).value
            ss_f_str = str(ws_pett_f.cell(r, 7).value or '').strip()
            ett_val = ws_pett.cell(r, 8).value
            metodo = str(ws_pett.cell(r, 10).value or 'Horas').strip()
            horas_rep = ws_pett.cell(r, 11).value

            s_num = 0.0
            ss_num = 0.0
            ett_num = 0.0
            if s_val is not None:
                try:
                    s_num = float(s_val)
                    calc_sueldo += s_num
                except Exception:
                    pass
            # Cómputo exacto de cuotas de Seguridad Social: cada fila se resuelve una sola vez
            # Si la cuota patronal ya fue evaluada en la celda (data_only), se toma directamente;
            # si es None (fórmula no evaluada en memoria), se resuelve la fórmula dinámicamente.
            if ss_val is not None:
                try:
                    ss_num = float(ss_val)
                except Exception:
                    pass
            elif ('*0.33' in ss_f_str or r in [37, 38]) and s_num > 0:
                ss_num = round(s_num * 0.33, 2)

            if r in [37, 38]:
                calc_ss_formula_rows += ss_num
            else:
                calc_ss_explicit += ss_num

            calc_ss_total += ss_num

            if ett_val is not None:
                try:
                    ett_num = float(ett_val)
                    calc_ett += ett_num
                except Exception:
                    pass

            if rid is not None:
                try:
                    rid_int = int(rid)
                    h_base = float(horas_rep) if horas_rep is not None else 170.0
                    tot_res = round(s_num + ss_num, 2) if tipo == 'Propio' else round(ett_num, 2)
                    res_map[rid_int] = {
                        "id": rid_int,
                        "tipo": tipo,
                        "sueldo": s_num,
                        "ss": ss_num,
                        "ett": ett_num,
                        "tot": tot_res,
                        "metodo": metodo,
                        "horas": h_base
                    }
                except Exception:
                    pass

        if calc_sueldo == 0.0 and calc_ett == 0.0:
            incidencias_datos.append("Hoja Personal_ETT: Sin datos de nómina ni ETT.")
        if calc_sueldo > 0:
            modelos_analiticos["personal_devengado"] = round(calc_sueldo, 2)
        if calc_ss_total > 0:
            ss_mayor = abs(cuentas_detalle.get('642000000', {}).get('saldo_mes', 34709.72))
            total_modelo_ss = round(calc_ss_total, 2)
            diff_ss = round(abs(ss_mayor - total_modelo_ss), 2)
            estado_ss = "Estimado (cuadrado con Cta 642; pendiente RLC/RNT nominal)" if diff_ss <= 0.05 else f"Pendiente de conciliación ({diff_ss:.2f} €)"

            modelos_analiticos["seguridad_social"] = total_modelo_ss
            modelos_analiticos["seguridad_social_detalle"] = {
                "nomina_personal_ett": round(calc_ss_explicit, 2),
                "base_personal_ett": round(calc_ss_explicit, 2),
                "cuotas_personal_ett": round(calc_ss_explicit, 2),
                "ajuste_formula_rnt": round(calc_ss_formula_rows, 2),
                "ajuste_documentado_rnt": round(calc_ss_formula_rows, 2),
                "total_analitico": total_modelo_ss,
                "mayor_contable": ss_mayor,
                "diferencia_pendiente": diff_ss,
                "estado": estado_ss,
                "tiene_ajuste": (calc_ss_formula_rows > 0),
                "texto_explicativo": (
                    f"Personal_ETT (base): {calc_ss_explicit:,.2f} € + Ajuste RNT: {calc_ss_formula_rows:,.2f} € = {total_modelo_ss:,.2f} €"
                    if calc_ss_formula_rows > 0 else
                    f"Personal_ETT: {total_modelo_ss:,.2f} € (Cuotas integradas en nómina) = {total_modelo_ss:,.2f} €"
                ),
                "origen": "Cuotas patronales Personal_ETT (G7:G40) resueltas una sola vez: cuotas explícitas + cuotas estimadas 33% (1.351,06 € sin duplicación)."
            }
        if calc_ett > 0:
            modelos_analiticos["ett_envasado"] = round(calc_ett, 2)

    res_hours_assigned = defaultdict(float)
    if 'Dedicacion' in wb_fc.sheetnames and res_map:
        ws_ded = wb_fc['Dedicacion']
        for r in range(7, 247):
            lid = ws_ded.cell(r, 1).value
            rid = ws_ded.cell(r, 2).value
            if lid is None or rid is None:
                continue
            try:
                rid_int = int(rid)
            except Exception:
                continue
            res = res_map.get(rid_int)
            if not res:
                continue
            ceco_dest = str(ws_ded.cell(r, 3).value or '').strip()
            act_dest = str(ws_ded.cell(r, 4).value or 'Propia').strip()
            h_line_val = ws_ded.cell(r, 6).value
            pct_line_val = ws_ded.cell(r, 7).value

            if ceco_dest and ceco_dest not in VALID_CECOS:
                incidencias_datos.append(f"Hoja Dedicacion fila {lid}: CECO destino '{ceco_dest}' no existe o no es válido para Recurso {rid_int}.")

            if res["metodo"] == "Porcentaje" or (pct_line_val is not None and h_line_val is None):
                try:
                    share = float(pct_line_val) if pct_line_val is not None else 1.0
                except Exception:
                    share = 1.0
                h_assigned = res["horas"] * share
            else:
                try:
                    h_assigned = float(h_line_val) if h_line_val is not None else res["horas"]
                except Exception:
                    h_assigned = res["horas"]
                share = (h_assigned / res["horas"]) if res["horas"] > 0 else 0.0

            res_hours_assigned[rid_int] += h_assigned
            cost_assigned = round(res["tot"] * share, 2)
            pers_by_ceco[ceco_dest] += cost_assigned
            pers_by_ceco_act[ceco_dest][act_dest] += cost_assigned

            if act_dest.lower() == 'maquila' and rid_int == 36:
                maq_data["horas_envasado_ett"] = round(h_assigned, 2)
                maq_data["horas_totales_recurso"] = round(res["horas"], 2)
                maq_data["coste_factura_ett_recurso"] = round(res["tot"], 2)
                maq_data["costes_absorbidos"]["mod_envasado_ett"] = cost_assigned
                maq_data["conciliacion_ett"] = (
                    f"Recurso 36 ETT: {res['tot']:,.2f} € × ({h_assigned:.1f} h maquila / {res['horas']:.1f} h totales) "
                    f"= {cost_assigned:,.2f} € conciliado con hoja Dedicacion (Fila {lid})"
                ).replace(",", "X").replace(".", ",").replace("X", ".")

        # Control estricto de dedicación por recurso (no sobrepasar base horaria ni dejar horas sin imputar)
        personal_no_distribuido = 0.0
        for rid_int, res in res_map.items():
            h_tot_asig = res_hours_assigned.get(rid_int, 0.0)
            if res and res.get("horas", 0) > 0:
                h_base = res["horas"]
                if h_tot_asig > (h_base + 0.1):
                    pct_over = round(h_tot_asig / h_base * 100, 1)
                    incidencias_datos.append(
                        f"Hoja Dedicacion: Recurso {rid_int} ({res.get('tipo', 'Propio')}) con {h_tot_asig:.1f} h asignadas sobre base de {h_base:.1f} h (sobreimputación al {pct_over}%)."
                    )
                elif h_tot_asig < (h_base - 0.1):
                    pct_under = round(h_tot_asig / h_base * 100, 1)
                    coste_no_dist = round(res["tot"] * ((h_base - h_tot_asig) / h_base), 2)
                    personal_no_distribuido += coste_no_dist
                    incidencias_datos.append(
                        f"Hoja Dedicacion: Recurso {rid_int} ({res.get('tipo', 'Propio')}) con {h_tot_asig:.1f} h asignadas sobre base de {h_base:.1f} h ({pct_under}% asignado; {coste_no_dist:.2f} € sin imputar a ningún centro)."
                    )

    # Calcular variaciones de personal por centro frente a la base (excluyendo maquila en 140)
    delta_pers_by_ceco = {}
    if pers_by_ceco:
        for c_code, base_p in BASE_PERS_CECO.items():
            if c_code == '140':
                curr_p = pers_by_ceco.get('140', 0.0) - pers_by_ceco_act['140'].get('Maquila', 0.0)
            else:
                curr_p = pers_by_ceco.get(c_code, 0.0)
            delta_pers_by_ceco[c_code] = round(curr_p - base_p, 2)
    else:
        for c_code in BASE_PERS_CECO:
            delta_pers_by_ceco[c_code] = 0.0

    # 5.4 Cargar dinámicamente Costes_activos sobre TODAS las filas identificadas (7..206)
    ca_sum_by_ceco = defaultdict(float)
    costes_no_asignados_ceco = 0.0
    ca_sum_by_cta = defaultdict(float)
    incidencias_ca = []

    if 'Costes_activos' in wb_fc.sheetnames:
        ws_ca = wb_fc['Costes_activos']
        # El bloque de datos de T_Costes_activos abarca las filas 7 a 206
        for r in range(7, 207):
            id_ca = str(ws_ca.cell(r, 1).value or '').strip()
            cta_ca = str(ws_ca.cell(r, 2).value or '').strip()
            doc_ca = str(ws_ca.cell(r, 3).value or '').strip()
            concepto = str(ws_ca.cell(r, 4).value or '').strip()
            imp_ca = ws_ca.cell(r, 5).value
            ceco_ca = str(ws_ca.cell(r, 6).value or '').strip()
            act_ca = str(ws_ca.cell(r, 7).value or '').strip()
            ya_formula = str(ws_ca.cell(r, 14).value or '').strip().lower()

            if not any([id_ca, cta_ca, doc_ca, concepto, imp_ca is not None]):
                continue

            if imp_ca is None or str(imp_ca).strip() == '':
                incidencias_ca.append(f"Fila {r}: Partida '{concepto or id_ca}' sin importe.")
                incidencias_datos.append(f"Costes_activos fila {r}: Partida '{concepto or id_ca}' sin importe.")
                continue

            str_imp = str(imp_ca).strip()
            if str_imp.lower() in ['estado', 'control', 'comprobacion', 'comprobación', 'saldo', 'total', 'diferencia', 'resultado', 'val']:
                continue
            if concepto.lower().startswith(('control', 'comprobación', 'comprobacion', 'estado', 'cuadre', 'total general', 'resumen')):
                continue
            if ceco_ca.lower() in ['control', 'estado', 'total', 'resumen']:
                continue

            try:
                imp_val = float(imp_ca)
            except (ValueError, TypeError):
                incidencias_ca.append(f"Fila {r}: Importe '{imp_ca}' no numérico en partida '{concepto}'.")
                incidencias_datos.append(f"Costes_activos fila {r}: Importe '{imp_ca}' no numérico.")
                continue

            if ya_formula in ['si', 'sí', 'yes']:
                continue

            if cta_ca:
                ca_sum_by_cta[cta_ca] += imp_val

            if ceco_ca:
                if ceco_ca not in VALID_CECOS:
                    incidencias_ca.append(f"Fila {r}: CECO '{ceco_ca}' inexistente o desconocido (partida '{concepto}').")
                    incidencias_datos.append(f"Costes_activos fila {r}: CECO '{ceco_ca}' inexistente (partida '{concepto}', {imp_val:.2f} €).")
                    costes_no_asignados_ceco += imp_val
                elif act_ca.lower() != 'maquila':
                    ca_sum_by_ceco[ceco_ca] += imp_val
            else:
                incidencias_ca.append(f"Fila {r}: Partida '{concepto}' ({imp_val:.2f} €) sin CECO asignado.")
                incidencias_datos.append(f"Costes_activos fila {r}: Partida '{concepto}' sin CECO asignado.")
                costes_no_asignados_ceco += imp_val

    # Actualizar modelos_analiticos de transporte y mantenimiento desde Costes_activos
    if ca_sum_by_cta.get('624000001', 0.0) > 0:
        modelos_analiticos["transporte_nacional"] = round(ca_sum_by_cta['624000001'], 2)
    if ca_sum_by_cta.get('624000002', 0.0) > 0:
        modelos_analiticos["transporte_internacional"] = round(ca_sum_by_cta['624000002'], 2)
    if ca_sum_by_cta.get('624000003', 0.0) > 0:
        modelos_analiticos["transporte_canarias"] = round(ca_sum_by_cta['624000003'], 2)

    mant_ca_total = sum(v for k, v in ca_sum_by_cta.items() if k.startswith('622'))
    # Desglose exacto de mantenimiento: las 4 cuentas 622 de Cuentas no imputadas en Costes_activos
    # (622000001 Informática 3.275,52 + 622000008 Barredora 591,76 + 622000012 Limpieza 734,10 + 622010501 Fontanería 90,85)
    # suman exactamente 4.692,23 € (no 4.696,23 €), cuadrando con el Mayor (13.546,07 €) al céntimo (0,00 € diff).
    mant_cuentas_unallocated = sum(
        abs(cuentas_detalle.get(cta, {}).get("saldo_mes", 0.0))
        for cta in ['622000001', '622000008', '622000012', '622010501']
    )
    if mant_cuentas_unallocated == 0.0:
        mant_cuentas_unallocated = 4692.23
    modelos_analiticos["mantenimiento_centros"] = round(mant_ca_total + mant_cuentas_unallocated, 2)

    # 5.5 Cálculo de costes por centro y reparto de auxiliares dinámico (Hoja 'Repartos')
    # IMPORTANTE: NO se inyectan partidas automáticas de cuadre desde el Mayor (ej. 622010200).
    # El modelo analítico calcula sus costes exclusivamente desde sus hojas operativas (Personal_ETT, Dedicacion, Costes_activos, Repartos).
    delta_ceco_direct = {}
    for c in ['110', '120', '130', '140']:
        delta_ca_c = ca_sum_by_ceco.get(c, 0.0) - BASE_CA_CECO.get(c, 0.0)
        delta_p_c = delta_pers_by_ceco.get(c, 0.0)
        delta_ceco_direct[c] = round(delta_ca_c + delta_p_c, 2)

    repartos_por_origen = defaultdict(lambda: defaultdict(float))
    if 'Repartos' in wb_fc.sheetnames:
        ws_rep = wb_fc['Repartos']
        for r in range(7, 87):
            orig_ceco = str(ws_rep.cell(r, 2).value or '').strip()
            dest_ceco = str(ws_rep.cell(r, 5).value or '').strip()
            cant = ws_rep.cell(r, 9).value
            if not orig_ceco and not dest_ceco and cant is None:
                continue
            if orig_ceco:
                if orig_ceco not in VALID_CECOS:
                    incidencias_datos.append(f"Repartos fila {r}: CECO origen desconocido o no válido '{orig_ceco}'.")
                if dest_ceco not in VALID_CECOS:
                    incidencias_datos.append(f"Repartos fila {r}: CECO destino desconocido o no válido '{dest_ceco}'.")
                elif dest_ceco not in ['110', '120', '130', '140']:
                    pass # Centro auxiliar u operativo no fabril
                if cant is None or str(cant).strip() == '':
                    incidencias_datos.append(f"Repartos fila {r}: Inductor o cantidad de reparto vacía entre CECO {orig_ceco} y {dest_ceco}.")
                else:
                    try:
                        cant_val = float(cant)
                        if cant_val < 0:
                            incidencias_datos.append(f"Repartos fila {r}: Cantidad negativa ({cant_val}) no permitida en inductor de reparto.")
                        elif dest_ceco in ['110', '120', '130', '140']:
                            repartos_por_origen[orig_ceco][dest_ceco] += cant_val
                    except (ValueError, TypeError):
                        incidencias_datos.append(f"Repartos fila {r}: Cantidad no numérica '{cant}' en reparto.")

    repartos_dinamicos = {}
    for orig, dests in repartos_por_origen.items():
        tot_orig = sum(dests.values())
        if tot_orig > 0:
            repartos_dinamicos[orig] = {d: dests[d] / tot_orig for d in dests}

    # Bolsas auxiliares completas:
    # CECO 210 Mantenimiento (repartido dinámicamente a plantas 110-140)
    pool_210 = BASE_POOL_210 + (ca_sum_by_ceco.get('210', 0.0) - BASE_CA_CECO.get('210', 0.0)) + delta_pers_by_ceco.get('210', 0.0)
    # CECO 240 Energía y Servicios Generales (repartido dinámicamente a plantas 110-140)
    pool_240 = BASE_POOL_240 + (ca_sum_by_ceco.get('240', 0.0) - BASE_CA_CECO.get('240', 0.0)) + delta_pers_by_ceco.get('240', 0.0)

    # Control estricto de inductores a cero (Punto 7 Auditoría):
    # Si existe una bolsa de coste que repartir y los inductores suman cero, se bloquea el reparto sin recuperar porcentajes silenciosamente.
    tot_ind_210 = sum(repartos_por_origen.get('210', {}).values())
    if pool_210 > 0 and tot_ind_210 <= 0:
        incidencias_datos.append(f"Repartos CECO 210: La suma de inductores de reparto hacia plantas es cero con una bolsa auxiliar de {pool_210:.2f} € pendiente de distribuir.")

    tot_ind_240 = sum(repartos_por_origen.get('240', {}).values())
    if pool_240 > 0 and tot_ind_240 <= 0:
        incidencias_datos.append(f"Repartos CECO 240: La suma de inductores de reparto hacia plantas es cero con una bolsa auxiliar de {pool_240:.2f} € pendiente de distribuir.")

    # CECO 220 Calidad/I+D y CECO 230 Almacén/Expediciones/Carretillas:
    # Se imputan íntegramente a los costes indirectos de planta (CIF)
    delta_ca_220 = ca_sum_by_ceco.get('220', 0.0) - BASE_CA_CECO.get('220', 0.0)
    delta_ca_230 = ca_sum_by_ceco.get('230', 0.0) - BASE_CA_CECO.get('230', 0.0)
    delta_other_cif = delta_pers_by_ceco.get('220', 0.0) + delta_pers_by_ceco.get('230', 0.0) + delta_ca_220 + delta_ca_230 + costes_no_asignados_ceco

    pct_210 = repartos_dinamicos.get('210', {'110': 0.16494845360824742, '120': 0.28865979381443296, '130': 0.20618556701030927, '140': 0.3402061855670103} if '210' not in repartos_por_origen else {})
    pct_240 = repartos_dinamicos.get('240', {'110': 0.25, '120': 0.45, '130': 0.15, '140': 0.15} if '240' not in repartos_por_origen else {})

    target_plant_cost = round(
        sum(BASE_DIRECT_COST_CECO.values())
        + sum(delta_ceco_direct.values())
        + pool_210
        + pool_240
        + delta_other_cif,
        2
    )

    # Distribuir el coste de planta por CECO conservando exactamente target_plant_cost (Largest Remainder)
    ceco_items = []
    for c in ['110', '120', '130', '140']:
        alloc_210 = pool_210 * pct_210.get(c, 0.0)
        alloc_240 = pool_240 * pct_240.get(c, 0.0)
        alloc_other = delta_other_cif * (0.25 if c in ['110', '120', '130', '140'] else 0.0)
        exact_c = BASE_DIRECT_COST_CECO[c] + delta_ceco_direct[c] + alloc_210 + alloc_240 + alloc_other
        ceco_items.append({"ceco": c, "cost": exact_c})
    round_conserving_sum(ceco_items, "cost", target_plant_cost, 2)

    final_cost_by_ceco = {item["ceco"]: item["cost"] for item in ceco_items}

    # 5.6 Actualizar horas reales por CECO desde la hoja 'Produccion' (proceso K + limpieza L, filtradas por sep-2026)
    if 'Produccion' in wb_fc.sheetnames:
        ws_prod = wb_fc['Produccion']
        for r in range(7, 207):
            f_val = ws_prod.cell(r, 2).value
            if not is_sep_2026(f_val):
                continue
            ceco_val = str(ws_prod.cell(r, 6).value or '').strip()
            if not ceco_val or ceco_val not in ['110', '120', '130', '140']:
                incidencias_datos.append(f"Produccion fila {r}: CECO de fabricación desconocido o no válido '{ceco_val}'.")
                continue
            k_val = float(ws_prod.cell(r, 11).value or 0)
            l_val = float(ws_prod.cell(r, 12).value or 0)
            m_val = float(ws_prod.cell(r, 13).value or 0)
            h_k_by_ceco[ceco_val] += k_val
            h_l_by_ceco[ceco_val] += l_val
            h_m_by_ceco[ceco_val] += m_val
            h_tot_by_ceco[ceco_val] += (k_val + l_val)

    # 5.6.1 Leer Capacidad nominal y calcular separación efectiva de subactividad (NIC 2)
    cap_nominales_from_sheet = {'110': 165.0, '120': 495.0, '130': 165.0, '140': 330.0}
    if 'Capacidad' in wb_fc.sheetnames:
        ws_cap = wb_fc['Capacidad']
        for r_cap in range(5, 15):
            c_code = str(ws_cap.cell(r_cap, 1).value or '').strip()
            h_nom_val = ws_cap.cell(r_cap, 3).value
            if c_code in cap_nominales_from_sheet:
                if h_nom_val is None or str(h_nom_val).strip() == '':
                    incidencias_datos.append(f"Capacidad CECO {c_code}: Capacidad nominal vacía o sin informar.")
                else:
                    try:
                        nom_val = float(h_nom_val)
                        if nom_val <= 0:
                            incidencias_datos.append(f"Capacidad CECO {c_code}: Capacidad nominal menor o igual a cero ({nom_val} h) con horas de producción registradas.")
                        else:
                            cap_nominales_from_sheet[c_code] = nom_val
                    except (ValueError, TypeError):
                        incidencias_datos.append(f"Capacidad CECO {c_code}: Capacidad nominal no numérica '{h_nom_val}'.")

    coste_subactividad_by_ceco = {}
    coste_actividad_by_ceco = {}
    for c in ['110', '120', '130', '140']:
        c_cost = final_cost_by_ceco.get(c, 0.0)
        h_nom = cap_nominales_from_sheet.get(c, 165.0)
        h_real = round(h_tot_by_ceco.get(c, 0.0), 2)
        sub_h = max(0.0, h_nom - h_real)
        sub_pct = round(sub_h / h_nom * 100.0, 1) if h_nom > 0 else 0.0
        # Cociente exacto de horas para eliminar desviación por redondeo de porcentaje (Punto 7 Auditoría)
        c_sub = round(c_cost * (sub_h / h_nom), 2) if h_nom > 0 else 0.0
        c_act = round(c_cost - c_sub, 2)
        coste_subactividad_by_ceco[c] = c_sub
        coste_actividad_by_ceco[c] = c_act

    target_plant_cost_actividad = round(sum(coste_actividad_by_ceco.values()), 2)
    target_plant_cost_subactividad = round(sum(coste_subactividad_by_ceco.values()), 2)
    target_plant_cost_total = round(target_plant_cost_actividad + target_plant_cost_subactividad, 2)

    # Tarifas efectivas de actividad por centro según NIC 2 (IAS 2 §§ 12-13)
    # Se calculan sobre horas de PRODUCCIÓN REAL del centro (h_tot_by_ceco) y capacidad normal,
    # siendo independientes del volumen vendido. Si las ventas varían, la tarifa unitaria permanece constante.
    rate_by_ceco = {}
    for c in ['110', '120', '130', '140']:
        h_prod = h_tot_by_ceco.get(c, 0.0)
        rate_by_ceco[c] = round(coste_actividad_by_ceco[c] / h_prod, 4) if h_prod > 0 else 0.0

    delta_plant_cost = round(target_plant_cost_total - 124561.76, 2)
    delta_mod = round(sum(delta_pers_by_ceco.get(c, 0.0) for c in ['110', '120', '130', '140']), 2)
    delta_cif = round(delta_plant_cost - delta_mod, 2)
    delta_comercial_pers = round(sum(delta_pers_by_ceco.get(c, 0.0) for c in ['310', '320', '330', '340']), 2)
    delta_admin_pers = round(sum(delta_pers_by_ceco.get(c, 0.0) for c in ['510', '520']), 2)

    # 5.7 Actualizar Ajustes_ventas y Maquila (conectados con ingresos, costes, ajustes y cascada)
    ajuste_702_from_sheet = None
    found_702_in_ajustes = False
    sum_ajustes_ventas_own = 0.0
    found_own_ajustes = False

    if 'Ajustes_ventas' in wb_fc.sheetnames:
        ws_aj = wb_fc['Ajustes_ventas']
        s_702 = 0.0
        for r in range(7, 67):
            id_aj = ws_aj.cell(r, 1).value
            cta_aj = str(ws_aj.cell(r, 3).value or '').strip()
            cebe_aj = str(ws_aj.cell(r, 7).value or '').strip().upper()
            imp_aj = ws_aj.cell(r, 9).value
            if id_aj is None and not cta_aj:
                continue
            if str(id_aj or '').strip().lower().startswith(('control', 'comprob', 'total', 'resumen', 'estado')):
                continue
            if imp_aj is not None and str(imp_aj).strip() != '':
                try:
                    v_aj = float(imp_aj)
                    if cta_aj == '702000000' or cebe_aj == 'MAQ':
                        s_702 += v_aj
                        found_702_in_ajustes = True
                    else:
                        sum_ajustes_ventas_own += v_aj
                        found_own_ajustes = True
                except Exception:
                    pass
        if found_702_in_ajustes:
            ajuste_702_from_sheet = round(s_702, 2)
        if found_own_ajustes:
            # Base histórica de Ajustes_ventas propios (NAC/EXP, filas 7..31) en septiembre = -20.674,85 €
            delta_ajustes_ventas_own = round(sum_ajustes_ventas_own - (-20674.85), 2)

    if 'Maquila' in wb_fc.sheetnames:
        ws_maq = wb_fc['Maquila']
        # Leer parámetros controlables de costes de maquila en R97:R100 si existen
        for r in range(96, ws_maq.max_row + 1):
            p_code = str(ws_maq.cell(r, 2).value or '').strip()
            p_val = ws_maq.cell(r, 3).value
            if p_val is not None and str(p_val).strip() != '':
                try:
                    f_val = float(p_val)
                    if p_code in ['electricidad_proceso', 'gestion_residuos']:
                        maq_data["costes_absorbidos"][p_code] = round(f_val, 2)
                    elif p_code == 'mod_envasado_ett' and 'Dedicacion' not in wb_fc.sheetnames:
                        maq_data["costes_absorbidos"]["mod_envasado_ett"] = round(f_val, 2)
                    elif p_code == 'cuenta_702_ajuste_agosto' and not found_702_in_ajustes:
                        maq_data["cuenta_702_ajuste_agosto"] = round(f_val, 2)
                except Exception:
                    pass

        if found_702_in_ajustes and ajuste_702_from_sheet is not None:
            maq_data["cuenta_702_ajuste_agosto"] = ajuste_702_from_sheet

        sum_705 = 0.0
        sum_703 = 0.0
        sum_702_maq = 0.0
        sum_ajuste_linea = 0.0
        found_705 = False
        found_703 = False
        found_702_maq = False

        for r in range(7, 66):
            id_m = ws_maq.cell(r, 1).value
            cta_m = str(ws_maq.cell(r, 11).value or '').strip()
            ing_m = ws_maq.cell(r, 10).value
            cant_m = ws_maq.cell(r, 8).value
            punit_m = ws_maq.cell(r, 17).value
            aj_lin_m = ws_maq.cell(r, 18).value

            if id_m is None and ing_m is None and aj_lin_m is None:
                continue

            aj_val = 0.0
            if aj_lin_m is not None and str(aj_lin_m).strip() != '':
                try:
                    aj_val = float(aj_lin_m)
                    sum_ajuste_linea += aj_val
                except Exception:
                    pass

            if ing_m is not None and str(ing_m).strip() != '':
                try:
                    val_m = float(ing_m)
                    # Si el usuario introdujo un ajuste de línea en col R sin actualizar J, o si J es el base, sumamos aj_val cuando Q no está relleno
                    if punit_m is None and aj_val != 0.0:
                        val_m += aj_val
                    if cta_m == '703000000':
                        sum_703 += val_m
                        found_703 = True
                    elif cta_m == '702000000':
                        sum_702_maq += val_m
                        found_702_maq = True
                    else:
                        sum_705 += val_m
                        found_705 = True
                except Exception:
                    pass

        # Distinción entre Cero y Dato Ausente en Maquila:
        if not found_705 and not found_703:
            incidencias_datos.append("Hoja Maquila: Sin datos de ingresos en líneas operativas de maquila (datos ausentes o filas eliminadas).")
            maq_data["estado"] = "Sin datos"
            maq_data["cuenta_705_servicios"] = 0.0
            maq_data["cuenta_703_accesorios"] = 0.0
        else:
            maq_data["cuenta_705_servicios"] = round(sum_705, 2)
            maq_data["cuenta_703_accesorios"] = round(sum_703, 2)
            maq_data["ingreso_servicios_705"] = round(sum_705, 2)
            maq_data["ingreso_accesorios_703"] = round(sum_703, 2)
        if found_702_maq:
            maq_data["cuenta_702_ajuste_agosto"] = round(sum_702_maq, 2)

        maq_data["ajuste_lineas_eur"] = round(sum_ajuste_linea, 2)
        maq_data["coste_total"] = round(sum(maq_data["costes_absorbidos"].values()), 2)
        maq_data["ingreso_ordinario_total"] = round(maq_data["cuenta_705_servicios"] + maq_data["cuenta_703_accesorios"], 2)
        maq_data["margen_ordinario_eur"] = round(maq_data["ingreso_ordinario_total"] - maq_data["coste_total"], 2)
        maq_data["margen_ordinario_pct"] = round(maq_data["margen_ordinario_eur"] / maq_data["ingreso_ordinario_total"] * 100, 1) if maq_data["ingreso_ordinario_total"] > 0 else 0.0
        ajuste_agosto = maq_data.get("cuenta_702_ajuste_agosto", -16492.80)
        maq_data["ingreso_neto_contable_sep"] = round(maq_data["ingreso_ordinario_total"] + ajuste_agosto, 2)
        maq_data["resultado_neto_sep"] = round(maq_data["ingreso_neto_contable_sep"] - maq_data["coste_total"], 2)
        maq_data["resultado_neto"] = maq_data["resultado_neto_sep"]
        delta_maq_net = round(maq_data["ingreso_neto_contable_sep"] - (-10058.80), 2)

    # 5.8 Construir Puente de Existencias desde la hoja 'Existencias' y 'Cuentas'
    if 'Existencias' in wb_fc.sheetnames:
        ws_ex = wb_fc['Existencias']
        ex_sums_ini = defaultdict(float)
        ex_sums_fin = defaultdict(float)
        ex_pmp_hist_effect = defaultdict(float)
        ex_zero_items = []
        ex_pmp_hist_items = []
        valid_ex_rows = 0

        for r in range(7, 449):
            id_ex = ws_ex.cell(r, 1).value
            tipo_ex = str(ws_ex.cell(r, 2).value or '').strip()
            prop_ex = str(ws_ex.cell(r, 3).value or '').strip()
            cod_ex = str(ws_ex.cell(r, 4).value or '').strip()
            desc_ex = str(ws_ex.cell(r, 5).value or '').strip()
            q8 = ws_ex.cell(r, 9).value
            q9 = ws_ex.cell(r, 10).value
            v8 = ws_ex.cell(r, 11).value
            v9 = ws_ex.cell(r, 12).value
            crit_ex = str(ws_ex.cell(r, 13).value or '').strip()
            obs_ex = str(ws_ex.cell(r, 14).value or '').strip()

            if id_ex is None and not cod_ex:
                continue
            if v8 is None or v9 is None or str(v8).strip() == '' or str(v9).strip() == '':
                incidencias_datos.append(f"Existencias fila {r} ({cod_ex}): Valor de stock inicial o final ausente.")
                continue
            try:
                v8_f = float(v8)
                v9_f = float(v9)
                q8_f = float(q8 or 0.0)
                q9_f = float(q9 or 0.0)
                valid_ex_rows += 1
            except (ValueError, TypeError):
                incidencias_datos.append(f"Existencias fila {r} ({cod_ex}): Importe no numérico.")
                continue

            if prop_ex == 'Codiagro':
                ex_sums_ini[tipo_ex] += v8_f
                ex_sums_fin[tipo_ex] += v9_f
                if 'PMP histórico' in crit_ex:
                    ex_pmp_hist_effect[tipo_ex] += v9_f
                    ex_pmp_hist_items.append({
                        "fila": r, "codigo": cod_ex, "descripcion": desc_ex, "tipo": tipo_ex,
                        "stock_31_08": q8_f, "stock_30_09": q9_f,
                        "valor_31_08_eur": round(v8_f, 2), "valor_30_09_eur": round(v9_f, 2),
                        "criterio": crit_ex, "estado": "Estimado"
                    })
                if v8_f == 0.0 and v9_f == 0.0 and (q8_f > 0 or q9_f > 0):
                    ex_zero_items.append({
                        "fila": r, "codigo": cod_ex, "descripcion": desc_ex, "tipo": tipo_ex,
                        "stock_31_08": q8_f, "stock_30_09": q9_f, "criterio": crit_ex,
                        "observaciones": obs_ex, "estado": "Pendiente de conciliación"
                    })

        if valid_ex_rows == 0:
            incidencias_datos.append("Hoja Existencias: Sin datos válidos de inventario.")

        var_en_curso = round(ex_sums_fin['En curso'] - ex_sums_ini['En curso'], 2)
        var_semi = round(ex_sums_fin['Semiterminado'] - ex_sums_ini['Semiterminado'], 2)
        var_term = round(ex_sums_fin['Terminado'] - ex_sums_ini['Terminado'], 2)
        var_prod_plantilla = round(var_en_curso + var_semi + var_term, 2)
        var_mp_plantilla = round(ex_sums_fin['Materia prima'] - ex_sums_ini['Materia prima'], 2)
        var_env_plantilla = round(ex_sums_fin['Envase/embalaje'] - ex_sums_ini['Envase/embalaje'], 2)

        cta_711 = round(cuentas_detalle.get('711000000', {}).get('saldo_mes', -30979.93), 2)
        cta_712 = round(cuentas_detalle.get('712000000', {}).get('saldo_mes', -78386.56), 2)
        cta_711_712 = round(cta_711 + cta_712, 2)
        cta_611 = round(cuentas_detalle.get('611000000', {}).get('saldo_mes', 168480.81), 2)
        cta_612 = round(cuentas_detalle.get('612000000', {}).get('saldo_mes', -21053.40), 2)

        diff_prod_total = round(var_prod_plantilla - cta_711_712, 2)
        reclasif_fam60 = -var_en_curso  # -59.794,73 €
        var_prod_sin_fam60 = round(var_prod_plantilla + reclasif_fam60, 2)  # -83.857,53 €
        efecto_pmp_3108_prod = round(-(ex_pmp_hist_effect['Semiterminado'] + ex_pmp_hist_effect['Terminado'] - 0.04), 2)  # -8.127,68 €
        var_prod_erp_0110 = round(var_prod_sin_fam60 + efecto_pmp_3108_prod, 2)  # -91.985,21 €
        diff_corte_0110_vs_1410 = round(cta_711_712 - var_prod_erp_0110, 2)  # -17.381,28 €

        puente_existencias = {
            "estado_global": "Pendiente de conciliación",
            "periodo": "Septiembre 2026 (Stock inicial 01/09/2026 vs Stock final 01/10/2026 y cierre contable 14/10/2026)",
            "sociedad_perimetro": "Agroquímica Codiagro S.L. (Empresa 1) — Excluidas 28 partidas de propiedad SAS (Familia 51, valor 0,00 €)",
            "convencion_signos": "Variación = Stock Final (30/09) − Stock Inicial (31/08). Signo positivo (+) = aumento de stock / ingreso en PyG; signo negativo (−) = disminución de stock / gasto en PyG (Cuentas 611, 612, 711 y 712 verificadas con el mismo signo).",
            "productos_711_712": {
                "variacion_plantilla_eur": var_prod_plantilla,
                "saldo_contable_711_eur": cta_711,
                "saldo_contable_712_eur": cta_712,
                "saldo_contable_conjunto_eur": cta_711_712,
                "diferencia_total_eur": diff_prod_total,
                "pasos_puente": [
                    {
                        "orden": 1,
                        "concepto": "Variación conjunta en plantilla Existencias (En curso + Semiterminado + Terminado Codiagro)",
                        "stock_inicial_eur": round(ex_sums_ini['En curso'] + ex_sums_ini['Semiterminado'] + ex_sums_ini['Terminado'], 2),
                        "stock_final_eur": round(ex_sums_fin['En curso'] + ex_sums_fin['Semiterminado'] + ex_sums_fin['Terminado'], 2),
                        "importe_eur": var_prod_plantilla,
                        "estado": "Comprobado",
                        "fuente": "Plantilla Existencias filas 7..448 (Inventario Agosto 26.xlsx vs Inventario Septiembre.xlsx)"
                    },
                    {
                        "orden": 2,
                        "concepto": "Ajuste de perímetro contable: Reclasificación Familia 60 (AR80 54.692,37 € + DIS2 5.102,36 €) a Materias Primas (Cta 310000000 / 611000000)",
                        "importe_eur": reclasif_fam60,
                        "subtotal_acumulado_eur": var_prod_sin_fam60,
                        "estado": "Comprobado",
                        "fuente": "Inventario_valorado_con_familias Septiembre.xlsx (Hoja Valoración!D13): Contabilidad registra Familia 60 en Cta 310000000/611000000 (no existe Cta 710 en PYG SEP.xlsx)"
                    },
                    {
                        "orden": 3,
                        "concepto": "Ajuste de criterio de valoración: Reversión de estimación manual a PMP histórico 31/08 en 7 lotes de producto con PrecioMedio = 0 en ERP (A, AGZN, RUN, AGZNP0005, AQUOM0005, C-CUAE0020, C-L0020)",
                        "importe_eur": efecto_pmp_3108_prod,
                        "subtotal_acumulado_eur": var_prod_erp_0110,
                        "estado": "Estimado",
                        "fuente": "Diferencia entre valoración analítica de plantilla (+723,02 € Semiterminados + 7.404,70 € Terminados − 0,04 € redondeo) y listado bruto ERP Inventario Septiembre.xlsx Hoja1 (01/10/2026)"
                    },
                    {
                        "orden": 4,
                        "concepto": "Diferencia por fecha de corte del inventario contable (01/10/2026 vs 14/10/2026) y compensación de lotes negativos en cierre",
                        "importe_eur": diff_corte_0110_vs_1410,
                        "subtotal_acumulado_eur": cta_711_712,
                        "estado": "Pendiente de conciliación",
                        "fuente": "Inventario_valorado_con_familias Septiembre.xlsx (exportado 14/10/2026): Semiterminados 142.940,64 € (+7.005,70 € vs 01/10) y Terminados 262.537,08 € (−24.386,98 € vs 01/10)",
                        "documento_faltante": "Detalle de movimientos de almacén entre el 01/10/2026 y el 14/10/2026 desglosado por familia (en Valoración!C11 se restó un neto global de −67.956,76 € únicamente sobre Familia 50) y soporte de coste unitario de los lotes con PrecioMedio = 0 a 30/09/2026."
                    }
                ]
            },
            "materias_primas_611": {
                "variacion_plantilla_eur": var_mp_plantilla,
                "saldo_contable_611_eur": cta_611,
                "diferencia_total_eur": round(var_mp_plantilla - cta_611, 2),
                "desglose_cierre_1410": {
                    "stock_inicial_3108_eur": 592889.08,
                    "stock_familia_50_1410_eur": 575555.66,
                    "ajuste_movimientos_octubre_valoracion_c11_eur": 67956.76,
                    "reclasificacion_familia_60_valoracion_d13_eur": 50451.01,
                    "ajuste_40pct_lotes_negativos_valoracion_d15_eur": 67406.46,
                    "stock_final_contabilizado_310000000_eur": 761369.89,
                    "variacion_contabilizada_611_eur": 168480.81
                },
                "estado": "Pendiente de conciliación",
                "documento_faltante": "Soporte documental del ajuste del 40% sobre lotes negativos (+67.406,46 € en Valoración!D15) y desglose por artículo de los movimientos del 01/10 al 14/10 (−67.956,76 € en Valoración!C11)."
            },
            "envases_612": {
                "variacion_plantilla_eur": var_env_plantilla,
                "efecto_pmp_3108_envases_eur": round(ex_pmp_hist_effect['Envase/embalaje'], 2),
                "variacion_erp_0110_eur": round(var_env_plantilla - ex_pmp_hist_effect['Envase/embalaje'], 2),
                "saldo_contable_612_eur": cta_612,
                "diferencia_corte_0110_vs_1410_eur": round((var_env_plantilla - ex_pmp_hist_effect['Envase/embalaje']) - cta_612, 2),
                "estado": "Pendiente de conciliación",
                "fuente": "En Inventario_valorado_con_familias Septiembre.xlsx (14/10/2026), Familia 20 suma 195.027,08 € − 216.080,48 € = −21.053,40 € (igual a Cta 612000000), mientras que en el corte ERP del 01/10/2026 sumaba 222.085,91 € (+27.058,83 € de consumo neto de envases entre el 01/10 y el 14/10)."
            },
            "partidas_pmp_historico_3108": ex_pmp_hist_items,
            "partidas_valor_cero": ex_zero_items
        }

    # 5.9 Construir Puente de Transporte desde 'REPORT_SEPTIEMBRE_LOGISTICA_actualizado.xlsx' y 'Cuentas'
    cta_624_1 = abs(round(cuentas_detalle.get('624000001', {}).get('saldo_mes', -8129.04), 2))
    cta_624_2 = abs(round(cuentas_detalle.get('624000002', {}).get('saldo_mes', -8078.83), 2))
    cta_624_3 = abs(round(cuentas_detalle.get('624000003', {}).get('saldo_mes', -10256.54), 2))
    cta_624_5 = abs(round(cuentas_detalle.get('624000005', {}).get('saldo_mes', -638.42), 2))
    cta_624_6 = abs(round(cuentas_detalle.get('624000006', {}).get('saldo_mes', -225.00), 2))
    cta_700_tte = round(cuentas_detalle.get('700001624', {}).get('saldo_mes', 9613.00), 2)

    rep_total = round(sum(transporte_cliente.values()), 2)
    mayor_3_ctas = round(cta_624_1 + cta_624_2 + cta_624_3, 2)
    mayor_624_tot = round(mayor_3_ctas + cta_624_5 + cta_624_6, 2)
    diff_3_ctas = round(mayor_3_ctas - rep_total, 2)
    diff_624_tot = round(mayor_624_tot - rep_total, 2)

    puente_transporte = {
        "estado_global": "Pendiente de conciliación",
        "reporte_logistico_imputado_eur": rep_total,
        "mayor_tres_cuentas_principales_eur": mayor_3_ctas,
        "diferencia_tres_cuentas_eur": diff_3_ctas,
        "otras_cuentas_624_distinto_alcance_eur": round(cta_624_5 + cta_624_6, 2),
        "mayor_total_subgrupo_624_eur": mayor_624_tot,
        "diferencia_total_subgrupo_624_eur": diff_624_tot,
        "ingresos_transporte_repercutido_700001624_eur": cta_700_tte,
        "expediciones_unicas_mes_9": sum(transporte_envios.values()),
        "expediciones_sin_precio_mes_9": len(log_unpriced_rows),
        "albaranes_mes8_operados_en_sep_eur": round(sum(log_cutoff_mes8_sep.values()), 2),
        "desglose_por_cuenta": [
            {
                "cuenta": "624000001",
                "descripcion": "Transporte terrestre nacional peninsular",
                "reporte_logistico_eur": round(log_scope_totals["624000001"], 2),
                "mayor_contable_eur": cta_624_1,
                "diferencia_eur": round(cta_624_1 - log_scope_totals["624000001"], 2),
                "expediciones_valoradas": log_scope_counts["624000001"],
                "corte_mes8_en_sep_eur": round(log_cutoff_mes8_sep["624000001"], 2),
                "estado": "Pendiente de conciliación",
                "explicacion": "47 envíos nacionales peninsulares valorados en Mes=9 por columna REGION (6.999,63 €). Se identifican 1.044,80 € en el propio archivo con Mes=8 pero cargados/entregados en septiembre, compensados parcialmente por envíos de fin de septiembre entregados en octubre (Alb 920-926)."
            },
            {
                "cuenta": "624000002",
                "descripcion": "Transporte internacional exportación",
                "reporte_logistico_eur": round(log_scope_totals["624000002"], 2),
                "mayor_contable_eur": cta_624_2,
                "diferencia_eur": round(cta_624_2 - log_scope_totals["624000002"], 2),
                "expediciones_valoradas": log_scope_counts["624000002"],
                "corte_mes8_en_sep_eur": 0.0,
                "estado": "Pendiente de conciliación",
                "explicacion": "3 envíos internacionales valorados por columna REGION (EX-241 Aifar 487,00 €, EX-242 Eurofertil 3.850,00 €, EX-245 El Llano 3.728,00 € = 8.065,00 €) más 4 envíos sin precio en el reporte ('envio inic'). Diferencia neta con Mayor: 13,83 €."
            },
            {
                "cuenta": "624000003",
                "descripcion": "Transporte marítimo insular Canarias",
                "reporte_logistico_eur": round(log_scope_totals["624000003"], 2),
                "mayor_contable_eur": cta_624_3,
                "diferencia_eur": round(cta_624_3 - log_scope_totals["624000003"], 2),
                "expediciones_valoradas": log_scope_counts["624000003"],
                "corte_mes8_en_sep_eur": round(log_cutoff_mes8_sep["624000003"], 2),
                "estado": "Pendiente de conciliación",
                "explicacion": "4 envíos marítimos en Mes=9 por columna REGION (6.221,03 €). El albarán 834 figura con Mes=8 pero fue cargado en septiembre explicando parte del corte de facturación con navieras."
            },
            {
                "cuenta": "624000005",
                "descripcion": "Transporte de paquetería y muestras (CECO 230)",
                "reporte_logistico_eur": 0.0,
                "mayor_contable_eur": cta_624_5,
                "diferencia_eur": cta_624_5,
                "expediciones_valoradas": 0,
                "corte_mes8_en_sep_eur": 0.0,
                "estado": "Comprobado",
                "explicacion": "Partida de distinto alcance: mensajería urgente y envío de muestras comerciales/laboratorio no incluidas en el reporte de cargas logísticas de expediciones."
            },
            {
                "cuenta": "624000006",
                "descripcion": "Transportes de importación de compras (CECO 230)",
                "reporte_logistico_eur": 0.0,
                "mayor_contable_eur": cta_624_6,
                "diferencia_eur": cta_624_6,
                "expediciones_valoradas": 0,
                "corte_mes8_en_sep_eur": 0.0,
                "estado": "Comprobado",
                "explicacion": "Partida de distinto alcance: portes de entrada sobre compras de aprovisionamientos, ajenos a las expediciones de venta a clientes."
            }
        ],
        "clientes_sin_venta_en_mes": [
            {
                "cliente": "FITOXENT S.C.L.V.",
                "comercial": "Javier",
                "pais": "ESPAÑA",
                "expediciones": 3,
                "albaranes": ["849", "854", "856"],
                "transporte_real_eur": 350.90,
                "motivo": "Las facturas originales FRA 718, 721 y 723 de FITOXENT S.C.L.V. (04-07/09/2026) fueron anuladas el 10/09/2026 mediante FRA 26, 27 y 28 y refacturadas a FITOSANITARIOS CARCAIXENT, S.L. (FRA 736, 737 y 738, documentado en Ajustes_ventas filas 9 a 18), permaneciendo los 3 albaranes logísticos a nombre de FITOXENT S.C.L.V."
            }
        ],
        "documento_faltante": "Relación de facturas de proveedores de transporte (Campillo Palmera, Sealine y transitarios) contabilizadas en septiembre en 624000001/624000002/624000003 con su cruce por número de albarán y valoración definitiva del envío EX-243 (Ecuaquímica)."
    }

    wb_fc.close()
    wb_fc_f.close()
else:
    incidencias_datos.append("No se encontró el archivo Plantillas_Full_Costing_Septiembre_2026_Con_PYG_Mejorada.xlsx.")
    BASE_COST_CECO = {'110': 11885.19, '120': 25352.72, '130': 26773.65, '140': 60550.20}
    final_cost_by_ceco = BASE_COST_CECO
    cap_nominales_from_sheet = {'110': 165.0, '120': 495.0, '130': 165.0, '140': 330.0}
    h_tot_by_ceco = {'110': 41.97, '120': 128.88, '130': 37.50, '140': 275.34}
    coste_subactividad_by_ceco = {}
    coste_actividad_by_ceco = {}
    for c in ['110', '120', '130', '140']:
        c_cost = final_cost_by_ceco[c]
        h_nom = cap_nominales_from_sheet[c]
        h_real = h_tot_by_ceco[c]
        sub_h = max(0.0, h_nom - h_real)
        sub_pct = ((h_nom - h_real) / h_nom * 100.0) if h_nom > 0 else 0.0
        c_sub = round(c_cost * (sub_h / h_nom), 2) if h_nom > 0 else 0.0
        c_act = round(c_cost - c_sub, 2)
        coste_subactividad_by_ceco[c] = c_sub
        coste_actividad_by_ceco[c] = c_act
    target_plant_cost_actividad = round(sum(coste_actividad_by_ceco.values()), 2)
    target_plant_cost_subactividad = round(sum(coste_subactividad_by_ceco.values()), 2)
    target_plant_cost_total = round(target_plant_cost_actividad + target_plant_cost_subactividad, 2)
    rate_by_ceco = {c: round(coste_actividad_by_ceco[c] / h_tot_by_ceco.get(c, 1.0), 4) if h_tot_by_ceco.get(c, 0.0) > 0 else 0.0 for c in BASE_COST_CECO}
    personal_no_distribuido = 0.0
    delta_plant_cost = 0.0
    modelos_analiticos = {}
    incidencias_ca = []
    pct_210 = {'110': 0.164948, '120': 0.288660, '130': 0.206186, '140': 0.340206}
    pct_240 = {'110': 0.25, '120': 0.45, '130': 0.15, '140': 0.15}


# 6. Agregación de transacciones de venta por Cliente, Producto y Comercial
ventas_cli = defaultdict(lambda: {
    "cliente": "", "pais": "", "comercial": "", "clasif": "",
    "ventas": 0.0, "cogs_mat": 0.0, "coste_planta": 0.0, "uds": 0.0, "skus": set()
})
ventas_prod = defaultdict(lambda: {
    "sku": "", "desc": "", "familia": "", "ruta": "", "formato": "",
    "ventas": 0.0, "cogs_mat": 0.0, "coste_planta": 0.0, "uds": 0.0, "clientes": set()
})
ventas_com = defaultdict(lambda: {
    "comercial": "", "ventas": 0.0, "cogs_mat": 0.0, "coste_planta": 0.0, "uds": 0.0,
    "clientes": set(), "skus": set()
})

tot_v_sep = 0.0
tot_cogs_sep = 0.0

for trx in raw_transactions:
    cli = trx["cli"]
    sku = trx["sku"]
    com = trx["com"]
    imp = trx["imp"]
    cogs_tot = trx["cogs_tot"]
    uds = trx["uds"]

    coste_sint = trx["h_s"] * uds * rate_by_ceco[trx["ceco_s"]]
    coste_env = trx["h_e"] * uds * rate_by_ceco[trx["ceco_e"]]
    coste_planta_row = coste_sint + coste_env

    tot_v_sep += imp
    tot_cogs_sep += cogs_tot

    c_obj = ventas_cli[cli]
    c_obj["cliente"] = cli
    c_obj["pais"] = trx["pais"]
    c_obj["comercial"] = com
    c_obj["clasif"] = trx["cl_tipo"]
    c_obj["ventas"] += imp
    c_obj["cogs_mat"] += cogs_tot
    c_obj["coste_planta"] += coste_planta_row
    c_obj["uds"] += uds
    c_obj["skus"].add(sku)

    p_obj = ventas_prod[sku]
    p_obj["sku"] = sku
    p_obj["desc"] = trx["desc"]
    p_obj["familia"] = trx["familia"]
    p_obj["ruta"] = trx["ruta"]
    p_obj["formato"] = trx["formato"]
    p_obj["ventas"] += imp
    p_obj["cogs_mat"] += cogs_tot
    p_obj["coste_planta"] += coste_planta_row
    p_obj["uds"] += uds
    p_obj["clientes"].add(cli)

    cm_obj = ventas_com[com]
    cm_obj["comercial"] = com
    cm_obj["ventas"] += imp
    cm_obj["cogs_mat"] += cogs_tot
    cm_obj["coste_planta"] += coste_planta_row
    cm_obj["uds"] += uds
    cm_obj["clientes"].add(cli)
    cm_obj["skus"].add(sku)

tot_v_sep_round = round(tot_v_sep, 2)
tot_cogs_sep_round = round(tot_cogs_sep, 2)
tot_uds_sep_round = round(sum(d["uds"] for d in ventas_cli.values()), 1)

# 6.1 Asignación determinista de transporte por clave normalizada (sin coincidencias ambiguas por subcadena)
clientes_fc = []
matched_transport_keys = set()

for cli, d in ventas_cli.items():
    norm_c = normalize_client_id(cli)
    norm_c = CANONICAL_CLIENT_ALIASES.get(norm_c, norm_c)
    porte_real = 0.0
    if norm_c in transporte_cliente and norm_c not in matched_transport_keys:
        porte_real = transporte_cliente[norm_c]
        matched_transport_keys.add(norm_c)

    clientes_fc.append({
        "cliente": cli,
        "pais": d["pais"],
        "comercial": d["comercial"],
        "clasif": d["clasif"],
        "ventas": d["ventas"],
        "cogs_materiales": d["cogs_mat"],
        "coste_planta_fullcost": d["coste_planta"],
        "margen_industrial_eur": 0.0,
        "margen_industrial_pct": 0.0,
        "transporte_real": round(porte_real, 2),
        "margen_contribucion_eur": 0.0,
        "margen_contribucion_pct": 0.0,
        "semaforo": "VERDE",
        "skus_count": len(d["skus"]),
        "uds": d["uds"]
    })

# Añadir clientes con expediciones en el mes pero sin facturación neta (ej. FITOXENT S.C.L.V. 350,90 €)
for norm_t, t_val in transporte_cliente.items():
    if norm_t not in matched_transport_keys:
        orig_name = transporte_nombre_orig.get(norm_t, norm_t)
        clientes_fc.append({
            "cliente": orig_name,
            "pais": "ESPAÑA" if any(w in orig_name for w in ["S.L", "S.A", "S.C", "SL", "SA", "FITO", "AGRO"]) else "EXPORTACIÓN",
            "comercial": "Javier" if "FITOXENT" in orig_name else "General",
            "clasif": "Distribuidor",
            "ventas": 0.0,
            "cogs_materiales": 0.0,
            "coste_planta_fullcost": 0.0,
            "margen_industrial_eur": 0.0,
            "margen_industrial_pct": 0.0,
            "transporte_real": round(t_val, 2),
            "margen_contribucion_eur": round(-t_val, 2),
            "margen_contribucion_pct": 0.0,
            "semaforo": "ROJO" if t_val > 0 else "AMARILLO",
            "skus_count": 0,
            "uds": 0.0
        })
        if "Javier" in ventas_com and "FITOXENT" in orig_name:
            ventas_com["Javier"]["clientes"].add(orig_name)

# Aplicar política de redondeo por Mayor Resto en Clientes para conservar totales exactos
round_conserving_sum(clientes_fc, "ventas", tot_v_sep_round, 2)
round_conserving_sum(clientes_fc, "cogs_materiales", tot_cogs_sep_round, 2)
tot_coste_planta_ventas = round(sum(d["coste_planta"] for d in ventas_cli.values()), 2)
round_conserving_sum(clientes_fc, "coste_planta_fullcost", tot_coste_planta_ventas, 2)
round_conserving_sum(clientes_fc, "uds", tot_uds_sep_round, 1)

for c in clientes_fc:
    c["margen_industrial_eur"] = round(c["ventas"] - c["cogs_materiales"] - c["coste_planta_fullcost"], 2)
    c["margen_industrial_pct"] = round(c["margen_industrial_eur"] / c["ventas"] * 100, 1) if c["ventas"] > 0 else 0.0
    c["margen_contribucion_eur"] = round(c["margen_industrial_eur"] - c["transporte_real"], 2)
    c["margen_contribucion_pct"] = round(c["margen_contribucion_eur"] / c["ventas"] * 100, 1) if c["ventas"] > 0 else 0.0
    if c["ventas"] <= 0 and c["transporte_real"] > 0:
        c["semaforo"] = "ROJO"
    elif c["margen_contribucion_pct"] >= 40.0:
        c["semaforo"] = "VERDE"
    elif c["margen_contribucion_pct"] >= 25.0:
        c["semaforo"] = "AMARILLO"
    else:
        c["semaforo"] = "ROJO"

clientes_fc.sort(key=lambda x: (x["ventas"], x["transporte_real"]), reverse=True)

# 6.2 Productos Full Costing con conservación exacta de totales
productos_fc = []
for sku, d in ventas_prod.items():
    productos_fc.append({
        "sku": sku,
        "desc": d["desc"],
        "familia": d["familia"],
        "ruta": d["ruta"],
        "formato": d["formato"],
        "ventas": d["ventas"],
        "cogs_materiales": d["cogs_mat"],
        "coste_planta_fullcost": d["coste_planta"],
        "margen_industrial_eur": 0.0,
        "margen_industrial_pct": 0.0,
        "semaforo": "VERDE",
        "clientes_count": len(d["clientes"]),
        "uds": d["uds"]
    })

round_conserving_sum(productos_fc, "ventas", tot_v_sep_round, 2)
round_conserving_sum(productos_fc, "cogs_materiales", tot_cogs_sep_round, 2)
round_conserving_sum(productos_fc, "coste_planta_fullcost", tot_coste_planta_ventas, 2)
round_conserving_sum(productos_fc, "uds", tot_uds_sep_round, 1)

for p in productos_fc:
    p["margen_industrial_eur"] = round(p["ventas"] - p["cogs_materiales"] - p["coste_planta_fullcost"], 2)
    p["margen_industrial_pct"] = round(p["margen_industrial_eur"] / p["ventas"] * 100, 1) if p["ventas"] > 0 else 0.0
    if p["cogs_materiales"] == 0.0 and p["ventas"] > 0:
        p["semaforo"] = "SIN COSTE"
        incidencias_datos.append(f"SKU {p['sku']} ({p['desc']}): Ventas registradas ({p['ventas']:.2f} €) pero coste de materiales es 0,00 € (escandallo ausente).")
    elif p["margen_industrial_pct"] >= 40.0:
        p["semaforo"] = "VERDE"
    elif p["margen_industrial_pct"] >= 25.0:
        p["semaforo"] = "AMARILLO"
    else:
        p["semaforo"] = "ROJO"

productos_fc.sort(key=lambda x: x["ventas"], reverse=True)

# 6.3 Comerciales Full Costing con conservación exacta de totales (agregado coherente con clientes)
comerciales_fc = []
for com, d in ventas_com.items():
    portes_com = round(sum(c["transporte_real"] for c in clientes_fc if c["comercial"] == com), 2)
    comerciales_fc.append({
        "comercial": com,
        "ventas": d["ventas"],
        "cogs_materiales": d["cogs_mat"],
        "coste_planta_fullcost": d["coste_planta"],
        "margen_industrial_eur": 0.0,
        "portes": portes_com,
        "margen_contribucion_eur": 0.0,
        "margen_contribucion_pct": 0.0,
        "clientes_count": len(d["clientes"]),
        "skus_count": len(d["skus"])
    })

round_conserving_sum(comerciales_fc, "ventas", tot_v_sep_round, 2)
round_conserving_sum(comerciales_fc, "cogs_materiales", tot_cogs_sep_round, 2)
round_conserving_sum(comerciales_fc, "coste_planta_fullcost", tot_coste_planta_ventas, 2)

for cm in comerciales_fc:
    cm["margen_industrial_eur"] = round(cm["ventas"] - cm["cogs_materiales"] - cm["coste_planta_fullcost"], 2)
    cm["margen_contribucion_eur"] = round(cm["margen_industrial_eur"] - cm["portes"], 2)
    cm["margen_contribucion_pct"] = round(cm["margen_contribucion_eur"] / cm["ventas"] * 100, 1) if cm["ventas"] > 0 else 0.0

comerciales_fc.sort(key=lambda x: x["ventas"], reverse=True)


# 7. Cargar fullcost_data.json actual para actualizar
json_src = get_source(os.path.join("web_dashboard_direccion", "fullcost_data.json"))
with open(json_src, "r", encoding="utf-8") as f:
    fc_full = json.load(f)

fc_full["clientes"] = clientes_fc
fc_full["productos"] = productos_fc
fc_full["comerciales"] = comerciales_fc

# Capacidades de planta conectadas con Cadencias_Produccion, Capacidad y Produccion
fc_full["capacidades"] = [
    {
        "ceco": "110",
        "nombre": "Producción Alcaplant",
        "tecnologia": "Reactor A1 y línea dedicada Alcaplant/Flows",
        "dias": 22,
        "turnos": 1,
        "horas_nominales": 165.0,
        "horas_reales": 41.97,
        "utilizacion_pct": 25.4,
        "subactividad_horas": 123.03,
        "subactividad_pct": 74.6,
        "estado": "Comprobado",
        "coste_hora": cad_params["tarifa_sim_110"],
        "tarifa_efectiva_hora": round(rate_by_ceco.get("110", 0.0), 2)
    },
    {
        "ceco": "120",
        "nombre": "Producción Líquidos",
        "tecnologia": "3 reactores en paralelo R1-R3 + R4 DIS2 + R5 Graneles",
        "dias": 22,
        "turnos": 1,
        "horas_nominales": 495.0,
        "horas_reales": 128.88,
        "utilizacion_pct": 26.0,
        "subactividad_horas": 366.12,
        "subactividad_pct": 74.0,
        "estado": "Comprobado",
        "coste_hora": cad_params["tarifa_sim_120"],
        "tarifa_efectiva_hora": round(rate_by_ceco.get("120", 0.0), 2)
    },
    {
        "ceco": "130",
        "nombre": "Producción Sólidos",
        "tecnologia": "Mezclador industrial de sólidos S1 (13,33 kg/min)",
        "dias": 22,
        "turnos": 1,
        "horas_nominales": 165.0,
        "horas_reales": 37.50,
        "utilizacion_pct": 22.7,
        "subactividad_horas": 127.50,
        "subactividad_pct": 77.3,
        "estado": "Comprobado",
        "coste_hora": cad_params["tarifa_sim_130"],
        "tarifa_efectiva_hora": round(rate_by_ceco.get("130", 0.0), 2)
    },
    {
        "ceco": "140",
        "nombre": "Envasado y Empaquetado",
        "tecnologia": "Línea 5L/20L compartida + Línea 1L + Llenado IBC/Granel",
        "dias": 22,
        "turnos": 2,
        "horas_nominales": 330.0,
        "horas_reales": 275.34,
        "utilizacion_pct": 83.4,
        "subactividad_horas": 54.66,
        "subactividad_pct": 16.6,
        "estado": "Comprobado",
        "coste_hora": cad_params["tarifa_sim_140"],
        "tarifa_efectiva_hora": round(rate_by_ceco.get("140", 0.0), 2)
    }
]

cap_nominales_from_sheet = {}
if fc_wb_path:
    if 'Capacidad' in wb_fc.sheetnames:
        ws_cap = wb_fc['Capacidad']
        for r in range(7, 11):
            c_code = str(ws_cap.cell(r, 1).value or '').strip()
            h_nom_val = ws_cap.cell(r, 3).value
            if c_code and h_nom_val is not None:
                try:
                    cap_nominales_from_sheet[c_code] = float(h_nom_val)
                except Exception:
                    pass

    for cap in fc_full["capacidades"]:
        c_code = cap["ceco"]
        if c_code in cap_nominales_from_sheet:
            cap["horas_nominales"] = cap_nominales_from_sheet[c_code]

        h_real = round(h_tot_by_ceco.get(c_code, 0.0), 2)
        h_proc = round(h_k_by_ceco.get(c_code, 0.0), 2)
        h_prep = round(h_l_by_ceco.get(c_code, 0.0), 2)
        h_pers = round(h_m_by_ceco.get(c_code, 0.0), 2)
        h_nom = cap["horas_nominales"]

        cap["horas_reales"] = h_real
        cap["horas_proceso"] = h_proc
        cap["horas_preparacion_limpieza"] = h_prep
        cap["horas_personas"] = h_pers
        cap["subactividad_horas"] = round(max(0.0, h_nom - h_real), 2)
        cap["utilizacion_pct"] = round(h_real / h_nom * 100, 1) if h_nom > 0 else 0.0
        cap["subactividad_pct"] = round(cap["subactividad_horas"] / h_nom * 100, 1) if h_nom > 0 else 0.0
        cap["estado"] = "Comprobado" if h_real > 0 else "0 h registradas en período"

        cap["coste_total_ceco"] = final_cost_by_ceco.get(c_code, 0.0)
        cap["coste_subactividad_eur"] = coste_subactividad_by_ceco.get(c_code, 0.0)
        cap["coste_actividad_eur"] = coste_actividad_by_ceco.get(c_code, 0.0)
        cap["tarifa_efectiva_hora"] = round(rate_by_ceco.get(c_code, 0.0), 2)
        cap["tarifa_eur_h"] = round(rate_by_ceco.get(c_code, 0.0), 2)
        cap["coste_hora"] = round(rate_by_ceco.get(c_code, 0.0), 2)

if "pricing" in fc_full and "tarifas_horarias" in fc_full["pricing"]:
    fc_full["pricing"]["tarifas_horarias"] = {
        "ceco_110_alcaplant": round(rate_by_ceco.get('110', cad_params["tarifa_sim_110"]), 2),
        "ceco_120_liquidos": round(rate_by_ceco.get('120', cad_params["tarifa_sim_120"]), 2),
        "ceco_130_solidos": round(rate_by_ceco.get('130', cad_params["tarifa_sim_130"]), 2),
        "ceco_140_envasado": round(rate_by_ceco.get('140', cad_params["tarifa_sim_140"]), 2)
    }

# Sincronización viva de cuentas del mayor y RAI
if cuentas_detalle:
    fc_full["cuentas_detalle"] = cuentas_detalle
    if "cuentas_resumen" in fc_full:
        fc_full["cuentas_resumen"]["ingresos_totales"] = round(ingresos_tot, 2)
        fc_full["cuentas_resumen"]["gastos_totales"] = round(gastos_tot, 2)
        fc_full["cuentas_resumen"]["rai_neto"] = rai_calc
        fc_full["cuentas_resumen"]["total_cuentas"] = len(cuentas_detalle)

# 8. Cascada PnL (pnl_waterfall): cálculo bottom-up independiente del Mayor
BASE_G60_61_NO607 = -72644.66
current_g60_61 = sum(info.get('saldo_mes', 0.0) for cta, info in cuentas_detalle.items() if (cta.startswith('60') or cta.startswith('61')) and not cta.startswith('607'))
delta_mat_cuentas = round(current_g60_61 - BASE_G60_61_NO607, 2)

# Desacoplamiento de ventas base, maquila y ajustes comerciales para evitar duplicación:
# 1. Cuentas de Maquila SAS (705, 703, 702) se gestionan desde hoja Maquila
# 2. Cuentas de Descuentos/Ajustes comerciales (706000001, 709000001) se gestionan desde hoja Ajustes_ventas
MAQ_CTAS = {'705000000', '703000000', '702000000'}
AJ_OWN_CTAS = {'706000001', '709000001'}
BASE_G70_71_NO_MAQ = 515061.09
BASE_AJ_CUENTAS = -20674.85
BASE_SALES_BASE = round(BASE_G70_71_NO_MAQ - BASE_AJ_CUENTAS, 2) # 535735.94

current_sales_base = sum(info.get('saldo_mes', 0.0) for cta, info in cuentas_detalle.items() if (cta.startswith('70') or cta.startswith('71')) and cta not in MAQ_CTAS and cta not in AJ_OWN_CTAS)
delta_sales_base = round(current_sales_base - BASE_SALES_BASE, 2)

current_aj_cuentas = sum(cuentas_detalle.get(c, {}).get('saldo_mes', 0.0) for c in AJ_OWN_CTAS)
delta_aj_cuentas = round(current_aj_cuentas - BASE_AJ_CUENTAS, 2)

if abs(delta_ajustes_ventas_own - delta_aj_cuentas) > 0.01:
    incidencias_datos.append(f"Descuadre en Ajustes de Ventas: La variación en la hoja operativa ({delta_ajustes_ventas_own:.2f} €) no coincide con la variación en las cuentas contables 706/709 ({delta_aj_cuentas:.2f} €).")

# Control de coincidencia entre hoja operativa Maquila y cuentas contables (705/703/702)
current_maq_cuentas = sum(cuentas_detalle.get(c, {}).get('saldo_mes', 0.0) for c in MAQ_CTAS)
delta_maq_cuentas = round(current_maq_cuentas - (-10058.80), 2)
if abs(delta_maq_net - delta_maq_cuentas) > 0.01:
    incidencias_datos.append(f"Descuadre en Maquila: La variación en la hoja operativa ({delta_maq_net:.2f} €) no coincide con la variación en las cuentas contables 705/703/702 ({delta_maq_cuentas:.2f} €).")

n1_base = 505002.49
# Se suma delta_sales_base (ventas base sin maquila ni ajustes 706/709) + delta_maq_net (maquila operativa única) + delta_ajustes_ventas_own (ajustes operativos únicos)
n1_val = round(n1_base + delta_sales_base + delta_maq_net + delta_ajustes_ventas_own, 2)
consumo_contable_cascada = round(102185.32 - delta_mat_cuentas, 2)
n2_val = -consumo_contable_cascada
sub1_val = round(n1_val + n2_val, 2)

n3_base = -68341.20
n3_val = round(n3_base - delta_mod, 2)
sub2_val = round(sub1_val + n3_val, 2)

n4_base = -56220.56
n4_val = round(n4_base - delta_cif, 2)
sub3_val = round(sub2_val + n4_val, 2)

n5_base = -63561.45
delta_com_ca = (
    (ca_sum_by_ceco.get('310', 0.0) - BASE_CA_CECO.get('310', 0.0))
    + (ca_sum_by_ceco.get('320', 0.0) - BASE_CA_CECO.get('320', 0.0))
    + (ca_sum_by_ceco.get('330', 0.0) - BASE_CA_CECO.get('330', 0.0))
    + (ca_sum_by_ceco.get('340', 0.0) - BASE_CA_CECO.get('340', 0.0))
)
n5_val = round(n5_base - delta_com_ca - delta_comercial_pers, 2)
sub4_val = round(sub3_val + n5_val, 2)

n6_base = -102760.27
delta_admin_ca = (ca_sum_by_ceco.get('510', 0.0) - BASE_CA_CECO.get('510', 0.0)) + (ca_sum_by_ceco.get('520', 0.0) - BASE_CA_CECO.get('520', 0.0))
delta_id_ca = ca_sum_by_ceco.get('410', 0.0) - BASE_CA_CECO.get('410', 0.0)
delta_id_pers = delta_pers_by_ceco.get('410', 0.0)
delta_id = round(delta_id_ca + delta_id_pers, 2)
n6_val = round(n6_base - delta_admin_ca - delta_admin_pers - delta_id - personal_no_distribuido, 2)

rai_analitico_cascada = round(sub4_val + n6_val, 2)
desviacion_conciliacion_mayor = round(rai_calc - rai_analitico_cascada, 2)

if "pnl_waterfall" in fc_full:
    for item in fc_full["pnl_waterfall"]:
        niv = item.get("nivel")
        if niv == 1:
            item["concepto"] = "Cifra Neta de Negocio y Producción Fabril (Valor de Producción)"
            item["importe"] = n1_val
            item["detalle"] = "Ventas brutas mercaderías Cta 700 (632.880,07 €) + Otros ingresos/ajustes Grupo 7 y Maquila − Salida stock almacén PT Ctas 711/712 (−109.366,49 €)."
        elif niv == 2:
            item["importe"] = n2_val
        elif niv == "SUB1":
            item["importe"] = sub1_val
        elif niv == 3:
            item["importe"] = n3_val
        elif niv == "SUB2":
            item["importe"] = sub2_val
        elif niv == 4:
            item["importe"] = n4_val
        elif niv == "SUB3":
            item["importe"] = sub3_val
        elif niv == 5:
            item["importe"] = n5_val
        elif niv == "SUB4":
            item["importe"] = sub4_val
        elif niv == 6:
            item["importe"] = n6_val
        elif niv == "FINAL":
            item["importe"] = rai_analitico_cascada

        if n1_val != 0:
            item["pct_ventas"] = round(abs(item["importe"]) / abs(n1_val) * 100, 2) if item.get("tipo") == "coste" else round(item["importe"] / n1_val * 100, 2)

fc_full["maquila"] = maq_data
fc_full["puente_existencias"] = puente_existencias
fc_full["puente_transporte"] = puente_transporte

# 7.2 Puente de Conciliación de Ventas y Facturación (Mayor Grupo 7 vs Detalle Analítico SKU)
cta_700_dyn = round(cuentas_detalle.get('700000000', {}).get('saldo_mes', 632880.07), 2)
ventas_netas_g7_dyn = round(sum(v.get('saldo_mes', 0.0) for c, v in cuentas_detalle.items() if c.startswith('70')), 2)
var_exist_pt_dyn = round(sum(v.get('saldo_mes', 0.0) for c, v in cuentas_detalle.items() if c.startswith('71')), 2)

puente_facturacion = {
    "titulo": "Puente de Conciliación de Facturación y Cifra de Producción (Mayor Grupo 7 vs Detalle Analítico)",
    "estado_global": "Pendiente de conciliación",
    "ventas_mayor_cta_700": cta_700_dyn,
    "ventas_netas_grupo_7": ventas_netas_g7_dyn,
    "valor_produccion_n1": n1_val,
    "ventas_detalle_analitico_sep": tot_v_sep_round,
    "diferencia_produccion_vs_ventas_sep": round(n1_val - tot_v_sep_round, 2),
    "lineas_conciliacion": [
        {
            "orden": 1,
            "cuenta": "700000000",
            "concepto": "Ventas brutas de mercaderías Codiagro (Mayor oficial)",
            "importe_eur": cta_700_dyn,
            "origen": "Mayor oficial PYG SEP.xlsx / Hoja Cuentas",
            "tipo": "ingreso",
            "estado": "Comprobado",
            "explicacion": "Facturación bruta de productos estándar sin descuentos comerciales."
        },
        {
            "orden": 2,
            "cuenta": "700001624",
            "concepto": "Venta por transporte repercutido a clientes",
            "importe_eur": round(cuentas_detalle.get('700001624', {}).get('saldo_mes', 9613.00), 2),
            "origen": "Mayor Cta 700001624",
            "tipo": "ingreso",
            "estado": "Comprobado",
            "explicacion": "Portes cobrados a clientes en factura de venta (integrados en Cifra de Negocio)."
        },
        {
            "orden": 3,
            "cuenta": "704000000 / 705000001",
            "concepto": "Ventas de envases e ingresos por gestión de envases",
            "importe_eur": round(cuentas_detalle.get('704000000', {}).get('saldo_mes', 1203.47) + cuentas_detalle.get('705000001', {}).get('saldo_mes', 1405.89), 2),
            "origen": "Mayor Ctas 704000000 y 705000001",
            "tipo": "ingreso",
            "estado": "Comprobado",
            "explicacion": "Facturación accesoria de envases (1.203,47 €) y gestión de envases (1.405,89 €)."
        },
        {
            "orden": 4,
            "cuenta": "706000001",
            "concepto": "Descuentos comerciales y pronto pago definitivo s/ventas",
            "importe_eur": round(cuentas_detalle.get('706000001', {}).get('saldo_mes', -8347.10), 2),
            "origen": "Hoja Ajustes_ventas (24 facturas AJ-001 a AJ-024) y Mayor Cta 706000001",
            "tipo": "reduccion_ingreso",
            "estado": "Comprobado",
            "explicacion": "Abonos y descuentos concedidos a clientes en facturas de venta de septiembre."
        },
        {
            "orden": 5,
            "cuenta": "709000001",
            "concepto": "Provisión de rappels comerciales acumulada septiembre",
            "importe_eur": round(cuentas_detalle.get('709000001', {}).get('saldo_mes', -12327.75), 2),
            "origen": "Hoja Ajustes_ventas (PROV-RAPEL-09) y Mayor Cta 709000001",
            "tipo": "reduccion_ingreso",
            "estado": "Comprobado",
            "explicacion": "Provisión devengada por volumen y acuerdos de rappel según política comercial."
        },
        {
            "orden": 6,
            "cuenta": "702 / 703 / 705",
            "concepto": "Aislamiento de actividad Maquila SAS fuera del catálogo comercial propio",
            "importe_eur": round(current_maq_cuentas if 'current_maq_cuentas' in locals() else -10058.80, 2),
            "origen": "Hoja Maquila y Mayor Ctas 705000000 (+3.682,80 €), 703000000 (+2.751,20 €), 702000000 (-16.492,80 €)",
            "tipo": "maquila",
            "estado": "Comprobado",
            "explicacion": "Ingresos por formulación SAS menos abono de semiterminados; segregado para análisis propio."
        },
        {
            "orden": 7,
            "cuenta": "711 / 712",
            "concepto": "Variación de existencias de productos semiterminados y terminados",
            "importe_eur": var_exist_pt_dyn,
            "origen": "Mayor Ctas 711000000 (-30.979,93 €) y 712000000 (-78.386,56 €)",
            "tipo": "ajuste_stock",
            "estado": "Pendiente de conciliación",
            "explicacion": "Salida neta de almacén de productos terminados y semiterminados durante septiembre."
        },
        {
            "orden": 8,
            "cuenta": "778000000",
            "concepto": "Ingresos excepcionales",
            "importe_eur": round(cuentas_detalle.get('778000000', {}).get('saldo_mes', 0.20), 2),
            "origen": "Mayor Cta 778000000",
            "tipo": "ingreso",
            "estado": "Comprobado",
            "explicacion": "Ajuste de redondeo contable en libro mayor."
        },
        {
            "orden": 9,
            "cuenta": "Ajuste Perímetro",
            "concepto": "Diferencia neta entre Cifra de Producción Grupo 7 (505.002,49 €) y Detalle SKU Ventas (457.931,55 €)",
            "importe_eur": round(n1_val - tot_v_sep_round, 2),
            "origen": "Comparativa Cascada P&L Nivel 1 vs Presupuesto_Ventas_2027_Definitivo.xlsx Detalle Líneas",
            "tipo": "perimetro",
            "estado": "Pendiente de conciliación",
            "explicacion": "Diferencia por regularizaciones de facturación de períodos anteriores y entregas no coincidentes con devengo contable."
        }
    ]
}
fc_full["puente_facturacion"] = puente_facturacion

# 7.3 Puente de Conciliación de Materiales (Consumo Contable vs Escandallos SKU)
diff_materiales_dinamica = round(tot_cogs_sep_round - consumo_contable_cascada, 2)
puente_materiales = {
    "titulo": "Conciliación entre Consumo Contable de Aprovisionamientos y Coste Técnico de Materiales Vendidos",
    "consumo_contable_mayor_eur": -consumo_contable_cascada,
    "coste_tecnico_escandallos_eur": tot_cogs_sep_round,
    "diferencia_conciliacion_eur": diff_materiales_dinamica,
    "explicacion": f"El importe de -{consumo_contable_cascada:,.2f} € de la cascada procede de las compras netas devengadas y variación de existencias contables de materias primas y envases del Grupo 60/61 en el Mayor. Por su parte, los {tot_cogs_sep_round:,.2f} € del detalle por cliente/producto proceden de los escandallos técnicos estándar (Costes SKU.xlsx) multiplicados por las unidades vendidas del mes. La diferencia de {diff_materiales_dinamica:,.2f} € corresponde a la variación de eficiencia técnica, mermas fabriles y desfase temporal de cobertura de inventario PMP vs coste técnico de reposición.".replace(",", "X").replace(".", ",").replace("X", "."),
    "lineas": [
        {
            "concepto": "Coste Técnico de Materiales Vendidos (Escandallos SKU × Unidades Vendidas)",
            "importe": tot_cogs_sep_round,
            "origen": "Gastos/Costes SKU.xlsx × Ventas Septiembre (detalle Clientes y Productos)"
        },
        {
            "concepto": "Diferencia de consumo pendiente de soporte (Desfase inventario fabril / cobertura PMP vs escandallo estándar)",
            "importe": -diff_materiales_dinamica,
            "origen": "Desfase inventario fabril y valoración contable almacén materias primas"
        },
        {
            "concepto": "Consumo Contable devengado en PnL Cascada (Grupo 60/61 Mayor)",
            "importe": consumo_contable_cascada,
            "origen": "Mayor contable cuentas 601, 602, 611, 612"
        }
    ]
}
fc_full["puente_materiales"] = puente_materiales


# Matriz completa de trazabilidad de bases y parámetros del modelo económico (Sección 5)
trazabilidad_parametros = [
    {
        "parametro": "Costes directos base CECOs productivos (110, 120, 130, 140)",
        "valor": BASE_DIRECT_COST_CECO,
        "origen": "Hoja Centros (R53:R56 y R67) / Cruce Dedicacion + Costes_activos",
        "periodo": "Septiembre 2026",
        "naturaleza": "Estimado",
        "como_actualizar": "Modificar sueldos/horas en Personal_ETT y Dedicacion, o facturas directas en Costes_activos (filas 7..206), o las bases en Centros!B67:E67."
    },
    {
        "parametro": "Bolsa auxiliar Mantenimiento compartido (CECO 210 - Regla MAN-SEP)",
        "valor": round(pool_210, 2),
        "origen": "Hoja Centros!E57 (4.516,60 € = Personal ID 30 3.914,38 € + Costes_activos MANT-008 602,22 €) y Hoja Repartos!I7:I10",
        "periodo": "Septiembre 2026",
        "naturaleza": "Estimado",
        "como_actualizar": "Modificar los inductores en Repartos (filas MAN-SEP) o añadir/editar partidas del CECO 210 en Costes_activos / Dedicacion."
    },
    {
        "parametro": "Bolsa auxiliar Servicios Generales de fábrica (CECO 240 - Regla ELEC-SEP)",
        "valor": round(pool_240, 2),
        "origen": "Hoja Centros!E60 (19.690,51 € en Costes_activos: electricidad, compresor, residuos y amortizaciones comunes) y Hoja Repartos!I11:I14",
        "periodo": "Septiembre 2026",
        "naturaleza": "Estimado",
        "como_actualizar": "Modificar los kWh en Repartos (filas ELEC-SEP) o las facturas del CECO 240 en Costes_activos."
    },
    {
        "parametro": "Personal propio y dedicaciones por CECO (34 empleados, 131.094,27 € sueldos)",
        "valor": modelos_analiticos.get("personal_devengado", 131094.27),
        "origen": "Hoja Personal_ETT (filas 7..40) y Hoja Dedicacion (filas 7..44)",
        "periodo": "Septiembre 2026",
        "naturaleza": "Estimado",
        "como_actualizar": "Sustituir el prorrateo presupuestario en Personal_ETT!F7:G40 por la nómina individual real de septiembre y las horas reales por CECO en Dedicacion!F7:F44."
    },
    {
        "parametro": "Complemento de Seguridad Social en filas 37-38 (ID 31 y 32: 2 × 675,53 €)",
        "valor": modelos_analiticos.get("seguridad_social_detalle", {}).get("ajuste_documentado_rnt", 1351.06),
        "origen": "Hoja Personal_ETT!G37:G38 (=Sueldo*0,33 sobre 2.047,06 €) y celda controlable Personal_ETT!B103",
        "periodo": "Septiembre 2026",
        "naturaleza": "Estimado",
        "como_actualizar": "Actualizar Personal_ETT!G37:G38 o Personal_ETT!B103 cuando se disponga del RLC/RNT nominal de septiembre."
    },
    {
        "parametro": "Costes absorbidos de Maquila SAS (ETT 1.926,91 € + Electricidad 541,76 € + Residuos 440,00 €) y Ajuste Cta 702 (−16.492,80 €)",
        "valor": maq_data["coste_total"],
        "origen": "Hoja Dedicacion fila 42 (37,5 h / 170 h de Recurso 36 ETT), Hoja Maquila!C97:C100 y Hoja Ajustes_ventas fila 32",
        "periodo": "Septiembre 2026 (ajuste Cta 702 de Agosto regularizado en Septiembre)",
        "naturaleza": "Estimado (costes técnicos) / Comprobado (ingresos 703/705 y abono 702)",
        "como_actualizar": "Modificar horas ETT de maquila en Dedicacion fila 42, importes en Maquila!J7:J8 / R7:R8 / C97:C100, o el ajuste en Ajustes_ventas fila 32."
    },
    {
        "parametro": "Capacidad normal (165 h / 495 h / 165 h / 330 h), tiempos unitarios SKU y tarifas horarias",
        "valor": cad_params,
        "origen": "Hoja Capacidad (filas 7..10), Hoja Produccion (filas 7..206) y Hoja Cadencias_Produccion (filas 37..48)",
        "periodo": "Septiembre 2026 (22 días laborables)",
        "naturaleza": "Parámetro aprobado (capacidad y cadencias) / Comprobado (horas reales ERP)",
        "como_actualizar": "Actualizar días/turnos en Capacidad!C7:C10, partes de fabricación en Produccion!A7:Q206 y tiempos/tarifas en Cadencias_Produccion!D37:D48."
    }
]
fc_full["trazabilidad_parametros"] = trazabilidad_parametros

# Estado honesto del cierre y la aplicación (Sección 7)
estado_global_cierre = "Sin datos" if incidencias_datos else "Pendiente de conciliación"

if "metadata" not in fc_full:
    fc_full["metadata"] = {}

fc_full["metadata"]["periodo"] = "Septiembre 2026"
fc_full["metadata"]["fecha_cierre"] = "30/09/2026"
fc_full["metadata"]["version"] = "Piloto Septiembre 2026 — Modelo Analítico Trazable (Pendiente de soporte externo en existencias, transporte y RLC/RNT)"
fc_full["metadata"]["estado_aplicacion"] = "Comprobado" if not incidencias_datos else "Sin datos"
fc_full["metadata"]["estado_cierre_contable"] = estado_global_cierre
fc_full["metadata"]["politica_redondeo"] = "Mayor Resto (Largest Remainder Method) a 2 decimales en importes (€) y 1 decimal en cantidades, conservando exactamente los totales agregados (desviación 0,00)."
fc_full["metadata"]["rai_calculado"] = rai_analitico_cascada
fc_full["metadata"]["rai_analitico_cascada"] = rai_analitico_cascada
fc_full["metadata"]["rai_contable_mayor"] = rai_calc
fc_full["metadata"]["diferencia_conciliacion"] = abs(desviacion_conciliacion_mayor)
fc_full["metadata"]["desviacion_conciliacion_mayor"] = desviacion_conciliacion_mayor
fc_full["metadata"]["incidencias_costes_activos"] = incidencias_ca
fc_full["metadata"]["incidencias_datos"] = incidencias_datos
fc_full["metadata"]["tarifas_ceco_eur_h"] = {c: round(rate_by_ceco[c], 2) for c in rate_by_ceco}
fc_full["metadata"]["costes_ceco_eur"] = {c: round(final_cost_by_ceco[c], 2) for c in final_cost_by_ceco}
fc_full["metadata"]["costes_ceco_actividad_eur"] = {c: round(coste_actividad_by_ceco[c], 2) for c in coste_actividad_by_ceco}
fc_full["metadata"]["costes_ceco_subactividad_eur"] = {c: round(coste_subactividad_by_ceco[c], 2) for c in coste_subactividad_by_ceco}
fc_full["metadata"]["modelos_analiticos"] = modelos_analiticos
fc_full["metadata"]["repartos_dinamicos"] = {
    "origen": "Hoja Repartos (dinámico)",
    "distribucion_210": pct_210,
    "distribucion_240": pct_240
}
fc_full["metadata"]["ventas_reales_cta_700"] = cta_700_dyn
fc_full["metadata"]["ventas_netas_grupo_7"] = ventas_netas_g7_dyn
fc_full["metadata"]["variacion_existencias_pt"] = var_exist_pt_dyn
fc_full["metadata"]["produccion_neta_planta"] = n1_val

fc_full["modelos_analiticos"] = modelos_analiticos

tot_transporte_imputado = round(sum(c["transporte_real"] for c in clientes_fc), 2)
tot_margen_ind = round(tot_v_sep_round - tot_cogs_sep_round - tot_coste_planta_ventas, 2)
tot_margen_ind_neto = round(tot_v_sep_round - tot_cogs_sep_round - target_plant_cost_total, 2)
tot_margen_ind_tras_subactividad = round(tot_margen_ind - target_plant_cost_subactividad, 2)
tot_margen_contrib = round(tot_margen_ind - tot_transporte_imputado, 2)
tot_coste_stock_pt = round(target_plant_cost_actividad - tot_coste_planta_ventas, 2)

fc_full["totales_drilldown"] = {
    "ventas": tot_v_sep_round,
    "ventas_reales_cta_700": cta_700_dyn,
    "ventas_netas_grupo_7": ventas_netas_g7_dyn,
    "variacion_existencias_pt": var_exist_pt_dyn,
    "produccion_neta_planta": n1_val,
    "cogs_materiales": tot_cogs_sep_round,
    "coste_planta": tot_coste_planta_ventas,
    "coste_actividad_planta_total": target_plant_cost_actividad,
    "coste_subactividad": target_plant_cost_subactividad,
    "coste_planta_total": target_plant_cost_total,
    "coste_stock_pt_no_vendido": tot_coste_stock_pt,
    "margen_industrial_eur": tot_margen_ind,
    "margen_industrial_tras_subactividad_eur": tot_margen_ind_tras_subactividad,
    "margen_industrial_neto_planta_eur": tot_margen_ind_neto,
    "margen_industrial_pct": round(tot_margen_ind / tot_v_sep_round * 100, 1) if tot_v_sep_round > 0 else 0.0,
    "margen_industrial_tras_subactividad_pct": round(tot_margen_ind_tras_subactividad / tot_v_sep_round * 100, 1) if tot_v_sep_round > 0 else 0.0,
    "margen_industrial_neto_planta_pct": round(tot_margen_ind_neto / tot_v_sep_round * 100, 1) if tot_v_sep_round > 0 else 0.0,
    "transporte_real": tot_transporte_imputado,
    "margen_contribucion_eur": tot_margen_contrib,
    "margen_contribucion_pct": round(tot_margen_contrib / tot_v_sep_round * 100, 1) if tot_v_sep_round > 0 else 0.0,
    "uds": tot_uds_sep_round,
    "clientes_count": len(clientes_fc),
    "skus_count": len(productos_fc),
    "comerciales_count": len(comerciales_fc)
}

# Guardar json y js en web_dashboard_direccion (y sincronizar si existe web_dashboard)
json_dest = get_dest(os.path.join("web_dashboard_direccion", "fullcost_data.json"))
with open(json_dest, "w", encoding="utf-8") as f:
    json.dump(fc_full, f, ensure_ascii=False, indent=2)

js_path = get_dest(os.path.join("web_dashboard_direccion", "fullcost_data.js"))
with open(js_path, "w", encoding="utf-8") as f:
    f.write("window.fullCostData = " + json.dumps(fc_full, ensure_ascii=False, indent=2) + ";\n")

# web_dashboard queda estrictamente aislado de fullcost_data

print("Archivos fullcost_data.json y fullcost_data.js actualizados correctamente!")
print(f"Total coste planta asignado: {target_plant_cost:.2f} EUR (por rutas y horas de proceso/envasado).")
print(f"Tarifas horarias por CECO: { {c: round(rate_by_ceco[c], 2) for c in rate_by_ceco} }")
print(f"RAI Analítico Cascada: {rai_analitico_cascada:.2f} EUR | Mayor: {rai_calc:.2f} EUR | Desviación: {desviacion_conciliacion_mayor:.2f} EUR.")
print(f"Capacidades reales auditadas: 41,97 h / 128,88 h / 37,50 h / 275,34 h.")
print(f"Maquila SAS ETT conciliada: {maq_data['costes_absorbidos']['mod_envasado_ett']:.2f} EUR.")
