import openpyxl, os, sys
import unicodedata

def norm(s):
    if not isinstance(s, str): return ''
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn').upper().strip()

from generar_excel_informe_cambios_q4 import all_modified_lines

wb = openpyxl.load_workbook('Presupuesto_Ventas_2027_Original_Sin_Aplanar.xlsx', data_only=True)
ws = wb['Previsión Matriz Horizontal'] if 'Previsión Matriz Horizontal' in wb.sheetnames else wb[wb.sheetnames[2]]

print(f"Checking {len(all_modified_lines)} modified lines:")

for ch in all_modified_lines:
    ch_com = norm(ch['comercial'])
    ch_cli = norm(ch['cliente'])
    ch_sku = str(ch['sku']).strip().upper()
    
    matches = []
    for r in range(2, ws.max_row+1):
        com = norm(ws.cell(r, 1).value or '')
        cli = norm(ws.cell(r, 3).value or '')
        sku = str(ws.cell(r, 5).value or '').strip().upper()
        pais = str(ws.cell(r, 2).value or '').strip().upper()
        
        sku_clean = sku
        if pais != 'ESPAÑA' and len(sku) > 2:
            sku_clean = sku[:-2]
            
        if (ch_cli == cli or ch_cli in cli or cli in ch_cli) and (ch_com in com or com in ch_com):
            if ch_sku == sku or ch_sku == sku_clean:
                matches.append((r, sku, ws.cell(r, 28).value, ws.cell(r, 30).value, ws.cell(r, 32).value))
                
    com_str = ch['comercial'][:7]
    cli_str = ch['cliente'][:25]
    sku_str = ch['sku'][:10]
    print(f"{com_str:7s} | {cli_str:25s} | {sku_str:10s} | Found {len(matches)}: {matches}")
