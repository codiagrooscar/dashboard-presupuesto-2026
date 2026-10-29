import openpyxl

wb_src = openpyxl.load_workbook('scratch/fullcost_author/checks_stage.xlsx', data_only=False)
wb_dst = openpyxl.load_workbook('outputs/fullcost_septiembre_2026/Plantillas_Full_Costing_Septiembre_2026_Con_PYG_Mejorada.xlsx', data_only=False)

def copy_cells(src_ws, dst_ws, r_start, r_end, c_start, c_end, replace_126_with_448=False):
    for r in range(r_start, r_end + 1):
        for c in range(c_start, c_end + 1):
            cell_src = src_ws.cell(r, c)
            if cell_src.value is not None:
                val = cell_src.value
                if replace_126_with_448 and isinstance(val, str):
                    val = val.replace('126', '448')
                cell_dst = dst_ws.cell(r, c)
                cell_dst.value = val

print("1. Restoring Inicio H24:I31 (Tolerancias)...")
copy_cells(wb_src['Inicio'], wb_dst['Inicio'], 24, 31, 8, 9)

print("2. Restoring Cuentas checks (E181:E185, E193:E201) and Col L (checks duplicados)...")
copy_cells(wb_src['Cuentas'], wb_dst['Cuentas'], 174, 201, 1, 6)
copy_cells(wb_src['Cuentas'], wb_dst['Cuentas'], 6, 173, 12, 12)

print("3. Restoring Centros checks (E43:E48) and Cols H & I...")
copy_cells(wb_src['Centros'], wb_dst['Centros'], 36, 48, 1, 6)
copy_cells(wb_src['Centros'], wb_dst['Centros'], 6, 31, 8, 9)

print("4. Restoring Personal_ETT checks (E94:E100) and Cols O-S...")
copy_cells(wb_src['Personal_ETT'], wb_dst['Personal_ETT'], 87, 100, 1, 6)
copy_cells(wb_src['Personal_ETT'], wb_dst['Personal_ETT'], 6, 86, 15, 19)

print("5. Restoring Capacidad checks (E18:E20, F28:F31) and Cols L-M...")
copy_cells(wb_src['Capacidad'], wb_dst['Capacidad'], 11, 31, 1, 6)
copy_cells(wb_src['Capacidad'], wb_dst['Capacidad'], 6, 10, 12, 13)

print("6. Restoring Repartos checks (E94:E97, K105:K124) and Cols O-P...")
copy_cells(wb_src['Repartos'], wb_dst['Repartos'], 87, 124, 1, 11)
copy_cells(wb_src['Repartos'], wb_dst['Repartos'], 6, 86, 15, 16)

print("7. Restoring Dedicacion auxiliary check Cols O-P...")
copy_cells(wb_src['Dedicacion'], wb_dst['Dedicacion'], 6, 258, 15, 16)

print("8. Restoring Produccion check block and auxiliary check Cols R-U...")
copy_cells(wb_src['Produccion'], wb_dst['Produccion'], 207, 262, 1, 8, replace_126_with_448=True)
copy_cells(wb_src['Produccion'], wb_dst['Produccion'], 6, 206, 18, 21, replace_126_with_448=True)

print("9. Restoring Costes_activos check block and auxiliary Cols Q-T...")
copy_cells(wb_src['Costes_activos'], wb_dst['Costes_activos'], 207, 409, 1, 8)
copy_cells(wb_src['Costes_activos'], wb_dst['Costes_activos'], 6, 206, 17, 20)

print("10. Restoring Maquila check block and auxiliary Cols Q-T...")
copy_cells(wb_src['Maquila'], wb_dst['Maquila'], 66, 92, 1, 8)
copy_cells(wb_src['Maquila'], wb_dst['Maquila'], 6, 65, 17, 20)

print("11. Restoring Ajustes_ventas check block...")
copy_cells(wb_src['Ajustes_ventas'], wb_dst['Ajustes_ventas'], 87, 113, 1, 8)

print("12. Creating Existencias check panel at rows 460-468...")
ws_ex = wb_dst['Existencias']
ws_ex['A460'] = 'Checks de esta hoja'
ws_ex['A461'] = 'OK: comprobado. ERROR: descuadre. PENDIENTE/SIN DATOS: falta evidencia. REVISAR: posible incidencia.'
ws_ex['A463'] = 'Comprobación'
ws_ex['B463'] = 'Esperado'
ws_ex['C463'] = 'Registrado'
ws_ex['D463'] = 'Diferencia'
ws_ex['E463'] = 'Estado'
ws_ex['F463'] = 'Qué revisar'

ws_ex['A464'] = 'Filas pendientes'
ws_ex['B464'] = 0
ws_ex['C464'] = '=COUNTIF($O$7:$O$448,"PENDIENTE")'
ws_ex['D464'] = '=IF(AND(ISNUMBER(B464),ISNUMBER(C464)),C464-B464,"")'
ws_ex['E464'] = '=IF(NOT(SUMPRODUCT(--($O$7:$O$448<>""))>0),"SIN DATOS",IF(OR(NOT(ISNUMBER(B464)),NOT(ISNUMBER(C464))),"PENDIENTE",IF(ABS(D464)<=0,"OK","PENDIENTE")))'
ws_ex['F464'] = 'Completa los campos indicados en Estado.'

ws_ex['A465'] = 'Filas a revisar'
ws_ex['B465'] = 0
ws_ex['C465'] = '=COUNTIF($O$7:$O$448,"REVISAR")+COUNTIF($O$7:$O$448,"REVISAR clave repetida")'
ws_ex['D465'] = '=IF(AND(ISNUMBER(B465),ISNUMBER(C465)),C465-B465,"")'
ws_ex['E465'] = '=IF(NOT(SUMPRODUCT(--($O$7:$O$448<>""))>0),"SIN DATOS",IF(OR(NOT(ISNUMBER(B465)),NOT(ISNUMBER(C465))),"PENDIENTE",IF(ABS(D465)<=0,"OK","REVISAR")))'
ws_ex['F465'] = 'Comprueba identificadores, claves y tipos.'

ws_ex['A466'] = 'Existencias MP'
ws_ex['B466'] = '=SUMIF(\'Cuentas\'!$A$7:$A$173,"611000000",\'Cuentas\'!$D$7:$D$173)'
ws_ex['C466'] = '=B454'
ws_ex['D466'] = '=IF(AND(ISNUMBER(B466),ISNUMBER(C466)),C466-B466,"")'
ws_ex['E466'] = '=IF(NOT(COUNTIFS($B$7:$B$448,"Materia prima",$C$7:$C$448,"Codiagro")>0),"SIN DATOS",IF(OR(NOT(ISNUMBER(B466)),NOT(ISNUMBER(C466))),"PENDIENTE",IF(ABS(D466)<=\'Inicio\'!$I$26,"OK","REVISAR")))'
ws_ex['F466'] = 'Puente de variación materias primas.'

ws_ex['A467'] = 'Otros aprovisionam.'
ws_ex['B467'] = '=SUMIF(\'Cuentas\'!$A$7:$A$173,"612000000",\'Cuentas\'!$D$7:$D$173)'
ws_ex['C467'] = '=B455'
ws_ex['D467'] = '=IF(AND(ISNUMBER(B467),ISNUMBER(C467)),C467-B467,"")'
ws_ex['E467'] = '=IF(NOT(COUNTIFS($B$7:$B$448,"Envase/embalaje",$C$7:$C$448,"Codiagro")>0),"SIN DATOS",IF(OR(NOT(ISNUMBER(B467)),NOT(ISNUMBER(C467))),"PENDIENTE",IF(ABS(D467)<=\'Inicio\'!$I$26,"OK","REVISAR")))'
ws_ex['F467'] = 'Puente de variación envases/embalajes.'

ws_ex['A468'] = 'Stock de producción'
ws_ex['B468'] = '=SUMIF(\'Cuentas\'!$A$7:$A$173,"710000000",\'Cuentas\'!$D$7:$D$173)+SUMIF(\'Cuentas\'!$A$7:$A$173,"711000000",\'Cuentas\'!$D$7:$D$173)+SUMIF(\'Cuentas\'!$A$7:$A$173,"712000000",\'Cuentas\'!$D$7:$D$173)'
ws_ex['C468'] = '=B456'
ws_ex['D468'] = '=IF(AND(ISNUMBER(B468),ISNUMBER(C468)),C468-B468,"")'
ws_ex['E468'] = '=IF(NOT(COUNTIFS($B$7:$B$448,"Terminado",$C$7:$C$448,"Codiagro")>0),"SIN DATOS",IF(OR(NOT(ISNUMBER(B468)),NOT(ISNUMBER(C468))),"PENDIENTE",IF(ABS(D468)<=\'Inicio\'!$I$26,"OK","REVISAR")))'
ws_ex['F468'] = 'Puente de variación productos.'

print("13. Updating Inicio row 76 to point to Existencias checks at E464:E468...")
ws_ini = wb_dst['Inicio']
ws_ini['B76'] = '=COUNTIF(\'Existencias\'!$E$464:$E$468,"ERROR")'
ws_ini['C76'] = '=COUNTIF(\'Existencias\'!$E$464:$E$468,"PENDIENTE")+COUNTIF(\'Existencias\'!$E$464:$E$468,"SIN DATOS")'
ws_ini['D76'] = '=COUNTIF(\'Existencias\'!$E$464:$E$468,"REVISAR")'
ws_ini['E76'] = '=IF(B76>0,"ERROR",IF(C76>0,"PENDIENTE",IF(D76>0,"REVISAR","OK")))'
ws_ini['F76'] = 'F5 y A460'

out_path = 'outputs/fullcost_septiembre_2026/Plantillas_Full_Costing_Septiembre_2026_Con_PYG_Mejorada.xlsx'
wb_dst.save(out_path)
print(f"Successfully saved {out_path}!")
